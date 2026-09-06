#!/usr/bin/env python3
"""
merge_and_build.py - Merge translated pages and build final outputs
Combines original steps 4-7: merge -> HTML -> TOC -> DOCX/EPUB/PDF

Usage: merge_and_build.py --temp-dir <path> [--title <title>] [--author <author>] [--lang <lang>]
"""

import os
import sys
import re
import glob
import shutil
import subprocess
import tempfile
import zipfile
import argparse
import html as _html_lib
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path

import json

import algorithm_float
import arxiv_backend
import glossary
import math_guard
import layout
import chromium_pdf
import equation_fit
from manifest import read_output_text, validate_for_merge

# Windows consoles default to a legacy codepage (e.g. cp949), which raises
# UnicodeEncodeError on em-dashes and CJK in our own progress output. Force
# UTF-8 so a real error is never masked by an encoding traceback.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding='utf-8', errors='replace')
    except (AttributeError, ValueError, OSError):
        pass


# =============================================================================
# Image structure validation helpers
# =============================================================================

# Markdown image: `![alt](url)` or `![alt](url "title")`.
# - Negative lookbehind on `\` skips escaped `\![...]` (renders as literal text).
# - Closing `)` is required — a missing `)` means the image won't render, so
#   such a fragment must NOT count as a preserved image reference.
_MD_IMG_RE = re.compile(r'(?<!\\)!\[[^\]]*\]\(\s*([^)\s]+)[^)]*\)')
_VALID_ATTR_NAME_RE = re.compile(r'^[A-Za-z_][A-Za-z0-9_:.\-]*$')


class _ImgTagCollector(HTMLParser):
    """Collects every <img> tag found in fed text. Uses stdlib HTMLParser, which
    correctly handles `>` inside quoted attribute values — unlike a plain
    `<img\\b[^>]*>` regex, which would truncate at the first `>`."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.records = []  # list of (raw_tag_text, attrs_list)

    def handle_starttag(self, tag, attrs):
        if tag == 'img':
            self.records.append((self.get_starttag_text(), list(attrs)))

    handle_startendtag = handle_starttag


def _scan_img_tags(text):
    """Return (Counter of <img> srcs, list of (raw_tag, bad_attr_name) tuples).

    Feeds the entire text to HTMLParser rather than pre-extracting tags via regex,
    so quoted attribute values containing `>` are handled correctly."""
    src_counts = Counter()
    bad_attrs = []
    parser = _ImgTagCollector()
    try:
        parser.feed(text)
        parser.close()
    except Exception as e:
        bad_attrs.append(('<unparseable input>', f'<parser error: {e}>'))
        return src_counts, bad_attrs
    for raw_tag, attrs in parser.records:
        for name, _ in attrs:
            if not _VALID_ATTR_NAME_RE.match(name):
                bad_attrs.append((raw_tag, name))
        for name, val in attrs:
            if name == 'src' and val:
                src_counts[val] += 1
    return src_counts, bad_attrs


def _scan_image_refs(text):
    """Return (Counter html_srcs, Counter md_srcs, list bad_attrs)."""
    html_srcs, bad_attrs = _scan_img_tags(text)
    md_srcs = Counter(_MD_IMG_RE.findall(text))
    return html_srcs, md_srcs, bad_attrs


def _validate_chunk_images(temp_dir):
    """Verify each output_chunk*.md preserves the image structure of its chunk*.md.

    Bad-attribute detection uses a per-chunk DELTA: a malformed <img> attribute
    is flagged only if it appears in the output chunk but not in the source
    chunk. This avoids false positives on code blocks that legitimately contain
    deliberately-broken <img> examples — both chunks carry the same example, so
    the delta is empty.

    Returns False on any divergence; collects all errors and prints them
    together so an agent can fix many chunks in one pass.
    """
    temp_path = Path(temp_dir)
    errors = []
    for src_chunk in sorted(temp_path.glob('chunk*.md')):
        if src_chunk.name.startswith('output_'):
            continue
        out_chunk = temp_path / f'output_{src_chunk.name}'
        if not out_chunk.exists():
            continue  # missing-output is the manifest validator's job
        src_html, src_md, src_bad = _scan_image_refs(src_chunk.read_text(encoding='utf-8'))
        out_html, out_md, out_bad = _scan_image_refs(out_chunk.read_text(encoding='utf-8'))

        src_bad_counts = Counter(name for _, name in src_bad)
        out_bad_counts = Counter(name for _, name in out_bad)
        new_bad_counts = out_bad_counts - src_bad_counts
        if new_bad_counts:
            new_bad_examples = [
                (raw_tag, attr_name)
                for raw_tag, attr_name in out_bad
                if new_bad_counts.get(attr_name, 0) > 0
            ]
            for raw_tag, attr_name in new_bad_examples:
                errors.append(
                    f"ERROR: {out_chunk.name} introduced malformed <img> tag (not present in source)\n"
                    f"  tag: {raw_tag}\n"
                    f"  problem: attribute name '{attr_name}' is not a valid HTML identifier\n"
                    f"  likely cause: an unescaped quote inside alt=\"...\" or title=\"...\" closed the attribute early\n"
                    f"  fix: in {out_chunk.name}, replace the inner quote with a curly quote in the target language or with &quot; / &#39;\n"
                    f"  source chunk for reference: {src_chunk.name}"
                )

        if src_html != out_html or src_md != out_md:
            errors.append(
                f"ERROR: {out_chunk.name} image references diverge from {src_chunk.name}\n"
                f"  missing <img src> (count): {sorted((src_html - out_html).items()) or 'none'}\n"
                f"  extra   <img src> (count): {sorted((out_html - src_html).items()) or 'none'}\n"
                f"  missing ![](path) (count): {sorted((src_md - out_md).items()) or 'none'}\n"
                f"  extra   ![](path) (count): {sorted((out_md - src_md).items()) or 'none'}\n"
                f"  fix: restore the missing image refs in {out_chunk.name} from {src_chunk.name}"
            )

    if errors:
        print("\n=== Image validation failed ===")
        for e in errors:
            print(e)
            print()
        return False
    return True


def _chunk_id_from_output(path):
    """output_chunk0001.md -> chunk0001.md"""
    name = os.path.basename(path)
    return name[len('output_'):] if name.startswith('output_') else name


def _latex_shape(text):
    """Structural fingerprint of the LaTeX in one chunk.

    Deliberately excludes the raw backslash count: the translation prompt
    explicitly allows deleting line-ending backslashes, so that number moves
    legitimately. What must NOT move is environment balance and the row/cell
    structure inside a tabular -- translation changes cell TEXT, never the
    grid around it.
    """
    envs = Counter(re.findall(r'\\begin\{([^}]+)\}', text))
    ends = Counter(re.findall(r'\\end\{([^}]+)\}', text))
    blocks = []
    for t in find_raw_latex_tables(text):
        bare = t['bare']
        blocks.append({
            'rows': bare.count('\\\\'),
            'cells': bare.count('&'),
            'rules': (bare.count('\\toprule') + bare.count('\\midrule')
                      + bare.count('\\bottomrule') + bare.count('\\hline')),
        })
    return {'begin': envs, 'end': ends, 'blocks': blocks}


def _validate_chunk_latex(temp_dir):
    """Verify each translated chunk kept the LaTeX skeleton of its source.

    A shell heredoc silently collapses `\\\\` to `\\`, which strips every row
    separator out of a tabular. The table then renders as one run-on row, or
    fails to convert at all -- and nothing else in the pipeline notices,
    because the text is all still there.
    """
    temp_path = Path(temp_dir)
    errors = []
    for out in sorted(temp_path.glob('output_chunk*.md')):
        stem = out.name[len('output_'):]
        srcf = temp_path / stem
        if not srcf.exists():
            continue
        try:
            a = _latex_shape(srcf.read_text(encoding='utf-8'))
            b = _latex_shape(out.read_text(encoding='utf-8'))
        except OSError as e:
            errors.append(f"ERROR: {out.name}: {e}")
            continue

        problems = []
        for env in sorted(set(a['begin']) | set(b['begin'])):
            if a['begin'][env] != b['begin'][env]:
                problems.append(f"\\begin{{{env}}}: source {a['begin'][env]} "
                                f"-> output {b['begin'][env]}")
        for env in sorted(set(a['end']) | set(b['end'])):
            if a['end'][env] != b['end'][env]:
                problems.append(f"\\end{{{env}}}: source {a['end'][env]} "
                                f"-> output {b['end'][env]}")
        if len(a['blocks']) != len(b['blocks']):
            problems.append(f"tabular blocks: source {len(a['blocks'])} "
                            f"-> output {len(b['blocks'])}")
        else:
            for i, (x, y) in enumerate(zip(a['blocks'], b['blocks']), 1):
                for key, human in (('rows', 'row separators (\\\\)'),
                                   ('cells', 'cell separators (&)'),
                                   ('rules', 'booktabs rules')):
                    if x[key] != y[key]:
                        problems.append(f"table {i} {human}: source {x[key]} "
                                        f"-> output {y[key]}")
        if problems:
            errors.append(
                f"ERROR: {out.name} lost LaTeX structure\n"
                + ''.join(f"  {p}\n" for p in problems[:8])
                + f"  cause: the translated file was almost certainly written "
                  f"through a shell\n"
                  f"         heredoc/printf, which collapses '\\\\' to '\\'.\n"
                  f"  fix: rewrite {out.name} with the LaTeX skeleton copied "
                  f"verbatim from\n"
                  f"       {stem} (use a file-writing tool or Python, never the "
                  f"shell), or\n"
                  f"       delete it and re-translate that chunk."
            )

    if errors:
        print("\n=== LaTeX structure validation failed ===")
        for e in errors:
            print(e)
        return False
    return True


def _validate_chunk_math(temp_dir):
    """Verify every math placeholder survived translation exactly once.

    A dropped token means a formula vanished from the book — invisible in the
    build log, obvious to a reader. So this is a hard failure naming the chunk,
    not a warning.

    Chunks with no sidecar are skipped, so temp dirs created before the math
    guard existed still merge unchanged.
    """
    temp_path = Path(temp_dir)
    sidecars = sorted(temp_path.glob('chunk*' + math_guard.SIDECAR_SUFFIX))
    if not sidecars:
        return True

    errors = []
    checked = 0
    for sidecar in sidecars:
        stem = sidecar.name[:-len(math_guard.SIDECAR_SUFFIX)]
        src = temp_path / f'{stem}.md'
        out = temp_path / f'output_{stem}.md'
        if not src.exists() or not out.exists():
            continue
        try:
            spans = math_guard.load_sidecar(temp_dir, f'{stem}.md')
        except (ValueError, OSError, json.JSONDecodeError) as e:
            errors.append(f"ERROR: {sidecar.name}: {e}")
            continue

        checked += 1
        report = math_guard.verify(src.read_text(encoding='utf-8'),
                                  out.read_text(encoding='utf-8'), spans)
        if report['missing'] or report['duplicated'] or report['foreign']:
            errors.append(
                f"ERROR: output_{stem}.md corrupted math placeholders\n"
                f"  dropped by translator : {report['missing'][:10] or 'none'}\n"
                f"  duplicated            : {report['duplicated'][:10] or 'none'}\n"
                f"  not from this chunk   : {report['foreign'][:10] or 'none'}\n"
                f"  fix: delete output_{stem}.md and re-translate that chunk. Every\n"
                f"       ⟦M####⟧ / ⟦C####⟧ / ⟦T####⟧ token must be copied\n"
                f"       through verbatim, exactly once."
            )

    if errors:
        print("\n=== Math placeholder validation failed ===")
        for e in errors:
            print(e)
            print()
        return False

    if checked:
        print(f"Math placeholder check: {checked} chunk(s) OK")
    return True


def _restore_math_for(temp_dir, output_path, content):
    """Substitute math tokens back to LaTeX at merge-read time.

    Deliberately non-destructive: rewriting output_chunk*.md in place would flip
    run_state's output-hash check for every chunk and make resume logic think
    every translation had changed.
    """
    chunk_name = _chunk_id_from_output(output_path)
    try:
        spans = math_guard.load_sidecar(temp_dir, chunk_name)
    except (ValueError, OSError, json.JSONDecodeError) as e:
        print(f"Warning: {e}")
        return content
    if spans is None:
        return content  # pre-upgrade temp dir: no-op
    return math_guard.restore(content, spans)


def _check_generated_html_sanity(html_path):
    """Sanity-check generated HTML for malformed <img> tags. Returns False on problems.

    Note: we deliberately do NOT flag `&lt;img` in the rendered HTML — books that
    discuss HTML in prose or code blocks legitimately render escaped `<img>` text,
    and that's not a corruption signal. Real corruption produces a malformed
    actual `<img>` tag, which the attribute-name check catches."""
    try:
        text = Path(html_path).read_text(encoding='utf-8')
    except Exception as e:
        print(f"ERROR: cannot read {html_path}: {e}")
        return False

    _, bad_attrs = _scan_img_tags(text)
    if not bad_attrs:
        return True

    print(f"ERROR: image sanity check failed on {Path(html_path).name}")
    for raw_tag, attr_name in bad_attrs:
        print(f"  - malformed <img>: {raw_tag}")
        print(f"    bad attribute name: '{attr_name}'")
    print(
        "  fix: inspect output.md and the corresponding output_chunk*.md;\n"
        "       if alt text contains literal quotes, replace with curly quotes or HTML entity"
    )
    return False

# Try to import BeautifulSoup for TOC generation
try:
    from bs4 import BeautifulSoup
    BS4_AVAILABLE = True
except ImportError:
    BS4_AVAILABLE = False

# Try to import markdown
try:
    import markdown
    MARKDOWN_AVAILABLE = True
except ImportError:
    MARKDOWN_AVAILABLE = False

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

# =============================================================================
# Language configuration — single source of truth for lang-dependent values
# =============================================================================

# The tables themselves live in layout.py so calibre_html_publish.py can
# share the one copy; they used to be duplicated there and had already
# drifted. Re-exported under the old names because callers and tests reach
# for merge_and_build.LANG_CONFIG / .get_lang_config.
LANG_CONFIG = layout.LANG_CONFIG
_DEFAULT_LANG_CONFIG = layout.DEFAULT_LANG_CONFIG
get_lang_config = layout.get_lang_config


def load_config(temp_dir):
    """Load configuration from config.txt"""
    config_file = os.path.join(temp_dir, 'config.txt')
    if not os.path.exists(config_file):
        print("Error: config.txt not found in temp directory.")
        sys.exit(1)

    config = {}
    with open(config_file, 'r', encoding='utf-8') as f:
        for line in f:
            if '=' in line and not line.strip().startswith('#'):
                key, value = line.strip().split('=', 1)
                config[key] = value
    return config


def natural_sort_key(text):
    """Natural sorting key for filenames with numbers"""
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r'(\d+)', text)]


# =============================================================================
# Step 4: Merge translated markdown files
# =============================================================================

_LEADING_BLANK_LINES_RE = re.compile(r'\A(?:[ \t]*\r?\n)+')


def trim_chunk_edges(content):
    r"""Drop blank lines around a chunk without touching its indentation.

    `.strip()` stood here, and it also removed the leading spaces of the
    first line. A chunk boundary can fall inside a code listing:
    VLA-Adapter's chunk0012 begins `            # RoPE`, and once those
    twelve spaces are gone markdown reads the line as a top-level heading.
    It reached the book as an H1 with its own table-of-contents entry,
    sitting between two halves of the same Python class.

    The translator was not involved -- `output_chunk0012.md` still had the
    indentation. Only the seam lost it, which is why nothing that reads a
    single chunk could ever have seen this.
    """
    return _LEADING_BLANK_LINES_RE.sub('', content.rstrip())


def merge_markdown_files(temp_dir):
    """Merge all translated output files into output.md"""
    print("=== Merging translated markdown files ===")

    output_md = os.path.join(temp_dir, 'output.md')

    # Always validate manifest, even if output.md exists (catch stale/corrupt outputs)
    ok, ordered_files, warnings = validate_for_merge(temp_dir)

    # Image structure validation runs unconditionally — bad chunks invalidate any cached output.md
    if not _validate_chunk_images(temp_dir):
        if os.path.exists(output_md):
            print("Removing stale output.md (built from chunks that failed image validation)")
            os.remove(output_md)
        return False

    # Math placeholders are compared BEFORE restoration, while both source and
    # translated chunks still hold tokens.
    if not _validate_chunk_math(temp_dir):
        if os.path.exists(output_md):
            print("Removing stale output.md (chunks failed math placeholder validation)")
            os.remove(output_md)
        return False

    # LaTeX skeleton check. Compared BEFORE merging, while each translated
    # chunk can still be matched against the source it came from.
    if not _validate_chunk_latex(temp_dir):
        if os.path.exists(output_md):
            print("Removing stale output.md (chunks failed LaTeX structure validation)")
            os.remove(output_md)
        return False

    # A temp dir cleaned with --cleanup has no chunk*.md left, so
    # validate_for_merge() can never succeed there. In that state output.md is
    # the ONLY surviving copy of the translation — deleting it would destroy
    # hours of work that cannot be rebuilt without re-translating.
    src_chunks_present = any(
        not p.name.startswith('output_') for p in Path(temp_dir).glob('chunk*.md')
    )

    if os.path.exists(output_md):
        if not ok and not src_chunks_present:
            print("Chunk sources absent (post-cleanup temp dir) — reusing existing output.md as-is")
            print("  NOTE: chunk-level validation skipped. Re-run convert.py to re-split if needed.")
            return True
        if not ok:
            print(f"WARNING: output.md exists but manifest validation failed — deleting stale output.md")
            os.remove(output_md)
        else:
            # Check if any output_chunk is newer than output.md (re-translated chunks)
            output_md_mtime = os.path.getmtime(output_md)
            newer_chunks = []
            if ordered_files:
                newer_chunks = [
                    os.path.basename(f) for f in ordered_files
                    if os.path.getmtime(f) > output_md_mtime
                ]
            if newer_chunks:
                print(f"Re-merging — {len(newer_chunks)} chunk(s) newer than output.md: {', '.join(newer_chunks[:5])}{'...' if len(newer_chunks) > 5 else ''}")
                os.remove(output_md)
            else:
                print(f"Skipping merge - output.md already exists and is up to date")
                return True

    if not ok:
        print("ERROR: Merge validation failed. Fix the issues above before merging.")
        return False

    if ordered_files is not None:
        # Manifest-based merge
        print(f"Merging {len(ordered_files)} translated files (manifest-ordered)")
        merged_content = ""
        for file_path in ordered_files:
            content = read_output_text(file_path)
            if content is None:
                print(f"ERROR: Cannot read {os.path.basename(file_path)} — aborting merge")
                return False
            content = trim_chunk_edges(content)
            if not content:
                # validate_for_merge already rejects blank outputs; this is a
                # last line of defense so a chunk can never vanish silently.
                print(f"ERROR: Blank output {os.path.basename(file_path)} — aborting merge")
                return False
            content = _restore_math_for(temp_dir, file_path, content)
            merged_content += content + "\n\n"
    else:
        # Legacy fallback: glob-based merge (no manifest)
        print("WARNING: No manifest.json found — using legacy glob-based merge.")
        print("  For hash validation, re-run convert.py to generate manifest.json")

        # Match chunk output files
        output_files = glob.glob(os.path.join(temp_dir, 'output_chunk*.md'))

        # Count original source files
        original_files = glob.glob(os.path.join(temp_dir, 'chunk*.md'))
        original_files = [f for f in original_files if not os.path.basename(f).startswith('output_')]

        if not output_files:
            print("Error: No translated markdown files found.")
            return False

        # Build expected output filename for each source file and verify 1:1 match
        source_basenames = sorted(
            [os.path.basename(f) for f in original_files],
            key=natural_sort_key
        )
        expected_outputs = set(f"output_{name}" for name in source_basenames)
        actual_outputs = set(os.path.basename(f) for f in output_files)

        missing = expected_outputs - actual_outputs
        orphaned = actual_outputs - expected_outputs

        if missing or orphaned:
            if missing:
                print(f"ERROR: Missing translations for: {', '.join(sorted(missing, key=natural_sort_key))}")
            if orphaned:
                print(f"ERROR: Orphaned outputs (no matching source): {', '.join(sorted(orphaned, key=natural_sort_key))}")
            return False

        # Verify no empty, unreadable, or whitespace-only output files
        for fp in output_files:
            if os.path.getsize(fp) == 0:
                print(f"ERROR: Empty output file: {os.path.basename(fp)}")
                return False
            text = read_output_text(fp)
            if text is None:
                print(f"ERROR: Unreadable output file: {os.path.basename(fp)}")
                return False
            if not text.strip():
                print(f"ERROR: Blank output file: {os.path.basename(fp)}")
                return False

        # Use source order to determine merge order (via expected output names)
        output_files = [
            os.path.join(temp_dir, f"output_{name}")
            for name in source_basenames
        ]
        print(f"Merging {len(output_files)} translated files (legacy glob)")

        merged_content = ""
        for file_path in output_files:
            content = read_output_text(file_path)
            if content is None:
                print(f"ERROR: Cannot read {os.path.basename(file_path)} — aborting merge")
                return False
            content = trim_chunk_edges(content)
            if not content:
                print(f"ERROR: Blank output {os.path.basename(file_path)} — aborting merge")
                return False
            content = _restore_math_for(temp_dir, file_path, content)
            merged_content += content + "\n\n"

    # Belt-and-braces: whichever merge branch ran, no placeholder may survive
    # into output.md. A leaked token would render literally in the final PDF.
    leftover = sorted({m.group(0) for m in math_guard.TOKEN_RE.finditer(merged_content)})
    if leftover:
        print(f"ERROR: {len(leftover)} math token(s) survived the merge: {leftover[:10]}")
        print("  cause: a sidecar is missing or stale for a chunk that contains tokens.")
        print("  fix: re-run convert.py to regenerate chunks and sidecars.")
        return False

    try:
        with open(output_md, 'w', encoding='utf-8', newline='\n') as f:
            f.write(merged_content)
        file_size = os.path.getsize(output_md)
        print(f"Merged into output.md ({file_size:,} bytes)")
        return True
    except Exception as e:
        print(f"Error saving merged file: {e}")
        return False


# =============================================================================
# Step 5: Convert markdown to HTML
# =============================================================================

# Reader extensions used for every markdown->X pandoc call.
#   tex_math_dollars        : $...$ / $$...$$ become real math (not literal text)
#   tex_math_single_backslash: also accept \(...\) and \[...\]
#   pipe_tables/grid_tables : tables must survive to every output format
#   -markdown_in_html_blocks: the only block-level HTML this pipeline emits is
#     the raw-LaTeX tables, rendered here as finished HTML with MathML inside.
#     Left on, pandoc reads markdown in them, and a literal `*` or `_` in a
#     formula pairs with the next one and splices an `<em>` through the middle
#     of the MathML: CafeQ's table 3 printed `45.6^{}` and lost the asterisk
#     its own caption explains. Nothing in these blocks is markdown, so the
#     parsing has nothing to do but damage.
PANDOC_FROM = ('markdown+smart+east_asian_line_breaks+tex_math_dollars'
               '+tex_math_single_backslash+pipe_tables+grid_tables+raw_html'
               '-markdown_in_html_blocks')

# Scripts that write no space between words, so a wrapped line has none to
# lose. Korean is East Asian to pandoc and to Unicode, and is NOT one of them.
_NO_INTERWORD_SPACE = ('zh', 'ja')


def pandoc_from(lang=None):
    r"""The reader extensions, minus any this language must not have.

    `east_asian_line_breaks` deletes the newline between two East Asian
    characters. That is right for Chinese and Japanese, where a line wrapped
    in the source carries no space to begin with. Korean separates words with
    spaces and pandoc classifies Hangul as East Asian too, so a paragraph
    wrapped across lines in the merged markdown came back with its words run
    together: `가로지르고 있든\n손잡이를` printed as `있든손잡이를`.

    No shipped book has shown it, for one reason: translator sub-agents happen
    to write each paragraph as a single long line. That is luck rather than
    design -- the merged markdown is an ordinary text file, a hand edit or a
    differently-behaved agent can wrap it, and the damage is silent because
    every count still agrees.
    """
    base = (lang or '').split('-')[0].lower()
    if base in _NO_INTERWORD_SPACE:
        return PANDOC_FROM
    return PANDOC_FROM.replace('+east_asian_line_breaks', '')

# [] = not yet resolved, [None] = definitively absent
_PANDOC_PATH = []


def resolve_pandoc():
    """Return an absolute pandoc path, or None. Cached after the first call.

    A bare `pandoc` lookup is not enough: on Windows the official installer
    drops pandoc in %LOCALAPPDATA%\\Pandoc without touching PATH, so probing
    only PATH silently downgrades the whole build to a table-less fallback.
    """
    if _PANDOC_PATH:
        return _PANDOC_PATH[0]

    found = None
    # pypandoc knows how to locate pandoc without PATH
    try:
        import pypandoc
        found = pypandoc.get_pandoc_path()
    except Exception:
        found = None

    if not found:
        found = shutil.which('pandoc')

    if not found:
        exe = 'pandoc.exe' if os.name == 'nt' else 'pandoc'
        for cand in [
            os.environ.get('PANDOC_PATH'),
            os.environ.get('PYPANDOC_PANDOC'),
            os.path.join(os.environ.get('LOCALAPPDATA', ''), 'Pandoc', exe),
            r"C:\Program Files\Pandoc\pandoc.exe",
            r"C:\Program Files (x86)\Pandoc\pandoc.exe",
            "/usr/local/bin/pandoc", "/usr/bin/pandoc", "/opt/homebrew/bin/pandoc",
        ]:
            if cand and os.path.isfile(cand):
                found = cand
                break

    if found:
        try:
            result = subprocess.run(
                [found, '--version'], capture_output=True, text=True,
                encoding='utf-8', errors='replace', timeout=30
            )
            if result.returncode != 0:
                found = None
            else:
                first = result.stdout.splitlines()[0] if result.stdout else '?'
                print(f"Found pandoc: {found} ({first})")
        except (OSError, subprocess.TimeoutExpired):
            found = None

    _PANDOC_PATH.append(found)
    return found


# Commands that reach the markdown intact and then print as literal text.
# \IEEEPARstart{T}{raining} is IEEE's drop-cap macro: the two arguments are the
# first letter and the rest of the word, so dropping it loses a real word.
# Brace-balanced, not line-anchored: \markboth{...} routinely wraps across
# two lines, and dropping only its first line left a stranded "FEB 2025}" in
# the middle of the page.
# \captionsetup{font={footnotesize}} chooses how a caption is TYPESET and puts
# nothing on the page, but pandoc keeps it as a raw inline. DeeR-VLA writes one
# in front of every caption, and that is enough: the figure formatter looks for
# a caption in the paragraph after the image, finds this instead, and gives up.
# Six of its seven figures reached the page with the caption loose in the prose
# — no number, no figcaption, and every "그림 N" in the text pointing at nothing.
# `\index{sub-gaussian}` marks a term for an index this pipeline does not
# build. randmat writes 79 of them and every one printed as literal text in the
# body; `\printindex` printed itself too. Dropping them is a real reduction —
# the paper has an index and the book will not — so `report_index_terms` says
# so at build time rather than letting it pass unremarked (K110).
_LEFTOVER_CMDS = ('maketitle', 'providecommand', 'newcommand', 'renewcommand',
                  'markboth', 'captionsetup', 'IEEEpeerreviewmaketitle',
                  'IEEEdisplaynontitleabstractindextext',
                  'IEEEtitleabstractindextext')
_INDEX_HEAD_RE = re.compile(r'\\(?:index|printindex|makeindex)\b')
# Directives with nothing in them for a reader, which pandoc hands through as
# raw inlines exactly the way `\index` is — so they need the same removal, span
# and all (K133). Everything here either takes no argument or takes one that is
# a declaration: a page style, a counter name, a length. Deliberately absent
# are the commands whose argument is CONTENT and which therefore need
# resolving rather than dropping — `\parhead` (a run-in heading), `\subref`
# and `\cref` (references), `\newcite` (a citation), `\answerYes` (a checklist
# answer), `\etal`, `\caption`, `\emph`. See K135 for that inventory.
_DIRECTIVE_HEAD_RE = re.compile(
    r'\\(?:pagestyle|thispagestyle|@addtoreset|allowdisplaybreaks'
    r'|algorithmicindent|algrenewcommand|titlerunning|authorrunning'
    r'|large|Large|normalsize)\b')


def drop_directive_spans(md_text):
    r"""Remove content-free directives and the raw inline around them."""
    return _drop_head_and_span(md_text, _DIRECTIVE_HEAD_RE)


def drop_index_terms(md_text):
    r"""Remove `\index{...}` AND the code span it sits in. (text, count).

    Not a member of `_LEFTOVER_CMDS`, and that is the whole point. pandoc hands
    these through as raw inlines, so the markdown says
    `` `\index{Condition number}` ``; dropping only the command leaves an empty
    pair of backticks, and an empty pair is a code-span DELIMITER that swallows
    everything to the next one — including the `$` on either side. Doing it the
    ordinary way took randmat from 0 unrendered formulas to 29 while fixing
    the 79 markers. `arxiv_backend.strip_latex_cruft` learned the same thing
    (`_EMPTY_CODE_RE`); the span has to go with the command.
    """
    return _drop_head_and_span(md_text, _INDEX_HEAD_RE)


def _drop_head_and_span(md_text, head_re):
    r"""Drop each match, its braced arguments, and the raw inline around it."""
    count = 0
    out = []
    at = 0
    while True:
        m = head_re.search(md_text, at)
        if not m:
            out.append(md_text[at:])
            return ''.join(out), count
        stop = m.end()
        # `\algrenewcommand{\algorithmicindent}{...}` takes two; take every
        # brace group that follows, not only the first.
        while stop < len(md_text) and md_text[stop:stop + 1] == '{':
            close = _balanced_group(md_text, stop)
            if close < 0:
                break
            stop = close
        if stop < len(md_text) and md_text[stop:stop + 1] == '{':
            out.append(md_text[at:m.end()])       # unbalanced: leave it alone
            at = m.end()
            continue
        start = m.start()
        # Take the wrapping raw-inline with it when nothing else is inside.
        tail = re.match(r'`(?:\{=[a-z]+\})?', md_text[stop:])
        if start and md_text[start - 1] == '`' and tail:
            start -= 1
            stop += tail.end()
        out.append(md_text[at:start])
        at = stop
        count += 1


_PARSTART_RE = re.compile(r'`?\\IEEEPARstart\{(.)\}\{([^}]*)\}`?')
# `ack` is NeurIPS's acknowledgements environment. pandoc has no reader for it,
# so the whole block — translated prose, funding, thanks — reached the markdown
# as raw LaTeX and was dropped without a word on the HTML path. The wrapper
# lines carry nothing; only what is between them is content.
_ENV_WRAPPER_RE = re.compile(
    r'^[ \t]*\\(?:begin|end)\{(?:IEEEkeywords|IEEEtitleabstractindextext|'
    r'abstract|ack|acks|acknowledgements|acknowledgments|'
    r'appendices|subappendices)\}[ \t]*$\n?',
    re.MULTILINE)
# An empty inline code span, with or without pandoc's raw-inline marker. It is
# what a command stripped out of a code span leaves behind, it renders as two
# stray backticks, and all three books printed one right after a heading:
# "(Estimating Layer Importance in MoE) `` 캘리브레이션 시점의 활성값이…".
# Page 8 of SINQ had the marker form too -- "정의한다. ``{=latex}여기서 σ…".
#
# Two things it must NOT touch, both found by trying them:
#   ``code with a ` inside``   -- a double-backtick span opens with exactly
#                                 this shape, so only a run standing alone
#                                 between spaces counts as empty
#   `a``b`                     -- adjacent spans, likewise not alone
# And a raw inline holding `<!-- -->` is load-bearing: pandoc puts it between
# `$\times$` and `7B` so a closing `$` followed by a digit still reads as
# maths, so only EMPTY ones go.
_EMPTY_RAW_INLINE_RE = re.compile(
    r'(?<![`\S])`[ \t]*`(?![`\S])'          # alone between spaces
    r'|(?<![`\S])`[ \t]*`\{=[a-z]+\}')      # or carrying pandoc's marker
# The marker with no span left in front of it. SINQ page 8 printed
# `\end{equation}$$ {=latex}여기서 …` -- the backticks had already gone, so
# neither form above caught it and the attribute reached the reader.
#
# This one CANNOT live in _EMPTY_RAW_INLINE_RE: that regex is applied to the
# slices between code regions, and a slice can begin at `{=html}` with the
# backtick that owns it in the previous slice. The lookbehind then sees the
# start of a string and strips the marker off a live raw inline -- which is
# how `<!-- -->`, the comment that keeps `$\times$` apart from `7B`, became
# an ordinary code span and printed itself 21 times in AlphaQ.
# An OPENING bracket may also sit in front of it: BERT writes
# "네 번째 토큰({=latex}`hairy`에 해당)". Requiring whitespace there left that one
# marker on the page. A code span can never END with `(`, `[` or `{`, so
# admitting those cannot strip the attribute off a live raw inline -- which is
# the failure the lookbehind exists to prevent.
_ORPHAN_RAW_ATTR_RE = re.compile(
    r'(?:(?<![`\S])|(?<=[(\[{]))\{=[a-z]+\}')
# `\setlength\abovedisplayskip{3pt}` written INSIDE the display maths, as
# DeeR-VLA does on all eight of its equations. texmath has no reader for it
# and refuses the whole formula, so the `$$` print as text and the equation
# is gone. The WHOLE LINE has to go: taking only the command leaves a blank
# line, a blank line ends the display block, and the dollars print anyway --
# measured, 8 unrendered spans became 12.
_SETLENGTH_RE = re.compile(
    r'[ \t]*\\setlength\s*(?:\{\s*\\[a-zA-Z]+\s*\}|\\[a-zA-Z]+)'
    r'\s*\{[^{}]*\}[ \t]*\n?')
_SIDESET_RE = re.compile(r'\\sideset\s*\{\s*\}\s*\{')


def rewrite_sideset(text):
    r"""`\sideset{}{_{X}}\sum` -> `\sum\nolimits_{X}`. (text, count).

    Not a deletion and not an approximation: `\sideset` with an empty left
    argument asks for the scripts BESIDE the operator, and `\nolimits` says
    exactly that. DeeR-VLA's own equation already uses `\sum\nolimits`
    further along the same line.

    texmath has no reader for `\sideset`, so the whole formula was refused
    and the `$$` printed as text. The lesson is wider than the command: the
    boundary is not where a command is unsupported, it is where the
    supported subset has no equivalent -- and that is a much smaller place.
    A `\sideset` carrying BOTH sides really has none, and is left alone.
    """
    out, cursor, count = [], 0, 0
    for m in _SIDESET_RE.finditer(text):
        if m.start() < cursor:
            continue
        close = _balanced_group(text, m.end() - 1)
        if close < 0:
            continue
        scripts = text[m.end():close - 1].strip()
        tail = re.match(r'\s*(\\[a-zA-Z]+)', text[close:])
        if not tail:
            continue
        if scripts.startswith(('_', '^')):
            replacement = '%s\\nolimits%s' % (tail.group(1), scripts)
        elif _SIDESET_MARK_RE.fullmatch(scripts):
            # The primed sum: `\sideset{}{'}\sum` is a RESTRICTED sum, and the
            # mark belongs beside the operator. `{\sum}'` puts it there and
            # texmath reads it. Maynard writes it 7 times and every one of
            # those formulas printed as source -- the rewrite was here, but it
            # only knew the second argument as a SCRIPT, so a prime walked
            # straight past it. Dropping the mark is not available: it is what
            # makes the sum restricted.
            replacement = '{%s}%s' % (tail.group(1), scripts)
        else:
            continue
        out.append(text[cursor:m.start()])
        out.append(replacement)
        cursor = close + tail.end()
        count += 1
    out.append(text[cursor:])
    return ''.join(out), count


# A raw inline carrying only a length: what `\vspace*{-2em}` becomes.
# Two shapes. Backticked is what pandoc usually leaves; BARE is what a
# half-stripped `\vspace*{-2.5mm}` leaves, and it arrives on a line of its own
# directly above an image. The bare form is anchored to the whole line, so a
# brace inside ordinary prose is never touched.
_LENGTH = r'\{\s*-?[\d.]+\s*(?:em|ex|pt|cm|mm|in|bp)\s*\}'
_SPACING_INLINE_RE = re.compile(
    r'[ \t]*`' + _LENGTH + r'`(?:\{=[a-z]+\})?'
    r'|^[ \t]*' + _LENGTH + r'[ \t]*$', re.M)
# A literal [word] that pandoc escaped as \[word\]; the markdown reader then
# reads \[...\] as DISPLAY MATH under tex_math_single_backslash and renders it
# in math italic. Only unescaped when the content is plainly not math.
_ESCAPED_BRACKET_RE = re.compile(r'\\\[([^\\\]$]{1,40}?)\\\]')


# `\subcaption{Cayley SGD.}` labels one panel of a multi-panel float.
# pandoc leaves it as an inline code span, sometimes with a `{=latex}`
# tail, and it prints verbatim next to the image. Its argument is the
# panel's caption text, which belongs on the page.
# The body is read with a balanced scan, not `[^{}]*`: CafeQ's panel label is
# `\subcaption[t]{Adam; $\lambda_{orth}=0$.}`, whose nested braces made the
# old pattern miss it entirely and print the command to the reader verbatim.
_SUBCAPTION_OPEN_RE = re.compile(
    r'`?\\subcaption\*?\s*(?:\[[^\]]*\])?\s*(?=\{)')


def _drop_latex_commands(text, names, stats):
    """Remove `\\cmd{...}{...}` calls whole, across newlines."""
    for name in names:
        while True:
            m = re.search(r'\\' + name + r'\b', text)
            if not m:
                break
            end = m.end()
            # Consume every brace group that follows, balanced.
            while end < len(text):
                nxt = re.match(r'[ \t\n]*\{', text[end:])
                if not nxt:
                    break
                close = _balanced_group(text, end + nxt.end() - 1)
                if close < 0:
                    break
                end = close
            text = text[:m.start()] + text[end:]
            stats['dropped'] += 1
    return text


# pandoc's TeX reader implements none of the pre-LaTeX2e font switches, and it
# does not ignore them either: `{\rm max}` makes the *entire* formula fail to
# parse, so it reaches the page as literal `$s = (w_{\rm max} - ...)$`. One
# such token inside a display equation costs the reader the whole equation.
#
# Applied everywhere except code. Scoping this to math spans looks safer and
# is not: pairing `$` across a paragraph desynchronises on the first stray one
# (a literal price, an escaped \$, an unbalanced formula), and every span after
# it is scanned off by one. Three `{\rm Q}` in CafeQ survived that way and took
# two displayed equations to the page as raw source. `{\rm ...}` is broken TeX
# wherever it lands in markdown; code is the one place it is a quoted example.
# The trailing `(?<!\\)` matters: `\ ` is an ESCAPED SPACE, an atomic token,
# not cosmetic padding. Shor writes `{\rm \ (mod\ }n`, and trimming that last
# space welded the backslash to what followed — `\n`, a control sequence no
# renderer knows — so 77 formulas in one paper printed as source.
_OLD_FONT_RE = re.compile(
    r'\{\s*\\(rm|bf|it|sf|tt|sc|cal)\s+([^{}]*?)(?<!\\)\s*\}')
_OLD_FONT_MAP = {'rm': 'mathrm', 'bf': 'mathbf', 'it': 'mathit',
                 'sf': 'mathsf', 'tt': 'mathtt', 'sc': 'mathrm',
                 'cal': 'mathcal'}
# The other spelling: the switch CALLED with a braced argument, `\cal{A}`.
# That is not valid LaTeX2e either, and the cost is the same — texmath refuses
# the whole formula, so GAN's definition of a subderivative printed as raw
# source twice on the page. The group form above does not match it: there is
# no `{` before the command and no space after it.
_OLD_FONT_CALL_RE = re.compile(r'\\(rm|bf|it|sf|tt|sc|cal)\s*(?=\{)')
# And the third spelling: the switch with NO group at all, `$\rm P$`. The group
# form needs a `{` before it and the call form needs one after it, so a bare
# switch slips past both and texmath refuses the formula. `\mathrm P` is valid
# without braces, so the same substitution works here.
_OLD_FONT_BARE_RE = re.compile(
    r'\\(rm|bf|it|sf|tt|sc|cal)(?=\s+[^\s{])')

# `\textsc` is the OTHER half of the problem, and it must be treated
# differently: in text mode it is correct LaTeX and pandoc renders small caps,
# but texmath has no reader for it and refuses the whole formula. BERT names
# its two model sizes `BERT$_{\textsc{BASE}}$` and `BERT$_{\textsc{LARGE}}$`,
# so thirty-five formulas printed to the reader as raw source. Asked directly,
# pandoc renders the same span the moment `\textsc` becomes `\mathrm`.
#
# Scoped to inline math spans on one line — the same conservative shape the
# guard uses — because rewriting it in TEXT mode would break the small caps
# that are working.
# Display blocks count too, and they cost more: Neural ODE writes
# `\textnormal` inside fourteen `align` environments, so fourteen displayed
# derivations printed to the reader as raw LaTeX — 637 leaked tokens, every
# one of them from this. `$$…$$` is matched non-greedily across lines because
# a display block legitimately spans them.
# An inline span may WRAP. `$[t_\textnormal{start}, t_\textnormal{end}]$` sits
# across a line break in Neural ODE, and an alternative that stops at `\n`
# walks past it — nineteen `\textnormal` survived the rewrite that way while
# the same command was being fixed everywhere else. A newline is allowed
# inside a span; a BLANK line is not, because that ends the paragraph.
_INLINE_MATH_SPAN_RE = re.compile(
    r'(?<!\\)\$\$.*?(?<!\\)\$\$'
    r'|(?<![\\$\w])\$(?!\s)'
    r'(?:[^$\n\\]|\\.|\n(?![ \t]*\r?\n)){1,800}?(?<!\\)\$(?!\d)'
    r'|(?<=\w)\$[_^]'
    r'(?:[^$\n\\]|\\.|\n(?![ \t]*\r?\n)){1,800}?(?<!\\)\$(?!\d)', re.DOTALL)
_TEXT_FONT_IN_MATH_RE = re.compile(
    r'\\text(normal|sc|bf|it|tt|rm|sf|up|md)\s*(?=\{)')
_TEXT_FONT_MAP = {'normal': 'mathrm', 'sc': 'mathrm', 'bf': 'mathbf',
                  'it': 'mathit', 'tt': 'mathtt', 'rm': 'mathrm',
                  'sf': 'mathsf', 'up': 'mathrm', 'md': 'mathrm'}


# `\mkern18mu` is math-mode glue: it sets space and has nothing in it to read.
# texmath has no reader for it and refuses the formula around it, which is the
# same trade `\setlength` made in K100 — a spacing directive costing a whole
# derivation. Removed inside math only; it is meaningless outside.
_MKERN_RE = re.compile(r'\\mkern\s*-?[\d.]+\s*mu\s*')
# `\vphantom{...}` reserves height and prints nothing; texmath has no reader.
# Matched by brace BALANCE, not by a fixed depth: Neural ODE's is
# `\vphantom{\frac{\partial p(\mathbf{z}(t), t)}{\partial \mathbf{z}(t)}}`,
# four levels deep, and a two-level pattern walks straight past it.
_VPHANTOM_HEAD_RE = re.compile(r'\\(?:vphantom|hphantom|phantom)\s*(?=\{)')
# `\ref` and `\eqref` inside a formula. texmath has no reader for either, so a
# derivation annotated `\mathrm{(by Eq \ref{eq:chain_rule})}` prints as source
# in full. The label index knows the number; use it, and drop the command when
# it does not.
_MATH_REF_HEAD_RE = re.compile(r'\\(?:eqref|[A-Za-z]*ref)\s*\{([^{}]*)\}')
# `\nicefrac{a}{b}` is the nicefrac package's slanted fraction. texmath refuses
# it, and `\frac` says the same thing in the subset it does read — eleven of
# Neural ODE's inline formulas printed as source over this one command.
_NICEFRAC_RE = re.compile(r'\\nicefrac(?=\s*\{)')
# The mark a `\sideset` can carry beside an operator instead of a script: a
# prime, a double prime, or a star. Anything longer is not this shape.
_SIDESET_MARK_RE = re.compile(r"[*'\u2032]{1,2}")
# `\idotsint` is amsmath's iterated integral. texmath has no reader for it and
# `\int\cdots\int` is the same thing written in the subset it does read.
_IDOTSINT_RE = re.compile(r'\\idotsint(?![A-Za-z])')
# `{X \atop Y}` is the plain-TeX stack. texmath refuses it; `\substack` stacks
# the same two operands and renders. Measured both ways.
_ATOP_RE = re.compile(r'\{((?:[^{}]|\{[^{}]*\})*?)\\atop'
                      r'((?:[^{}]|\{[^{}]*\})*?)\}')
# `\multicolumn{1}{c}{X}` inside a math array spans ONE column, so it is pure
# alignment and X is the whole of its content. texmath refuses the formula
# over it. A span of 2 or more is left alone: dropping that would leave the
# row short of cells, which corrupts the array instead of rescuing it.
_MATH_MULTICOL1_RE = re.compile(
    r'\\multicolumn\s*\{\s*1\s*\}\s*\{[^{}]*\}\s*'
    r'\{((?:[^{}]|\{[^{}]*\})*)\}')
# `\hat\mathbf{x}` -- an accent whose argument is another command rather than a
# single token. LaTeX accepts it; texmath does not, and refuses the formula.
# Measured: `$\hat\mathbf{x}_0$` is rejected, `$\hat{\mathbf{x}}_0$` renders.
# Braces are added, never removed, so the meaning cannot change.
# `\text{\mathtt{[CLS] ...}}` -- a math font command wrapped in a text-mode
# box. texmath refuses it; `\mathtt{...}` alone renders and looks the same,
# because the inner command already sets the font. Measured both ways.
_TEXT_WRAPPING_MATHFONT_RE = re.compile(
    r'\\text(?:rm|normal|it|bf|sf|tt|up|md|sc)?\s*\{\s*'
    r'(\\math(?:tt|bf|rm|it|sf|cal|bb|frak|scr)\s*'
    r'\{(?:[^{}]|\{[^{}]*\})*\})\s*\}')
_BARE_ACCENT_RE = re.compile(
    r'\\(hat|bar|tilde|vec|dot|ddot|check|breve|acute|grave|widehat'
    r'|widetilde|widebar|overline|underline|mathring)\s*'
    r'(\\[A-Za-z]+\s*\{(?:[^{}]|\{[^{}]*\})*\})')
# `\qedhere` puts the QED box on the last display line and `\notag` suppresses
# an equation number. Both are typesetting directives with nothing to read, and
# both cost the whole formula -- four of randmat's displays, including the one
# carrying the proof's final inequality. The proof terminator is already
# handled as its own mark in the markdown, so nothing visible is lost.
# `\linebreak[3]` and friends carry an OPTIONAL argument, so the bracket has to
# go with the command; left behind, `[3]` prints as text in the middle of a
# formula. They are line-breaking hints with no content, like the rest here.
_MATH_DIRECTIVE_RE = re.compile(
    r'\\(?:qedhere|notag|nonumber|allowdisplaybreaks|displaybreak'
    r'|linebreak|nolinebreak|newline|pagebreak|nopagebreak)'
    r'(?![A-Za-z])\s*(?:\[[^\]\n]*\]\s*)?')


# `\mathrm{event at time $t$}` -- inside a text-mode argument the author
# switches back to math for one symbol. It is ordinary LaTeX, and it defeats
# every flat `$`-pairing scanner downstream, this module's own included: the
# span closes at the inner `$` and the dollars pair off by one from there.
# Inside `\mathrm{}` the argument is already set as text, so dropping the inner
# delimiters says the same thing -- measured against pandoc, which refuses the
# nested form and renders the flattened one.
# The nested formula carries braces of its own -- `$\mathbf{z}(t)$` -- so none
# of the three parts may exclude them outright. A brace-free pattern caught
# `$t$` and walked past `$\mathbf{z}(t)$`, leaving the whole display raw.
# The nested brace group must exclude `$` too, or the pattern runs from a
# `\text{` in one formula, through that formula's closing `$` and the next
# one's opening `$`, to a `}` further down.
_INNER = r'(?:[^{}$]|\{[^{}$]*\})'
_NESTED_MATH_IN_TEXT_RE = re.compile(
    r'(\\(?:math|text)(?:rm|normal|it|bf|sf|tt|up|md|sc)?\s*\{)'
    r'(' + _INNER + r'{0,120}?)\$(' + _INNER + r'{1,120})\$'
    r'(' + _INNER + r'{0,120}?)(\})')


def _drop_balanced_command(text, head_re):
    """Remove `\\cmd{...}` and its argument by brace BALANCE. (text, count)."""
    count = 0
    while True:
        m = head_re.search(text)
        if not m:
            return text, count
        close = _balanced_group(text, text.index('{', m.end() - 1))
        if close < 0:
            return text, count
        text = text[:m.start()] + text[close:]
        count += 1


def resolve_math_references(md_text, temp_dir):
    r"""Turn `\ref{key}` inside a formula into its number. (text, count).

    A derivation annotated `\mathrm{(by Eq \ref{eq:chain_rule})}` costs the
    whole display: texmath has no reader for `\ref`, so the entire block prints
    as LaTeX source. The number is already known here.
    """
    numbers = build_label_numbers(temp_dir)
    hits = [0]

    def fix_span(m):
        def sub(ref):
            hits[0] += 1
            number = numbers.get(ref.group(1).strip())
            return number if number else ''

        return _MATH_REF_HEAD_RE.sub(sub, m.group(0))

    return _INLINE_MATH_SPAN_RE.sub(fix_span, md_text), hits[0]


def split_nested_math_text(md_text):
    r"""`\text{A$X$B}` -> `\text{A}X\text{B}`. Returns (text, count).

    Measured against pandoc, not reasoned about, because the shape alone does
    not decide it: ResNet's `\text{3$\times$3, 64}` renders as written, and
    Neural ODE's `\mathrm{event at time $t$}` does not. Deleting the inner
    delimiters fixes the second and BREAKS the first — `\times` in text mode
    means nothing — which is how this was found: a clean ResNet went to 112
    leaked tokens on the flattening version of this pass.

    Closing the text group and reopening it is the one transformation both
    accept, and it is exact: A and B stay text, X stays maths. It also removes
    the nesting, so every flat `$`-pairing scanner downstream — this module's
    own included — stops mis-closing on these.

    Boundary: one nested span per argument. `\text{a$x$b$y$c}` is left exactly
    as it was rather than widened for; a pattern that reached further is what
    caused the ResNet regression above.
    """
    total = 0
    for _ in range(4):                        # an argument may hold several
        md_text, n = _NESTED_MATH_IN_TEXT_RE.subn(r'\1\2}\3\1\4}', md_text)
        total += n
        if not n:
            break
    return md_text, total


# The paper's own shorthand. `\def \< {\langle}` is a control SYMBOL, so no
# letter boundary applies to it; `\def \E {\mathbb{E}}` is a control WORD and
# must not be found inside `\Ell`.
_MATH_MACRO_DEF_RE = re.compile(
    r'\\(?:newcommand|renewcommand|providecommand)\s*\*?\s*'
    r'(?:\{\s*(\\[A-Za-z]+|\\[^A-Za-z\s])\s*\}|(\\[A-Za-z]+|\\[^A-Za-z\s]))'
    r'\s*(?:\[\s*\d+\s*\][^{]*)?(?=\{)'
    r'|\\def\s*(\\[A-Za-z]+|\\[^A-Za-z\s])\s*(?=\{)'
    r'|\\DeclareMathOperator\s*\*?\s*\{\s*(\\[A-Za-z]+)\s*\}\s*(?=\{)')


_TEXMATH_READS = {}


def _texmath_reads(commands):
    r"""{command: bool} — can texmath render `$\cmd$`? Asked once, cached.

    One pandoc call for the whole batch, matched back by the TeX annotation it
    writes beside each formula. When pandoc cannot be reached every answer is
    False, which reproduces the conservative behaviour this replaces: a macro
    whose target may be unreadable is dropped rather than followed.
    """
    todo = [c for c in commands if c not in _TEXMATH_READS]
    if not todo:
        return {c: _TEXMATH_READS.get(c, False) for c in commands}
    pandoc = resolve_pandoc()
    if not pandoc:
        for c in todo:
            _TEXMATH_READS[c] = False
        return {c: False for c in commands}
    doc = '\n\n'.join('$%s$' % c for c in todo)
    try:
        proc = subprocess.run(
            [pandoc, '-f', 'markdown+tex_math_dollars', '-t', 'html',
             '--mathml'],
            input=doc, capture_output=True, text=True, encoding='utf-8',
            errors='replace', timeout=60)
        html = proc.stdout or ''
    except Exception:                                        # noqa: BLE001
        html = ''
    rendered = {' '.join(m.split()) for m in
                re.findall(r'<annotation\b[^>]*>(.*?)</annotation>', html,
                           re.DOTALL)}
    for c in todo:
        _TEXMATH_READS[c] = c in rendered
    return {c: _TEXMATH_READS.get(c, False) for c in commands}


def read_math_macros(temp_dir):
    r"""{name: body} for the paper's own zero-argument math shorthand.

    arxiv_backend already collects these into math_macros.tex -- and nothing
    ever read the file back, so the definitions were written to disk and
    dropped. Vershynin writes `\def \< {\langle}` and then uses `\<` in 56
    formulas; texmath has never heard of `\<`, so all 56 printed as source.

    Only zero-argument definitions are expanded. One taking `#1` needs a real
    macro expander, and guessing at one would corrupt formulas that render
    correctly today.
    """
    macros = {}
    for name in ('math_macros.tex', 'flat.tex'):
        path = os.path.join(temp_dir or '', name)
        if not os.path.isfile(path):
            continue
        try:
            with open(path, 'r', encoding='utf-8', errors='replace') as fh:
                text = fh.read()
        except OSError:
            continue
        if name == 'flat.tex':
            text = text.split(r'\begin{document}')[0]
        for m in _MATH_MACRO_DEF_RE.finditer(text):
            token = m.group(1) or m.group(2) or m.group(3) or m.group(4)
            open_at = text.find('{', m.end() - 1)
            if not token or open_at < 0:
                continue
            close = _balanced_group(text, open_at)
            if close < 0:
                continue
            body = text[open_at + 1:close - 1]
            if '#' in body or not body.strip():
                continue
            if m.group(4):                    # \DeclareMathOperator
                body = r'\operatorname{%s}' % body
            macros.setdefault(token, body)
        # An alias is only useful when its target is readable. `\let\gev\GeV`
        # with no usable `\GeV` turns a name texmath cannot read into a
        # different name it cannot read, so it is dropped rather than followed
        # into nothing (K121).
        #
        # But "not defined in this document" is the wrong test for readable,
        # and it was deleting the definitions it exists to protect:
        # `\def \< {\langle}` has exactly the shape of a dangling alias and
        # points at a command texmath knows perfectly well. Nine papers lost
        # macros to it — randmat lost `\<` and `\>` and printed 48 formulas as
        # source. Of the 22 distinct targets across the corpus, 17 render.
        # Shape cannot separate the two cases, so ask (K132).
        candidates = sorted({v.strip() for v in macros.values()
                             if re.fullmatch(r'\\[A-Za-z]+', v.strip())
                             and v.strip() not in macros})
        usable = _texmath_reads(candidates)
        for name in [k for k, v in macros.items()
                     if v.strip() in candidates and not usable.get(v.strip())]:
            del macros[name]
        if macros:
            break
    return macros


def expand_math_macros(md_text, temp_dir):
    r"""Replace the paper's shorthand with what it stands for, inside math.

    Safe at this point in the pipeline precisely because the translator never
    saw any of it: every formula travelled as a placeholder and was restored
    from its sidecar a moment ago. Returns (text, count).
    """
    macros = read_math_macros(temp_dir)
    if not macros:
        return md_text, 0
    subs = []
    for token, body in sorted(macros.items(), key=lambda kv: -len(kv[0])):
        tail = r'(?![A-Za-z])' if token[1:].isalpha() else ''
        subs.append((re.compile(re.escape(token) + tail), body))
    hits = [0]

    def expand(span):
        for _ in range(4):                    # shorthand can nest
            before = span
            for pattern, body in subs:
                # A callable replacement: a body full of backslashes must not
                # be read back as escape sequences.
                span = pattern.sub(lambda _m, b=body: b, span)
            if span == before:
                break
        return span

    def on_span(m):
        out = expand(m.group(0))
        if out != m.group(0):
            hits[0] += 1
        return out

    return _INLINE_MATH_SPAN_RE.sub(on_span, md_text), hits[0]


def unwrap_text_boxed_math_fonts(md_text):
    r"""`\text{\mathtt{X}}` -> `\mathtt{X}`. Returns (text, count).

    Runs AFTER the legacy font pass, not with the other span rewrites. BERT
    writes `\text{\tt {[CLS] ...}}`; `\tt` only becomes `\mathtt` in
    normalize_math_commands, so an unwrap placed before that never sees the
    shape it is looking for -- which is exactly how this first failed to fire.
    """
    count = [0]

    def fix_span(m):
        span = m.group(0)
        for _ in range(3):                    # boxes can nest
            span, hit = _TEXT_WRAPPING_MATHFONT_RE.subn(
                lambda inner: inner.group(1), span)
            if not hit:
                break
            count[0] += hit
        return span

    return _INLINE_MATH_SPAN_RE.sub(fix_span, md_text), count[0]


def rewrite_text_fonts_in_math(md_text):
    r"""Make a formula readable to texmath. Returns (text, count).

    Two shapes, both of which cost the WHOLE formula: a text-mode font switch
    (`\textsc`, `\textnormal`) and math glue (`\mkern`). Neither has a reader,
    and a formula with one in it reaches the reader as raw LaTeX.
    """
    count = [0]

    def fix_span(m):
        def swap(inner):
            count[0] += 1
            return '\\' + _TEXT_FONT_MAP[inner.group(1)]

        def drop(_inner):
            count[0] += 1
            return ' '

        def to_frac(_inner):
            count[0] += 1
            return '\\frac'
        span = _TEXT_FONT_IN_MATH_RE.sub(swap, m.group(0))
        span = _MKERN_RE.sub(drop, span)
        span, dropped = _drop_balanced_command(span, _VPHANTOM_HEAD_RE)
        count[0] += dropped
        span = _MATH_DIRECTIVE_RE.sub(drop, span)

        def to_iterated(_inner):
            count[0] += 1
            return r'\int\cdots\int'
        span = _IDOTSINT_RE.sub(to_iterated, span)

        def to_substack(inner):
            count[0] += 1
            return '\\substack{%s \\\\ %s}' % (inner.group(1).strip(),
                                               inner.group(2).strip())
        span = _ATOP_RE.sub(to_substack, span)

        def unspan(inner):
            count[0] += 1
            return inner.group(1)
        span = _MATH_MULTICOL1_RE.sub(unspan, span)

        def brace(inner):
            count[0] += 1
            return '\\%s{%s}' % (inner.group(1), inner.group(2))
        span = _BARE_ACCENT_RE.sub(brace, span)
        return _NICEFRAC_RE.sub(to_frac, span)

    return _INLINE_MATH_SPAN_RE.sub(fix_span, md_text), count[0]
_CODE_REGION_RE = re.compile(
    r'^[ \t]*(?P<fence>```+|~~~+).*?^[ \t]*(?P=fence)[ \t]*$'
    r'|`[^`\n]+`',
    re.MULTILINE | re.DOTALL)
# Raw LaTeX tables reach the reader through this module's own renderer, not
# through pandoc, and a cell in one is text mode.
_RAW_TABLE_RE = re.compile(
    r'\\begin\{(?P<tenv>table\*?|tabular\*?|tabularx|longtable)\}'
    r'.*?\\end\{(?P=tenv)\}', re.DOTALL)
# Math inside such a table, which is still math and still has to be modernised.
_MATH_REGION_RE = re.compile(
    r'(?<!\\)\$\$.+?(?<!\\)\$\$'
    r'|(?<!\\)\$(?:\\.|[^$\\\n])+?(?<!\\)\$'
    r'|\\\(.+?\\\)'
    r'|\\\[.+?\\\]'
    r'|\\begin\{(?P<menv>equation|align|alignat|flalign|gather|multline'
    r'|eqnarray|displaymath|math)\*?\}.+?\\end\{(?P=menv)\*?\}',
    re.DOTALL)


# An accent whose argument is another command, written without braces:
# `\widetilde\mathbf{A}`. LaTeX takes the following command as the argument
# and typesets it; texmath wants a brace there, gives up on the WHOLE span,
# and the formula reaches the page as literal TeX. VLA-Adapter printed six
# equations that way, and `leak_probe` then counted 75 fragments of them.
#
# Measured against pandoc 3.10.2: `\widetilde\mathbf{A}^0_t`,
# `\widehat\mathbf{B}`, `\bar\mathcal{C}`, `\vec\boldsymbol{d}` and
# `\tilde\mathrm{e}` all fail, and every one renders once the argument is
# braced. So the fix is the brace, and it belongs to the family rather than
# to the one command this paper happened to use.
_MATH_ACCENTS = ('widetilde', 'widehat', 'overline', 'overrightarrow',
                 'underline', 'bar', 'hat', 'tilde', 'vec', 'dot', 'ddot',
                 'check', 'breve', 'acute', 'grave', 'mathring')
_MATH_STYLES = ('mathbf', 'mathrm', 'mathcal', 'mathbb', 'mathsf', 'mathtt',
                'mathit', 'mathfrak', 'mathscr', 'boldsymbol', 'bm', 'symbf')
# One level of nesting inside the style's argument, so `\mathbf{A_{t}}` is
# still matched whole and never cut in half.
_ACCENT_ON_COMMAND_RE = re.compile(
    r'\\(' + '|'.join(_MATH_ACCENTS) + r')\s*'
    r'(\\(?:' + '|'.join(_MATH_STYLES) + r')'
    r'\{(?:[^{}]|\{[^{}]*\})*\})')


def normalize_math_commands(md_text):
    """Modernise the pre-LaTeX2e font switches. Returns (text, stats).

    Everywhere except the inside of a raw LaTeX table. `\\mathbf` and its
    siblings are math-mode commands and LaTeX rejects them in text mode, so
    rewriting a text-mode `{\\bf ...}` does not modernise it -- it corrupts
    it. A `tabular` cell is text mode: CafeQ's table 4 carried `{\\bf 46.6}`
    there, the rewrite turned it into `\\mathbf{46.6}`, and the table
    renderer, handed a cell it could not parse, dropped the row. Three
    numbers the paper reports left the book while the build printed
    `8 converted, 0 failed`.

    Only the tables are held back, not every text-mode span, because the
    rewrite is what makes the math legible to texmath: left alone, `{\\rm Q}`
    reaches the HTML as unrendered `$...$`. Scoping the rewrite to math
    document-wide is not available either -- that needs `$` to pair, and in
    CafeQ's prose it does not, so five Korean sentences parse as formulas.
    A table is a region this module can find exactly, which is why the line
    is drawn there.
    """
    stats = {'fonts': 0, 'accents': 0}

    def rewrite_font(m):
        stats['fonts'] += 1
        inner = m.group(2)
        # `{\rm \min}` -- the body is already an operator command, and
        # \mathrm{\min} is not valid TeX. Keep the command alone.
        if inner.startswith('\\'):
            return inner
        return '\\%s{%s}' % (_OLD_FONT_MAP[m.group(1)], inner)

    def rewrite_call(m):
        stats['fonts'] += 1
        return '\\%s' % _OLD_FONT_MAP[m.group(1)]

    def rewrite_plain(text):
        # Nested groups need more than one pass: {\rm a {\bf b}}.
        previous = None
        while previous != text:
            previous = text
            text = _OLD_FONT_RE.sub(rewrite_font, text)
        # `\cal{A}` keeps its braces; only the command name changes, so one
        # pass is enough and it must run AFTER the group form — otherwise
        # `{\rm x}` would have its command rewritten before the group rule
        # could see the shape it matches on.
        text = _OLD_FONT_CALL_RE.sub(rewrite_call, text)
        # Last, the bare switch with no group either side. It has to run after
        # both, or it would rewrite the command inside `{\rm x}` before the
        # group rule could recognise that shape.
        text = _OLD_FONT_BARE_RE.sub(rewrite_call, text)
        # Then brace an accent's command argument, AFTER the font rules and
        # not before: `\bar\cal{C}` only becomes `\bar\mathcal{C}` above, and
        # the accent rule has to be shown that form to catch it. A callable,
        # never a replacement string: the replacement carries a backslash.
        text, braced = _ACCENT_ON_COMMAND_RE.subn(
            lambda m: '\\%s{%s}' % (m.group(1), m.group(2)), text)
        stats['accents'] += braced
        return text

    held = sorted([(m.start(), m.end(), 'code')
                   for m in _CODE_REGION_RE.finditer(md_text)]
                  + [(m.start(), m.end(), 'table')
                     for m in _RAW_TABLE_RE.finditer(md_text)])

    pieces, cursor = [], 0
    for start, end, kind in held:
        if start < cursor:                # nested in a span already handled
            continue
        pieces.append(rewrite_plain(md_text[cursor:start]))
        block = md_text[start:end]
        if kind == 'table':
            # The cells stay as written; the formulas in them do not.
            block = _MATH_REGION_RE.sub(lambda m: rewrite_plain(m.group(0)),
                                        block)
        pieces.append(block)              # code: left exactly as written
        cursor = end
    pieces.append(rewrite_plain(md_text[cursor:]))
    return ''.join(pieces), stats


def rewrite_subcaptions(md_text, stats):
    """`\\subcaption{Cayley SGD.}` -> a bold paragraph of its own.

    Two things have to happen, not one. The command has to go, or it prints to
    the reader verbatim -- and it has to stop sharing a line with its image,
    because format_figure_blocks only recognises an image that is alone on its
    line. Left where pandoc put it, CafeQ's Figure 3 was three unlabelled
    pictures: no number, no caption, nothing tying them to the text.
    """
    out, cursor = [], 0
    while True:
        m = _SUBCAPTION_OPEN_RE.search(md_text, cursor)
        if not m:
            break
        close = _balanced_group(md_text, md_text.index('{', m.end() - 1))
        if close < 0:
            out.append(md_text[cursor:m.end()])
            cursor = m.end()
            continue
        inner = md_text[md_text.index('{', m.end() - 1) + 1:close - 1].strip()
        tail = md_text[close:]
        trail = re.match(r'`?(?:\{=latex\})?', tail)
        stats['dropped'] += 1
        # Whatever sat in front of it on the line stays put; the label moves
        # down into its own paragraph.
        lead = md_text[cursor:m.start()]
        out.append(lead.rstrip(' \t'))
        if inner:
            out.append('\n\n**%s**' % inner)
        cursor = close + trail.end()
    out.append(md_text[cursor:])
    return ''.join(out)


_FOOTNOTE_DEF_RE = re.compile(r'(?m)^\[\^([^\]\s]+)\]:[ \t]*')
_FOOTNOTE_REF_RE = re.compile(r'\[\^([^\]\s]+)\](?!:)')
_FIRST_HEADING_RE = re.compile(r'(?m)^#{1,6} ')


def _footnote_def_end(text, at):
    """Where the note definition whose body starts at `at` ends.

    pandoc's rule: continuation lines are indented, and a blank line only
    continues the note when an indented line follows it.
    """
    i = text.find('\n', at)
    while i >= 0:
        j = text.find('\n', i + 1)
        line = text[i + 1:j if j >= 0 else len(text)]
        if line.strip():
            if not line.startswith(('    ', '\t')):
                return i
        else:
            k = text.find('\n', j + 1) if j >= 0 else -1
            nxt = text[j + 1:k if k >= 0 else len(text)] if j >= 0 else ''
            if not nxt.startswith(('    ', '\t')):
                return i
        i = j
    return len(text)


def rescue_orphan_footnotes(md_text):
    r"""Render note definitions nothing references. Returns (text, count).

    An IEEE paper carries its front matter in `\thanks`: submission dates,
    every author's affiliation, the equal-contribution and corresponding-author
    notes, the funding, the DOI. pandoc reads each one as a footnote whose
    REFERENCE lives in the title block -- which the backend drops on purpose,
    because the title and authors come from the metadata. The definitions are
    left with nothing pointing at them, and pandoc drops an unreferenced note
    silently: TinyVLA translated 1271 characters of front matter and printed
    none of it. Nothing counted them, because a note that was never referenced
    is missing from every stage at once (K83).

    They are moved, not deleted: ahead of the first heading, which is where
    page-1 footnotes belong in the original. A note that IS referenced is left
    exactly where it is -- it is a working footnote and must stay one.
    """
    refs = set(_FOOTNOTE_REF_RE.findall(md_text))
    spans, bodies = [], []
    for m in _FOOTNOTE_DEF_RE.finditer(md_text):
        if m.group(1) in refs:
            continue
        end = _footnote_def_end(md_text, m.end())
        spans.append((m.start(), end))
        body = md_text[m.end():end]
        bodies.append(re.sub(r'(?m)^(?:    |\t)', '', body).strip())
    if not spans:
        return md_text, 0

    kept, cursor = [], 0
    for start, end in spans:
        kept.append(md_text[cursor:start])
        cursor = end
    kept.append(md_text[cursor:])
    md_text = re.sub(r'\n{3,}', '\n\n', ''.join(kept))

    block = '::: titlenotes\n' + '\n\n'.join(bodies) + '\n:::\n\n'
    heading = _FIRST_HEADING_RE.search(md_text)
    at = heading.start() if heading else len(md_text.rstrip()) + 1
    md_text = md_text[:at] + block + md_text[at:]
    return md_text, len(bodies)


def normalize_latex_leftovers(md_text):
    """Clear LaTeX that survived conversion and prints as literal text.

    Returns (text, stats). Conservative by design: it removes only commands
    that carry no reader-visible content, and rewrites only the two that do.
    """
    stats = {'parstart': 0, 'dropped': 0, 'brackets': 0}

    def parstart(m):
        stats['parstart'] += 1
        return m.group(1) + m.group(2)

    md_text = _PARSTART_RE.sub(parstart, md_text)

    md_text = rewrite_subcaptions(md_text, stats)

    def counted_drop(pattern, text):
        found = len(pattern.findall(text))
        stats['dropped'] += found
        return pattern.sub('', text)

    md_text, stats['index_terms'] = drop_index_terms(md_text)
    md_text, directives = drop_directive_spans(md_text)
    stats['dropped'] += directives
    md_text = _drop_latex_commands(md_text, _LEFTOVER_CMDS, stats)
    md_text = counted_drop(_ENV_WRAPPER_RE, md_text)
    # `\vspace*{-2em}` survives as a raw inline holding only its length. It is
    # page geometry with nothing in it to read, and leaving it costs more than
    # the stray `{-2em}` it prints: the figure formatter stops recognising the
    # image line it sits on, so CafeQ's figure 1 lost its `그림 1` label and
    # its caption printed as an ordinary paragraph.
    md_text = counted_drop(_SPACING_INLINE_RE, md_text)
    md_text = counted_drop(_SETLENGTH_RE, md_text)
    md_text, stats['sideset'] = rewrite_sideset(md_text)

    def keep_argument(m):
        stats['dropped'] += 1
        return m.group(1) or ''

    # Whole text, raw tables included — that is where it does its damage.
    md_text = _FONTSIZE_RE.sub(keep_argument, md_text)
    # Whole-text, before the slicing below: only here can the lookbehind see
    # whether a backtick owns this marker.
    md_text = counted_drop(_ORPHAN_RAW_ATTR_RE, md_text)
    # Outside code only: a fenced block may legitimately contain anything,
    # including two backticks, and rewriting inside one changes a listing.
    pieces, cursor = [], 0
    for code in _CODE_REGION_RE.finditer(md_text):
        pieces.append(counted_drop(_EMPTY_RAW_INLINE_RE,
                                   md_text[cursor:code.start()]))
        pieces.append(code.group(0))
        cursor = code.end()
    pieces.append(counted_drop(_EMPTY_RAW_INLINE_RE, md_text[cursor:]))
    md_text = ''.join(pieces)

    def unbracket(m):
        inner = m.group(1)
        # Math would carry operators, digits-with-symbols, or backslashes.
        # `%` and `°` belong here too: they are units, never a formula on
        # their own, and `\[%\]` -- the escaped `[\%]` of an Overhead column
        # -- was read back as display maths and rendered as nothing.
        if re.fullmatch(r'[\w \-/,.%°]+', inner):
            stats['brackets'] += 1
            return '[' + inner + ']'
        return m.group(0)

    md_text = _ESCAPED_BRACKET_RE.sub(unbracket, md_text)
    return md_text, stats


# =============================================================================
# Original document structure
# =============================================================================
#
# A translated paper read on its own gives no clue where you are in the
# original. IEEEtran (and every other class) numbers sections automatically, so
# the numbers exist nowhere in the markdown -- but flat.tex still has the
# heading ladder, and reproducing it is deterministic.
#
# Two rules keep this honest:
#   * strip LaTeX comments first. This paper has a `%\subsection{...}` the
#     authors commented out; counting it would consume a letter and shift every
#     heading after it.
#   * if the ladders do not line up 1:1, number nothing. A heading labelled "D"
#     that is really E is worse than no label at all.
#   * do not invent the numbering scheme. It belongs to the document class --
#     IEEEtran prints "III-B", ICML's article prints "2.1", and many classes
#     print nothing at all -- so it is read off the source PDF instead.

_HEADING_LINE_RE = re.compile(r'^(#{1,6}) +(.+?)\s*$', re.MULTILINE)
# MathML keeps the TeX source of every formula in an <annotation> element for
# copy-paste. Anything that turns heading HTML into plain text has to drop it
# first, or a heading with math in it reads as the glyph followed by its own
# source -- "γ\gamma의 ..." in the TOC, the bookmarks and the print TOC.
_ANNOTATION_RE = re.compile(r'<annotation\b[^>]*>.*?</annotation>', re.DOTALL)


# `\begin{comment}...\end{comment}` (comment.sty, verbatim.sty) hides its
# contents as completely as a `%` does — LaTeX typesets none of it, and pandoc
# correctly drops it. What reads flat.tex afterwards did not: Maynard leaves a
# 54-line block containing `\section{Motivation}`, so the heading list came
# back with eleven entries against the translation's ten and the build gave up
# on section numbering entirely — "refusing to guess" about a section the
# author had already deleted. The same block also hides two theorems and their
# labels, which would have numbered every theorem after them one too high.
_COMMENT_ENV_RE = re.compile(
    r'\\begin\s*\{comment\}.*?\\end\s*\{comment\}', re.DOTALL)


def strip_tex_comments(text):
    r"""Remove % comments and `comment` environments, honouring \\% escapes."""
    text = _COMMENT_ENV_RE.sub('', text)
    out = []
    for line in text.split('\n'):
        i, n = 0, len(line)
        while i < n:
            if line[i] == '\\' and i + 1 < n:
                i += 2
                continue
            if line[i] == '%':
                line = line[:i]
                break
            i += 1
        out.append(line)
    return '\n'.join(out)


# `\subsection{Additional \texttt{PL\_Alpha\_Hill} Comparisons}` -- the title
# argument nests braces, so it has to be read with a balanced scan. A plain
# `\{([^}]*)\}` stops at the inner `}` and yields a title that is both
# truncated and unbalanced.
_HEADING_CMD_RE = re.compile(r'\\((?:sub)*)(section|paragraph)(\*?)\s*\{')

# Wrappers whose argument IS the text. \texorpdfstring{tex}{pdf} keeps the
# first; the second exists only because the PDF bookmark cannot take math.
_TITLE_UNWRAP_RE = re.compile(
    r'\\(?:textbf|textit|textrm|texttt|textsc|emph|mbox|text|'
    r'operatorname|mathrm|lowercase|uppercase)\s*\{([^{}]*)\}')
_ENSUREMATH_RE = re.compile(r'\\ensuremath\s*\{([^{}]*)\}')
_TEXORPDF_RE = re.compile(r'\\texorpdfstring\s*\{')
_SIMPLE_MACRO_RE = re.compile(
    r'\\(?:new|renew|provide)command\s*\*?\s*\{\\([A-Za-z]+)\}\s*\{')
_TEX_ESCAPES = (('\\&', '&'), ('\\_', '_'), ('\\%', '%'), ('\\#', '#'),
                ('\\$', '$'), ('~', ' '))


def read_simple_macros(tex):
    r"""{name: body} for the paper's own zero-argument \newcommands.

    CafeQ writes `\newcommand{\tx}{\ensuremath{M}}` and then
    `\subsection{Constraints on \tx}`. pandoc expands it in the body, so the
    translated heading reads "Constraints on $M$" -- but flat.tex still says
    `\tx`, and the bilingual suffix would show the macro name.
    """
    out = {}
    for m in _SIMPLE_MACRO_RE.finditer(tex):
        open_at = tex.index('{', m.end() - 1)
        close = _balanced_group(tex, open_at)
        if close < 0:
            continue
        body = tex[open_at + 1:close - 1]
        if '#' not in body:                       # takes no argument
            out.setdefault(m.group(1), body)
    return out


def clean_heading_title(raw, macros=None):
    """Reduce a LaTeX section title to the text a reader would see."""
    text = raw
    for _ in range(4):                            # macros can nest
        if macros:
            expanded = re.sub(
                r'\\([A-Za-z]+)(?![A-Za-z])',
                lambda m: macros.get(m.group(1), m.group(0)), text)
        else:
            expanded = text
        # \texorpdfstring{$\gamma$}{gamma} -> keep the TeX form
        while True:
            m = _TEXORPDF_RE.search(expanded)
            if not m:
                break
            first_open = expanded.index('{', m.end() - 1)
            first_close = _balanced_group(expanded, first_open)
            if first_close < 0:
                break
            second_close = _balanced_group(expanded, first_close) \
                if first_close < len(expanded) and expanded[first_close] == '{' else first_close
            expanded = (expanded[:m.start()]
                        + expanded[first_open + 1:first_close - 1]
                        + expanded[second_close:])
        # `Constraints on $\tx$` expands to `$\ensuremath{M}$`, and wrapping
        # the body in dollars again gives `$$M$$` -- display math. pandoc
        # then emits a centred block inside the heading, and CafeQ printed
        # `M에 대한 제약 (Constraints on` on one line with a lone centred `M`
        # 27pt below it and the `)` after that. The source already opened
        # math here, so take its dollars rather than adding a pair.
        expanded = re.sub(r'\$\s*\\ensuremath\s*\{([^{}]*)\}\s*\$',
                          r'$\1$', expanded)
        expanded = _ENSUREMATH_RE.sub(r'$\1$', expanded)
        expanded = _TITLE_UNWRAP_RE.sub(r'\1', expanded)
        if expanded == text:
            break
        text = expanded
    text = re.sub(r'\\label\s*\{[^{}]*\}', '', text)
    for old, new in _TEX_ESCAPES:
        text = text.replace(old, new)

    # Strip unknown commands OUTSIDE math only. `Sensitivity of $\\gamma$` would
    # otherwise become `Sensitivity of $$` -- the command removed, the empty
    # delimiters left behind.
    def strip_outside_math(chunk):
        return re.sub(r'\\[A-Za-z]+\s*', '', chunk)

    pieces, cursor = [], 0
    for span in re.finditer(r'\$[^$\n]+\$', text):
        pieces.append(strip_outside_math(text[cursor:span.start()]))
        pieces.append(span.group(0))
        cursor = span.end()
    pieces.append(strip_outside_math(text[cursor:]))
    text = ''.join(pieces)
    return re.sub(r'\s+', ' ', text).strip()


def read_tex_headings(temp_dir):
    """[(level, original_title, is_numbered)] from flat.tex, in document order."""
    flat = os.path.join(temp_dir, 'flat.tex')
    if not os.path.exists(flat):
        return []
    try:
        with open(flat, 'r', encoding='utf-8', errors='replace') as fh:
            tex = strip_tex_comments(fh.read())
    except OSError:
        return []
    macros = read_simple_macros(tex)
    out = []
    for m in _HEADING_CMD_RE.finditer(tex):
        open_at = m.end() - 1
        close = _balanced_group(tex, open_at)
        if close < 0:
            continue
        title = tex[open_at + 1:close - 1].strip()
        # A macro *definition* body, not a heading: a real title never has #1.
        if '#' in title or not title:
            continue
        subs = m.group(1).count('sub')
        title = clean_heading_title(title, macros)
        if not title:
            continue
        if m.group(2) == 'section':
            # Starred sections are unnumbered in LaTeX -- Abstract, Index
            # Terms. Numbering them would shift every real section after them.
            out.append((subs + 1, title, m.group(3) != '*'))
        else:
            # \paragraph and \subparagraph sit below the default secnumdepth
            # of every class these papers use, so LaTeX prints them with no
            # number. They still occupy a rung -- pandoc turns them into ####
            # headings -- and leaving them out is what made the ladder come up
            # short on all three of SINQ, CafeQ and AlphaQ, disabling section
            # numbering entirely.
            out.append((4 + subs, title, False))
    return out


def _source_pdf(temp_dir):
    """The PDF this temp dir was built from, per config.txt."""
    config = os.path.join(temp_dir, 'config.txt')
    try:
        with open(config, 'r', encoding='utf-8', errors='replace') as fh:
            for line in fh:
                if line.startswith('input_file='):
                    path = line.split('=', 1)[1].strip()
                    if path.lower().endswith('.pdf') and os.path.isfile(path):
                        return path
    except OSError:
        pass
    return None


def _normalize_heading(text):
    """Fold a heading to something two renderings of it can agree on."""
    text = re.sub(r'\$[^$]*\$', ' ', text)          # math renders differently
    text = re.sub(r'[^0-9A-Za-z]+', ' ', text)
    return re.sub(r'\s+', ' ', text).strip().lower()


# "2.1.1." / "III." / "A.3" / "4" -- whatever the class chose to print.
_PDF_PREFIX_RE = re.compile(
    r'^\s*((?:[0-9]+|[A-Z]|[IVXLC]+)(?:[.\-][0-9A-Z]+)*\.?)[ \t]+(\S.*)$')


def read_pdf_section_prefixes(temp_dir, tex_heads):
    """[prefix or ''] per heading, read off the original PDF.

    Returns (prefixes, stats). An empty string means the paper prints that
    heading without a number -- which is a real answer, not a failure. None is
    returned for the whole list when the PDF cannot be consulted at all.
    """
    stats = {'matched': 0, 'unnumbered': 0, 'missing': 0,
             'wrapped': 0, 'reason': None}
    pdf_path = _source_pdf(temp_dir)
    if not pdf_path:
        stats['reason'] = 'no source PDF recorded in config.txt'
        return None, stats
    try:
        import pymupdf                                     # optional dependency
    except ImportError:
        stats['reason'] = 'pymupdf not installed'
        return None, stats

    try:
        doc = pymupdf.open(pdf_path)
        try:
            lines = []
            for page in doc:
                lines.extend(page.get_text('text').split('\n'))
        finally:
            doc.close()
    except Exception as exc:                               # noqa: BLE001
        stats['reason'] = 'could not read %s (%s)' % (os.path.basename(pdf_path), exc)
        return None, stats

    # title -> prefix. First occurrence wins: a heading is printed before it is
    # cited, and the table of contents (if any) agrees with the body anyway.
    numbered, plain = {}, set()
    for raw in lines:
        line = raw.strip()
        if not line or len(line) > 120:
            continue
        m = _PDF_PREFIX_RE.match(line)
        if m:
            key = _normalize_heading(m.group(2))
            if key and key not in numbered:
                numbered[key] = m.group(1).strip()
        key = _normalize_heading(line)
        if key:
            plain.add(key)

    prefixes = []
    for _level, title, _is_numbered in tex_heads:
        key = _normalize_heading(title)
        if key in numbered:
            prefixes.append(numbered[key])
            stats['matched'] += 1
            continue
        if key in plain:
            prefixes.append('')
            stats['unnumbered'] += 1
            continue
        # The column was too narrow and the PDF wrapped the title, so the
        # extracted line holds only its beginning.
        partial = _longest_prefix_match(key, numbered)
        if partial is not None:
            prefixes.append(numbered[partial])
            stats['matched'] += 1
            stats['wrapped'] += 1
            continue
        # The title carries maths, which `_normalize_heading` deletes on the
        # LaTeX side and the PDF renders as an ordinary letter.
        extended = _math_extended_match(title, key, numbered)
        if extended is not None:
            prefixes.append(numbered[extended])
            stats['matched'] += 1
            stats['wrapped'] += 1
            continue
        prefixes.append(None)
        stats['missing'] += 1
    return prefixes, stats


def _longest_prefix_match(key, numbered, minimum=14):
    """The longest heading in `numbered` that `key` starts with."""
    best = None
    for candidate in numbered:
        if len(candidate) >= minimum and key.startswith(candidate):
            if best is None or len(candidate) > len(best):
                best = candidate
    return best


def _math_extended_match(title, key, numbered, minimum=10):
    r"""The PDF heading that `key` is a prefix of, when maths was deleted.

    `_normalize_heading` drops `$...$` because two renderings of a formula
    never agree character for character. On the PDF side there is nothing to
    drop: `\section{Smooth choice of $y$}` prints as "Smooth choice of y", so
    the LaTeX key ends where the PDF key carries one more letter and neither
    the exact test nor the wrapped-title test can bridge it. Three of
    Maynard's ten sections are written that way, and all three came back as
    "not found" — which also cost the print TOC and the PDF outline their
    numbers until the same gap was closed on that side (K126).

    Only for titles that actually contain maths, and only when exactly one
    candidate extends the key: a heading that is a genuine prefix of another
    ("Notation" before "Notation and conventions") must stay unmatched rather
    than take its neighbour's number.
    """
    if '$' not in (title or '') or len(key) < minimum:
        return None
    hits = [c for c in numbered if c != key and c.startswith(key)]
    return hits[0] if len(hits) == 1 else None


def _already_says(translated, original):
    r"""Does the translated heading already carry the original's words?

    The translator glosses a term on first use, so `2.1.2. 타일링(tiling)`
    came back already carrying its English — and the bilingual suffix then
    added it a second time: `타일링(tiling) (Tiling)`, in the heading and
    again in the table of contents. Compare with spaces and case removed,
    which is what makes `No-Overhead SINQ` match `(No-Overhead SINQ)`.
    """
    flat = lambda s: re.sub(r'\s+', '', s).lower()
    return flat(original) in flat(translated)


# A heading's original-language gloss is there so a reader can match it to the
# paper; a cross-reference left raw inside it helps nobody. GAN's "Convergence
# of Algorithm \ref{alg:AGF}" reached the page as "(Convergence of Algorithm
# {alg:AGF})": the Korean half had its reference resolved, because
# resolve_references runs BEFORE this pass, and the gloss is lifted from the
# original afterwards, where nothing has touched it.
# `[a-z]{2,12}` and a key of word characters both assumed a naming convention.
# Vershynin labels his sections `{s: sums matrices}` -- a one-letter prefix and
# spaces in the key -- so ten headings kept a raw `{s: introduction}` beside
# their translated title.
_GLOSS_REF_RE = re.compile(
    r'\s*\\[a-zA-Z]*ref\*?\s*\{([^{}]*)\}'
    r'|\s*\{([a-z]{1,12}:[^{}]{1,60})\}')


def _gloss_reference(original, numbers):
    r"""The source heading as a reader should see it.

    A heading like `For Section~\ref{s: introduction}` carries a pointer, not a
    word. Print the number it stands for -- the surrounding text already reads
    "For Section" -- and fall back to dropping it when the label is unknown,
    which is what this did for every reference before.
    """
    def sub(m):
        key = (m.group(1) or m.group(2) or '').strip()
        number = numbers.get(key) if numbers else None
        return ' %s' % number if number else ''

    return _GLOSS_REF_RE.sub(sub, original).strip()


def theorem_declarations(temp_dir):
    r"""[(pandoc_number, printed_label)] for every theorem-like, in order.

    Two tallies over the same walk. `pandoc_number` is what pandoc wrote into
    the chunk: it honours the shared counter of `\newtheorem{lmm}[thrm]{Lemma}`
    but knows nothing of `[section]`, so it runs 1..N across the paper. The
    printed label is what the paper prints — the same counter with the section
    reset and prefix `_counter_label` already applies to equations.

    A starred declaration gets `''`: `\newtheorem*{rmk}{Remark}` prints no
    number at all, and pandoc invents one anyway — six of them in Maynard, and
    no check could see them because nothing `\ref`s an unnumbered environment.
    """
    flat = os.path.join(temp_dir, 'flat.tex')
    if not os.path.exists(flat):
        return []
    try:
        with open(flat, 'r', encoding='utf-8', errors='replace') as fh:
            tex = strip_tex_comments(fh.read())
    except OSError:
        return []

    envs = read_theorem_environments(tex)
    starred = _read_starred_theorem_envs(tex)
    if not envs and not starred:
        return []
    parents = read_counter_parents(tex)
    fixed = read_fixed_counter_prefix(tex)

    names = sorted(set(envs) | set(starred), key=lambda n: (-len(n), n))
    scan = re.compile(
        r'\\((?:sub)*)section(\*?)\s*\{'
        r'|\\begin\{(' + '|'.join(re.escape(n) for n in names) + r')\}')

    section_head = ''
    pandoc_n, paper_n = {}, {}
    out = []
    for m in scan.finditer(tex):
        if m.group(1) is not None and m.group(2) is not None:
            if m.group(2) == '*' or m.group(1).count('sub'):
                continue
            section_head = str(int(section_head or 0) + 1)
            for counter, parent in parents.items():
                if parent == 'section':
                    paper_n[counter] = 0
            continue
        env = m.group(3)
        if env in starred:
            pandoc_n[env] = pandoc_n.get(env, 0) + 1
            out.append((pandoc_n[env], ''))
            continue
        group = envs.get(env, 'theorem')
        pandoc_n[group] = pandoc_n.get(group, 0) + 1
        paper_n[group] = paper_n.get(group, 0) + 1
        out.append((pandoc_n[group],
                    _counter_label(group, paper_n[group], parents,
                                   section_head, fixed)))
    return out


def _read_starred_theorem_envs(tex):
    r"""Names declared with `\newtheorem*`, which print without a number."""
    out = set()
    for m in _NEWTHEOREM_RE.finditer(tex):
        if m.group(1) and m.group(2).strip():
            out.add(m.group(2).strip())
    return out


def number_theorem_statements(md_text, temp_dir, lang_cfg=None):
    r"""Give each `**Theorem 1**` the number the paper prints. (text, stats).

    The reference half of this was fixed first and on its own made the book
    WORSE: prose saying "정리 1.1" over a declaration line reading "정리 1" is
    less usable than two numbers that agree with each other and with nothing
    else. Both halves or neither.

    `**정리 1**` is characters in `output.md`, and this build rewrites that
    line as a matter of routine — the claim that pandoc owns theorem numbers
    (K113) was about pandoc's output, not about the book (K130).

    Refuses wholesale, the way `number_sections` does, and on a stronger
    condition: the site COUNT must match and every number already printed must
    equal the flat tally. If pandoc numbered anything differently from the way
    modelled here, nothing is rewritten and the reason is reported.
    """
    stats = {'numbered': 0, 'unnumbered': 0, 'skipped_reason': None}
    wanted = theorem_declarations(temp_dir)
    if not wanted:
        stats['skipped_reason'] = 'no theorem environments declared'
        return md_text, stats

    words = (lang_cfg or {}).get('theorem_words') or ()
    if not words:
        stats['skipped_reason'] = 'no theorem vocabulary for this language'
        return md_text, stats
    alt = '|'.join(re.escape(w) for w in
                   sorted(words, key=lambda w: (-len(w), w)))
    site_re = re.compile(r'\*\*(' + alt + r')\s+(\d+)\*\*')

    sites = list(site_re.finditer(md_text))
    if len(sites) != len(wanted):
        stats['skipped_reason'] = (
            '%d theorem-like(s) in flat.tex vs %d numbered statement(s) in the '
            'translation — refusing to guess' % (len(wanted), len(sites)))
        return md_text, stats
    for m, (flat_n, _label) in zip(sites, wanted):
        if int(m.group(2)) != flat_n:
            stats['skipped_reason'] = (
                'statement %r carries %s where the tally says %d — refusing '
                'to guess' % (m.group(0), m.group(2), flat_n))
            return md_text, stats

    pieces, cursor = [], 0
    for m, (_flat_n, label) in zip(sites, wanted):
        pieces.append(md_text[cursor:m.start()])
        if label:
            pieces.append('**%s %s**' % (m.group(1), label))
            stats['numbered'] += 1
        else:
            pieces.append('**%s**' % m.group(1))
            stats['unnumbered'] += 1
        cursor = m.end()
    pieces.append(md_text[cursor:])
    return ''.join(pieces), stats


def number_sections(md_text, temp_dir, bilingual=True):
    """Prefix each heading with its number in the original, and the original
    title. Returns (text, stats)."""
    stats = {'numbered': 0, 'skipped_reason': None}
    tex_heads = read_tex_headings(temp_dir)
    if not tex_heads:
        stats['skipped_reason'] = 'no flat.tex (not an arXiv-sourced build)'
        return md_text, stats

    md_heads = list(_HEADING_LINE_RE.finditer(md_text))
    if len(md_heads) != len(tex_heads):
        stats['skipped_reason'] = (
            f'{len(tex_heads)} headings in flat.tex vs {len(md_heads)} in the '
            f'translation — refusing to guess')
        return md_text, stats
    for m, (level, _title, _numbered) in zip(md_heads, tex_heads):
        if len(m.group(1)) != level:
            stats['skipped_reason'] = (
                f'heading levels diverge at {m.group(2)!r} — refusing to guess')
            return md_text, stats

    labels, pdf_stats = read_pdf_section_prefixes(temp_dir, tex_heads)
    if labels is None:
        stats['skipped_reason'] = (
            'cannot check numbering against the original (%s) -- refusing to '
            'invent a scheme' % pdf_stats['reason'])
        return md_text, stats
    found = pdf_stats['matched'] + pdf_stats['unnumbered']
    if found < max(3, int(0.6 * len(tex_heads))):
        stats['skipped_reason'] = (
            'only %d of %d headings could be located in the original PDF '
            '— refusing to guess' % (found, len(tex_heads)))
        return md_text, stats
    stats['pdf'] = pdf_stats
    labels = ['' if lb is None else lb for lb in labels]
    gloss_numbers = build_label_numbers(temp_dir)
    pieces, cursor = [], 0
    for m, label, (_lvl, original, _num) in zip(md_heads, labels, tex_heads):
        translated = m.group(2).strip()
        if re.match(r'^(?:[IVXLC]+|[A-Z]|\d+)[.)] ', translated):
            continue  # already numbered; do not double up
        text = f'{label} {translated}'.strip()
        if bilingual and original and not _already_says(translated, original):
            gloss = _gloss_reference(original, gloss_numbers)
            if gloss:
                text += f' ({gloss})'
        pieces.append(md_text[cursor:m.start()])
        pieces.append(f'{m.group(1)} {text}')
        cursor = m.end()
        stats['numbered'] += 1
    pieces.append(md_text[cursor:])
    return ''.join(pieces), stats


# =============================================================================
# Figure captions
# =============================================================================
#
# The arXiv path emits a bare image followed by a bold paragraph, which prints
# as ordinary body text -- there is nothing to tell a reader where the caption
# stops and the argument resumes. Wrapping the pair in <figure>/<figcaption>
# gives the print sheet something to style, and the float order in flat.tex
# gives the number.

# [ \t]*$ rather than \s*$: a trailing \s* swallows the blank lines
# after the image, and then the caption-gap guard below has nothing left to
# measure -- an uncaptioned figure could reach forward and take the next
# paragraph as its caption.
# The trailing raw span is a spacing directive the source put on the image's
# own line -- `` `{-2em}`{=latex} ``. Requiring the line to end at the image
# meant a figure carrying one was never recognised as a figure AT ALL: no
# number, no printed label, no anchor, and its caption left behind as loose
# prose. CafeQ's figure 1 went that way while three cross-references kept
# pointing at it. Only a brace-wrapped span is allowed in, so a raw span
# holding real content cannot be swallowed here.
_FIG_IMAGE_RE = re.compile(
    r'^!\[([^\]]*)\]\((images/fig(\d+)[^)]*)\)(?:\{[^}]*\})?'
    r'(?:[ \t]*`\{[^`\n]*\}`\{=[a-z]+\})?[ \t]*$',
    re.MULTILINE)

# How wide to draw one panel of a multi-panel float, by panel count.
#
# A panel drawn at full text width is about 125mm tall against a 257mm text
# block, so two cannot share a page: SINQ printed three panels of one figure
# on three pages, and seven of its thirty-six pages held a single picture and
# fourteen characters of caption.
#
# Each width is the largest that still lets the whole float sit on one page,
# from panels x (0.125 x width + 15mm of caption and margin) <= 257mm. Beyond
# four panels no width both fits and stays legible, so they fill a page at a
# time; 38% of the 174mm text width is 66mm, still wider than the ~40mm these
# same panels get in the printed original.
_PANEL_WIDTH = {2: 80, 3: 55}
_PANEL_WIDTH_MANY = 38


def _figure_number_from_path(path, index_fallback):
    m = re.search(r'fig(\d+)', path)
    return int(m.group(1)) if m else index_fallback


# A vertical-space directive the source parked between an image and its
# caption: `{-2em}` on its own, or `` `{-2em}`{=latex} `` once pandoc has
# wrapped it as a raw span. It holds no words, so it is not a caption -- but
# it sits exactly where the caption would be, and the blank line after it
# ended the search. CafeQ's figure 1 lost its number, its printed label and
# the anchor three cross-references pointed at, for that alone: the same
# figure in the same paper kept all three in the previous build, where the
# directive happened to arrive without a blank line after it.
_SPACING_SPAN_RE = re.compile(
    r'^(?:`(?P<raw>[^`\n]*)`\{=[a-z]+\}|(?P<bare>\{[^{}\n]*\}))[ \t]*(?=\n|$)')
_SPACING_BODY_RE = re.compile(
    r'^\{?\s*(?:\\(?:vspace|hspace|vskip|hskip)\*?)?\s*'
    r'[-+]?\d*\.?\d*\s*(?:em|ex|pt|cm|mm|in|bp|sp|mu|\\baselineskip)?\s*\}?$')


def _spacing_only_prefix(tail):
    """The leading run of tail that is pure spacing, or '' if there is none."""
    m = _SPACING_SPAN_RE.match(tail)
    if not m:
        return ''
    body = m.group('raw') if m.group('raw') is not None else m.group('bare')
    if not _SPACING_BODY_RE.match((body or '').strip()):
        return ''
    return m.group(0)


_GRAPHIC_RE = re.compile(
    r'\\includegraphics\s*(?:\[([^\]]*)\])?\s*\{([^}]*)\}')


def _graphic_stem(path, options=''):
    """'figures/corr_alignment_Qwen3-1.7B.pdf' -> 'corralignmentqwen317b'.

    A page beyond the first joins the stem, because the backend gives that
    panel its own file. CafeQ draws two panels of one figure from page 1 and
    page 4 of the same PDF; keyed on the path alone they collapse into one
    panel and the second image is left unclaimed.
    """
    stem = os.path.basename(path).rsplit('.', 1)[0]
    page = re.search(r'\bpage\s*=\s*(\d+)', options or '')
    if page and page.group(1) != '1':
        stem += '_p' + page.group(1)
    return re.sub(r'[^0-9a-z]', '', stem.lower())


_FLOAT_ENV_RE = re.compile(
    r'\\begin\{((?:SC|wrap|sideways|long|floating)?(?:figure|table)\*?)\}'
    r'(.*?)\\end\{\1\}', re.DOTALL)
# \caption, plus \captionof{figure} for a float built out of a plain box.
_CAPTION_CMD_RE = re.compile(
    r'\\caption(?:of)?\s*(?:\{(?:figure|table)\}\s*)?(?:\[[^\]]*\])?\s*\{')


# The commented tail of a line, for counting braces that TeX never sees.
_COMMENT_LINE_RE = re.compile(r'(?<!\\)%[^\n]*')

# Figures the paper DRAWS rather than includes. There is no image file for one
# anywhere in the source, so no stage here can render it; its absence from the
# book is a limitation to report, not a fault to stop the build over.
_CODE_DRAWN_FIGURES = {'tikzpicture', 'pgfpicture', 'pspicture', 'picture'}

_PROOF_ENV_RE = re.compile(
    r'^[ \t]*\\begin\{proof\}(?:\s*\[[^\]]*\])?[ \t]*$\n?'
    r'|^[ \t]*\\end\{proof\}[ \t]*$\n?', re.MULTILINE)
_LIST_ENV_RE = re.compile(
    r'^[ \t]*\\begin\{(enumerate|itemize)\}(?:\s*\[[^\]]*\])?[ \t]*$\n'
    r'(.*?)'
    r'^[ \t]*\\end\{\1\}[ \t]*$\n?', re.MULTILINE | re.DOTALL)
_ITEM_RE = re.compile(r'^[ \t]*\\item[ \t]*', re.MULTILINE)
_LISTING_ENV_RE = re.compile(
    r'^[ \t]*\\begin\{(lstlisting|verbatim|minted)\}'
    r'(?:\s*\[[^\]]*\])?(?:\s*\{[^{}]*\})?[ \t]*$\n'
    r'(.*?)'
    r'^[ \t]*\\end\{\1\}[ \t]*$\n?', re.MULTILINE | re.DOTALL)


def unwrap_prose_environments(md_text):
    r"""Turn leftover LaTeX prose environments into markdown. (text, count).

    pandoc's markdown reader takes `\begin{env}…\end{env}` as ONE raw LaTeX
    block and the HTML writer drops it whole, without a word. Neural ODE's
    appendix lost a proof, a Python listing, a numbered list and a figure that
    way — all of them already translated.

    They are not exotic: a proof is paragraphs, a list is a list, a listing is
    a code block. Written as markdown they render everywhere, DOCX included,
    which raw LaTeX never does.
    """
    count = [0]

    def drop_wrapper(m):
        count[0] += 1
        return ''

    def as_list(m):
        count[0] += 1
        kind, body = m.group(1), m.group(2)
        marker = '1. ' if kind == 'enumerate' else '- '
        return '\n' + _ITEM_RE.sub(marker, body).strip('\n') + '\n\n'

    def as_code(m):
        count[0] += 1
        return '\n```\n' + m.group(2).strip('\n') + '\n```\n\n'

    md_text = _LISTING_ENV_RE.sub(as_code, md_text)
    md_text = _LIST_ENV_RE.sub(as_list, md_text)
    # The proof's own wrapper only; its paragraphs stay where they are.
    md_text = _PROOF_ENV_RE.sub(drop_wrapper, md_text)
    return md_text, count[0]

# `\fontsize{6pt}{1pt}\selectfont{system}` — a size switch with its argument.
# It sets type and shows nothing of its own, but pandoc's tabular reader loses
# the ROW it opens: ResNet's per-class detection table converted "successfully"
# and reached the page with its header and every class column gone. Keep the
# argument, drop the switch. The trailing group is optional because the switch
# is as often used bare, to change size for the rest of the cell.
# pandoc hands it over as a backticked raw inline, so the BACKTICKS have to go
# with it. Removing the command alone left `` on the page — two stray marks
# the empty-span cleanup could not take, because it only lifts a pair standing
# alone between spaces and this one had text hard against it.
_FONTSIZE_RE = re.compile(
    r'[ \t]*`?\\fontsize\s*\{[^{}]*\}\s*\{[^{}]*\}\s*\\selectfont\s*'
    r'(?:\{((?:[^{}]|\{[^{}]*\})*)\})?[ \t]*`?(?:\{=[a-z]+\})?')


def _in_latex_comment(text, pos):
    """Is `pos` on the commented-out part of its line?

    `float_units` takes comment-stripped text and so never had to ask. The
    passes that read the MERGED MARKDOWN cannot strip: the raw floats have to
    reach pandoc byte for byte. They need the question answered in place.

    A `%` opens a comment unless it is escaped, and a backslash escapes only
    when it is not itself escaped -- so what decides is whether the run of
    backslashes in front of the `%` is even.
    """
    line_start = text.rfind('\n', 0, pos) + 1
    scan = line_start
    while True:
        at = text.find('%', scan, pos)
        if at < 0:
            return False
        back = at
        while back > line_start and text[back - 1] == '\\':
            back -= 1
        if (at - back) % 2 == 0:
            return True
        scan = at + 1


_THE_COUNTER_RE = re.compile(
    r'\\(?:re)?newcommand\s*\{?\s*\\the(figure|table)\s*\}?\s*'
    r'(\{(?:[^{}]|\{[^{}]*\})*\})')
_SETCOUNTER_RE = re.compile(r'\\setcounter\s*\{(figure|table)\}\s*\{(-?\d+)\}')


def counter_events(tex):
    r"""Explicit counter declarations, in source order.

    A paper can letter its appendix floats by hand instead of scoping the
    counter to a section. VLA-Adapter gives each of its nine appendix
    sections a `\renewcommand{\thefigure}{A\arabic{figure}}` and a
    `\setcounter{figure}{0}`, lettering A through I, so what it prints as
    Figure A1 is this counter's ninth figure. Numbering straight through
    sent seventeen cross-references to the wrong float, and `source_probe`
    caught it by reading the number off the original PDF.

    Returns [(position, kind, 'prefix'|'set', value)]. A redefinition whose
    prefix is itself a command is skipped rather than guessed at: the point
    is to read what the source declares, not to evaluate TeX.
    """
    events = []
    for m in _THE_COUNTER_RE.finditer(tex):
        kind, body = m.group(1), m.group(2)
        inner = re.search(r'\\arabic\s*\{\s*' + kind + r'\s*\}', body)
        if not inner:
            continue          # not `<prefix>\arabic{kind}`; leave it alone
        prefix = body[1:inner.start()].strip()
        if '\\' in prefix:
            continue          # the prefix is a command; do not evaluate it
        events.append((m.start(), kind, 'prefix', prefix))
    for m in _SETCOUNTER_RE.finditer(tex):
        events.append((m.start(), m.group(1), 'set', int(m.group(2))))
    events.sort(key=lambda e: e[0])
    return events


def float_units(tex):
    """One entry per figure/table number the paper actually issues.

    Counting float ENVIRONMENTS is the obvious reading and it is wrong twice
    over. LaTeX numbers a float when \\caption runs, so a float can be worth
    two numbers or none:

      * AlphaQ puts two `minipage`s inside one `table*`, each with its own
        \\caption. That is two tables, and reading it as one numbered every
        later table two too low.
      * SINQ leaves two figures commented out. Those number nothing, and
        counting them numbered every figure after the first one too high.

    Both shipped, because the caption side of the pipeline had stripped
    comments and the cross-reference side had not: the caption under the plot
    read "그림 6" while the sentence pointing at it read "그림 7".

    Pass comment-stripped `tex`. Returns [{'kind','number','start','stop',
    'labels'}] where start..stop is the slice of the float owned by that
    caption, and `number` is None for a float that has no caption at all.
    """
    units, counters = [], {'figure': 0, 'table': 0}
    # A float counter can be scoped to the section, in which case the paper
    # prints `Table 3.1` and restarts at every section. This function sees
    # floats but not sections, so the boundaries have to be walked alongside
    # them; without it Shor's four table and figure references named numbers
    # the paper does not print.
    parents = read_counter_parents(tex)
    scoped = {k for k in ('figure', 'table') if parents.get(k) == 'section'}
    sections = []
    if scoped:
        depth0 = 0
        for m in re.finditer(r'\\(?:sub)*section(\*?)\s*\{', tex):
            if m.group(1):                 # starred: numbers nothing
                continue
            if m.group(0).count('sub') == 0:
                depth0 += 1
                sections.append((m.start(), str(depth0)))
    section_head, next_section = '', 0
    # The other way a paper letters its floats: by declaring it, rather than
    # by scoping the counter to a section. Walked alongside the floats for the
    # same reason the sections are.
    events, next_event = counter_events(tex), 0
    prefixes = {'figure': '', 'table': ''}

    for float_match in _FLOAT_ENV_RE.finditer(tex):
        while next_section < len(sections) \
                and sections[next_section][0] < float_match.start():
            section_head = sections[next_section][1]
            for name in scoped:
                counters[name] = 0
            next_section += 1
        while next_event < len(events) \
                and events[next_event][0] < float_match.start():
            _at, event_kind, action, value = events[next_event]
            if action == 'prefix':
                prefixes[event_kind] = value
            else:
                counters[event_kind] = value
            next_event += 1
        kind = 'table' if 'table' in float_match.group(1).lower() else 'figure'
        body, base = float_match.group(2), float_match.start(2)
        panels = _panel_spans(body)
        # A \caption inside a subfigure is a subcaption; it letters the panel
        # rather than numbering the float.
        captions = [m for m in _CAPTION_CMD_RE.finditer(body)
                    if not any(s <= m.start() < e for s, e, _t in panels)]
        # Split at caption STARTS, which is the one rule that survives both
        # layouts in the wild: content-then-caption and caption-then-content.
        bounds = [0] + [m.start() for m in captions[1:]] + [len(body)]
        for index in range(max(1, len(captions))):
            if captions:
                counters[kind] += 1
            region = tex[base + bounds[index]:base + bounds[index + 1]]
            if not captions:
                number = None                 # a float that numbers nothing
            elif kind in scoped:
                number = _counter_label(kind, counters[kind], parents,
                                        section_head)
            elif prefixes[kind]:
                # The paper declared the prefix itself: `A\arabic{figure}`.
                number = '%s%d' % (prefixes[kind], counters[kind])
            else:
                number = counters[kind]
            units.append({
                'kind': kind,
                'number': number,
                'start': base + bounds[index],
                'stop': base + bounds[index + 1],
                'labels': [l.strip() for l in
                           re.findall(r'\\label\{([^}]+)\}', region)],
            })
    return units


def read_float_units(temp_dir):
    """float_units() over a temp dir's flat.tex, or [] when there is none."""
    flat = os.path.join(temp_dir or '', 'flat.tex')
    if not temp_dir or not os.path.exists(flat):
        return []
    try:
        with open(flat, 'r', encoding='utf-8', errors='replace') as fh:
            return float_units(strip_tex_comments(fh.read()))
    except OSError:
        return []


def figure_panels(temp_dir):
    """{image filename: {'float': 2, 'letter': 'b', 'panels': 3, 'caption': True}}

    A float with three \\includegraphics is ONE figure with three panels, not
    three figures. Panels are matched to extracted files by name: SINQ has 17
    \\includegraphics but 13 image files, because tikz pictures and unresolved
    graphics leave nothing behind, so counting positions drifts after the
    first one that produced no file.

    Returns None when flat.tex or images/ is missing, meaning "fall back".
    """
    flat = os.path.join(temp_dir or '', 'flat.tex')
    images_dir = os.path.join(temp_dir or '', 'images')
    if not temp_dir or not os.path.exists(flat) or not os.path.isdir(images_dir):
        return None
    try:
        with open(flat, 'r', encoding='utf-8', errors='replace') as fh:
            tex = strip_tex_comments(fh.read())
        files = sorted(os.listdir(images_dir))
    except OSError:
        return None

    by_stem = {}
    for name in files:
        # extracted as fig0003_random_walk_hyp.png
        body = re.sub(r'^fig\d+_', '', name).rsplit('.', 1)[0]
        by_stem.setdefault(re.sub(r'[^0-9a-z]', '', body.lower()), name)

    out = {}
    for unit in float_units(tex):
        if unit['kind'] != 'figure':
            continue
        resolved = []
        for options, graphic in _GRAPHIC_RE.findall(tex[unit['start']:unit['stop']]):
            name = by_stem.get(_graphic_stem(graphic, options))
            # The same graphic can be included twice in one float
            # (an overlay, a repeated panel); it is still one panel.
            if name and name not in out and name not in resolved:
                resolved.append(name)
        for index, name in enumerate(resolved):
            out[name] = {'float': unit['number'],
                         'letter': chr(ord('a') + index) if len(resolved) > 1 else None,
                         'panels': len(resolved),
                         'caption': (unit['number'] is not None
                                     and index == len(resolved) - 1)}
    return out or None


def figures_with_captions(temp_dir):
    """{figure_number} for the floats that carry a \\caption in the source.

    Most LaTeX captions open with \\textbf{...} and so arrive as
    `**Lead-in.** rest`, which is easy to spot -- but a caption that simply
    lacks bold is still a caption, and figure 10 of a real paper was losing its
    caption for exactly that reason. Asking flat.tex which floats HAVE one
    turns a guess into a lookup.

    Returns None when flat.tex is absent, meaning "fall back to the heuristic".
    """
    flat = os.path.join(temp_dir or '', 'flat.tex')
    if not temp_dir or not os.path.exists(flat):
        return None
    return {unit['number'] for unit in read_float_units(temp_dir)
            if unit['kind'] == 'figure' and unit['number'] is not None}


_INCLUDEGRAPHICS_RE = re.compile(
    r'\\includegraphics\s*\[([^\]]*)\]\s*\{([^{}]*)\}')
_TRIM_OPT_RE = re.compile(r'trim\s*=\s*\{([^}]*)\}|trim\s*=\s*([^,\]]+)')
_LEN_RE = re.compile(r'(-?[\d.]+)\s*(cm|mm|in|pt|bp|px)?')
_PT_PER = {'cm': 28.3464567, 'mm': 2.83464567, 'in': 72.0,
           'pt': 1.0, 'bp': 1.0, 'px': 1.0, None: 1.0}


def _trim_of(options):
    r"""The four `trim=` lengths in points, LaTeX order: left bottom right top."""
    m = _TRIM_OPT_RE.search(options)
    if not m or 'clip' not in options:
        return None
    parts = (m.group(1) or m.group(2) or '').replace(',', ' ').split()
    if len(parts) != 4:
        return None
    out = []
    for part in parts:
        hit = _LEN_RE.match(part.strip())
        if not hit:
            return None
        out.append(float(hit.group(1)) * _PT_PER.get(hit.group(2), 1.0))
    return out


def apply_graphics_trim(md_text, temp_dir):
    r"""Crop each image the way its `\includegraphics[trim=...,clip]` asked.

    Returns (text, cropped). The extracted PNG is the whole source page, so a
    figure the paper crops arrives with whatever the author cropped OFF still
    on it: CafeQ's figure 1 carries the plot's own debug title band --
    `dperf ~ qerr; m=leaderboard_all_nodrop_mean; n=4; q=-1` -- which the
    original hides. The reader sees something the paper does not show.
    """
    try:
        import fitz
    except ImportError:
        return md_text, 0
    flat_path = os.path.join(temp_dir or '', 'flat.tex')
    if not os.path.isfile(flat_path):
        return md_text, 0
    with open(flat_path, encoding='utf-8', errors='replace') as fh:
        flat = strip_tex_comments(fh.read())

    wanted = {}
    for m in _INCLUDEGRAPHICS_RE.finditer(flat):
        trim = _trim_of(m.group(1))
        if not trim:
            continue
        page = re.search(r'page\s*=\s*(\d+)', m.group(1))
        stem = os.path.splitext(os.path.basename(m.group(2)))[0]
        wanted[(stem, int(page.group(1)) if page else 1)] = (trim, m.group(2))

    if not wanted:
        return md_text, 0

    html_srcs, md_srcs, _ = _scan_image_refs(md_text)
    cropped = 0
    for ref in sorted(set(html_srcs) | set(md_srcs)):
        name = os.path.basename(ref)
        hit = re.match(r'fig\d+_(.+?)(?:_p(\d+))?\.[a-z]+$', name)
        if not hit:
            continue
        key = (hit.group(1), int(hit.group(2) or 1))
        if key not in wanted:
            continue
        trim, source = wanted[key]
        image = os.path.normpath(os.path.join(temp_dir, ref))
        origin = _source_file(temp_dir, source)
        if not os.path.isfile(image) or not origin:
            continue
        out_ref = '%s_trim%s' % os.path.splitext(ref)
        out_path = os.path.normpath(os.path.join(temp_dir, out_ref))
        # Re-render the cropped region from the source rather than cutting
        # the PNG: same resolution, and no dependence on which Pixmap
        # constructors this build of PyMuPDF happens to offer.
        if not os.path.isfile(out_path):
            try:
                doc = fitz.open(origin)
                page = doc[min(key[1], doc.page_count) - 1]
                rect = page.rect
                scale = fitz.Pixmap(image).width / rect.width
                left, bottom, right, top = trim
                clip = fitz.Rect(rect.x0 + left, rect.y0 + top,
                                 rect.x1 - right, rect.y1 - bottom)
                if clip.is_empty or clip.width < 2 or clip.height < 2:
                    doc.close()
                    continue
                page.get_pixmap(matrix=fitz.Matrix(scale, scale),
                                clip=clip).save(out_path)
                doc.close()
            except Exception as exc:
                print(f"WARNING: could not crop {ref}: {exc}")
                continue
        md_text = md_text.replace(ref, out_ref)
        cropped += 1
    return md_text, cropped


def _source_file(temp_dir, source):
    """The extracted-source file an `\\includegraphics` path points at."""
    stem = os.path.splitext(os.path.basename(source))[0]
    for root, _dirs, files in os.walk(os.path.join(temp_dir, 'arxiv_src')):
        for name in files:
            if os.path.splitext(name)[0] == stem:
                return os.path.join(root, name)
    return None


def _figure_caption_md(label, body):
    r"""`**label** body`, safe to sit inside an image's alt text."""
    body = (body or '').strip()
    caption = f'**{label}** {body}'.strip() if body else f'**{label}**'
    # Square brackets would close the alt-text span early. Escaping is safe
    # here: normalize_latex_leftovers has already run, so nothing will re-read
    # `\[` as display math afterwards.
    caption = caption.replace('[', r'\[').replace(']', r'\]')
    # One line: implicit_figures needs the image alone in its paragraph, and a
    # newline inside the alt text would end it.
    return re.sub(r'\s*\n\s*', ' ', caption)


def _last_of_float(matches, n, panels):
    """Is the n-th match (1-based) the last panel of its float?"""
    def float_of(match):
        panel = (panels or {}).get(os.path.basename(match.group(2)))
        return panel['float'] if panel and panel['panels'] > 1 else None

    if n >= len(matches):
        return True
    return float_of(matches[n - 1]) != float_of(matches[n])


def format_figure_blocks(md_text, lang_cfg=None, temp_dir=None):
    """Fold each image + caption paragraph into one markdown image.

    Emitted as MARKDOWN, not raw HTML. pandoc's `implicit_figures` turns a
    paragraph holding nothing but an image into <figure>/<figcaption> natively,
    and the alt text keeps going through the markdown and math readers on the
    way. Hand-built HTML looked equivalent and was not: raw HTML is dropped
    entirely on the DOCX path (every figure vanished, 5.4MB -> 25KB) and skips
    the math reader, so `$...$` inside a caption printed literally.

    Returns (text, count).
    """
    lang_cfg = lang_cfg or {}
    fig_label = lang_cfg.get('figure_label', 'Figure')
    captioned = figures_with_captions(temp_dir)
    panels = figure_panels(temp_dir)
    matches = list(_FIG_IMAGE_RE.finditer(md_text))
    if not matches:
        return md_text, 0

    pieces, cursor, count = [], 0, 0
    open_group, group_caption = False, ''
    for n, m in enumerate(matches, 1):
        if m.start() < cursor:
            continue
        alt, path = m.group(1), m.group(2)
        panel = (panels or {}).get(os.path.basename(path))
        if panel:
            number, letter = panel['float'], panel['letter']
            wants_caption = panel['caption']
        else:
            number, letter = _figure_number_from_path(path, n), None
            wants_caption = None          # decided below, as before

        # The caption is the paragraph after the image. Skip whatever gap is
        # there, but only across a blank line or two, so an uncaptioned figure
        # cannot reach forward and steal the next section's opening sentence.
        rest = md_text[m.end():]
        gap = re.match(r'\s*', rest).group(0)
        caption_md, sub_md, end = '', '', m.end()
        # the image line's own newline, plus at most three blank lines
        if gap.count('\n') <= 4:
            tail = rest[len(gap):]
            # Step over a spacing directive standing between the image and
            # its caption, and over the blank line that follows it. Without
            # this the directive IS the caption the search finds, and the
            # real one is left behind as ordinary prose.
            spacing = _spacing_only_prefix(tail)
            if spacing:
                lead = re.match(r'\s*', tail[len(spacing):]).group(0)
                if gap.count('\n') + lead.count('\n') <= 4:
                    gap += spacing + lead
                    tail = tail[len(spacing) + len(lead):]
            # Prefer the source's own answer; the bold lead-in is only a
            # fallback for builds with no flat.tex to consult.
            if wants_caption is None:
                take = (number in captioned) if captioned is not None \
                    else tail.startswith('**')
            else:
                take = wants_caption
            # The next paragraph is the next panel of this same float, never a
            # caption. Folding it in cost two of SINQ's thirteen images.
            if _FIG_IMAGE_RE.match(tail):
                take = False
            # A panel of a multi-panel float can carry its own \subcaption,
            # which arrives as a short bold paragraph. It belongs to this
            # panel whether or not this panel also carries the float caption.
            if panel and panel['panels'] > 1 and not _FIG_IMAGE_RE.match(tail):
                first = re.match(r'\*\*[^\n*][^\n]{0,78}?\*\*[ \t]*(?=\n|$)', tail)
                if first:
                    sub_md = first.group(0).strip()
                    consumed = len(gap) + len(sub_md)
                    end = m.end() + consumed
                    tail = tail[len(sub_md):]
                    gap = re.match(r'\s*', tail).group(0)
                    tail = tail[len(gap):]
                    if take and _FIG_IMAGE_RE.match(tail):
                        take = False
            if take and tail.strip():
                stop = re.search(r'\n\s*\n', tail)
                caption_md = (tail[:stop.start()] if stop else tail).strip()
                end = (end if sub_md else m.end()) + len(gap) + len(caption_md)

        if number is None:
            # The float carries no \caption, so the paper never numbered it.
            # Printing "Figure N" here would invent a number the reader cannot
            # find and push every later figure out of step.
            pieces.append(md_text[cursor:m.start()])
            pieces.append(f'\n\n![{alt}]({path})\n\n')
            cursor = m.end()
            count += 1
            continue

        if letter and not caption_md:
            # A panel that does not carry the float's caption says only which
            # panel it is. Numbering every one of them printed
            # `그림 6 (Fig. 6)` four times over what the paper prints once,
            # and a reader counting figures found four where there is one.
            label = f'({letter})'
        else:
            label = f'{fig_label} {number}'
            if fig_label != 'Figure':
                label += f' (Fig. {number})'
            if letter:
                # The caption refers to its panels as (a)/(b)/(c) -- and
                # \subref resolves to those letters -- so the panel carrying
                # the caption still has to say which one it is.
                label += f' ({letter})'
        # The label is bold so the print sheet can pick it out with
        # `figcaption strong`; the caption keeps whatever markdown it had.
        grouped = bool(panel and panel['panels'] > 1)
        if grouped and caption_md:
            # The float's caption belongs under the whole float, not under
            # whichever panel happened to carry it in the source. SINQ's
            # figure 2 explained panel (a) in a caption printed beneath panel
            # (c), a page later; the panel keeps only its own subcaption.
            group_caption = _figure_caption_md(
                f'{fig_label} {number}'
                + (f' (Fig. {number})' if fig_label != 'Figure' else ''),
                caption_md)
            caption = _figure_caption_md(f'({letter})' if letter else label,
                                         sub_md)
        else:
            caption = _figure_caption_md(
                label, ' '.join(p for p in (sub_md, caption_md) if p))

        # No per-image width on a panel: the group's row divides itself
        # between them. An inline `width:33%` here means 33% of the panel's
        # OWN box, which left each panel on a line of its own and spread one
        # float down a whole page.
        attr = ''
        pieces.append(md_text[cursor:m.start()])
        if grouped and not open_group:
            pieces.append('\n\n::: figuregroup\n')
            open_group = True
        pieces.append(f'\n\n![{caption}]({path}){attr}\n\n')
        if open_group and (not grouped or _last_of_float(matches, n, panels)):
            if group_caption:
                pieces.append('\n%s\n' % group_caption)
                group_caption = ''
            pieces.append('\n:::\n\n')
            open_group = False
        cursor = end
        count += 1
    pieces.append(md_text[cursor:])
    return ''.join(pieces), count


# =============================================================================
# Citations and cross-references
# =============================================================================
#
# Two kinds of marker survive the arXiv path and used to print verbatim:
#
#   [@brohan2023rt-2]   pandoc citation syntax. The paper ships a precompiled
#                       main.bbl rather than a .bib, so there is nothing for
#                       --citeproc to read -- but the inlined \bibitem list IS
#                       the numbering, so keys resolve against it exactly.
#   (fig:compare)       what \ref{fig:compare} degrades to. flat.tex still has
#                       every \label in float order, which gives the real
#                       figure and table numbers with no guessing.
#
# Both are resolved on the merged markdown, so an already-translated book can
# be fixed by rebuilding rather than re-translating.

_CITE_RE = re.compile(r'\[(@[^\]]+)\]')
_CITE_KEY_RE = re.compile(r'@([A-Za-z0-9_:\-\.]+)')
# The reference word that sits in front of the number. \ref leaves the one the
# author typed ("See Tab.~\ref{tab:x}"); \cref generates it, so pandoc emits
# nothing and the sentence arrives as "in ( (eq:y))". Absorbing it either way
# is what stops the output reading "See Tab. 표 16".
_XREF_LEAD = (r'(?:\b(?:Figs?|Figures?|Tabs?|Tables?|Secs?|Sections?|Eqs?|Eqn|'
              r'Equations?|App|Appendix|Appendices|Algs?|Algorithms?|Lemmas?|'
              r'Theorems?|Defs?|Definitions?)\.?[ \t\u00a0~]*)?')
# [^()] matters: without it 'increase (Sec.\u00a0(sec:method))' lets the
# leading \( swallow the outer bracket, 'Sec' is taken for the kind and the
# label is captured as '(sec:method' -- which resolves to nothing.
_XREF_BODY = (
    r'\(\s*(fig|figure|tab|table|eq|eqn|equation|sec|section|subsec|'
    r'app|appendix|alg|algorithm|thm|theorem|lem|lemma|def|definition|'
    r'prop|proposition|cor|corollary)[:.]\s*([^()\s]+?)\s*\)')
_XREF_RE = re.compile(_XREF_LEAD + _XREF_BODY, re.IGNORECASE)


def template_affixes(formats, words):
    r"""The literal words a reference template puts around the number.

    Chinese writes a section reference `第3.2节`, so `ref_formats` carries
    `第{number}{label}`. A translator writes that same 第 in front of the
    placeholder and that same 节 after it. Neither was absorbed: 第 is not a
    label word at all, and the closing 节 is only reachable after the
    substitution, where `(?!\w)` can never hold because the next character in
    Chinese is another ideograph. DeeR-VLA printed `第 第3.2节 节在任意...`.

    Only non-ASCII affixes are returned. The equation template is
    `{label} ({number})`, whose literal head is ` (` -- punctuation, not a
    word a translator repeats, and absorbing a bracket would eat the
    reference's own parenthesis.
    """
    lead, trail = [], []
    for slot, template in (formats or {}).items():
        head, _sep, after = template.partition('{number}')
        opener = head.replace('{label}', ' ').strip()
        if opener and not opener.isascii():
            lead.append(opener)
        label = (words or {}).get(slot)
        if label and '{label}' in after and not label.isascii():
            trail.append(label)
        closer = after.replace('{label}', ' ').strip()
        if closer and not closer.isascii():
            trail.append(closer)
    return lead, trail


def _xref_regex(labels, trailing=()):
    """The reference pattern, also absorbing the target language's own words.

    A translator writes "표 (tab:main)에서", not "Tab. (tab:main)". Absorbing
    only the English word left the Korean one standing beside the label this
    emits -- "표 표 12". The lookbehind keeps compounds intact: the word has to
    start where it stands, so "수식 (eq:x)" is not split at "식".

    `trailing` is for a language whose reference form CLOSES with a word.
    Cleaning that up after the substitution means guessing, because 节 also
    opens 节点 and a script with no word boundaries cannot tell a duplicate
    from the next word. Here the placeholder's own `)` bounds it, so there is
    nothing to guess. It is captured rather than dropped: `xref_sub` puts it
    back unless the resolved reference already ends in it.
    """
    words = sorted({w.strip() for w in labels if w and w.strip()},
                   key=len, reverse=True)
    close = sorted({w.strip() for w in trailing if w and w.strip()},
                   key=len, reverse=True)
    tail = ''
    if close:
        tail = (r'(?:[ \t\xa0~]*(%s))?'
                % '|'.join(re.escape(w) for w in close))
    if not words:
        if not tail:
            return _XREF_RE
        return re.compile(_XREF_LEAD + _XREF_BODY + tail, re.IGNORECASE)
    native = (r'(?:(?<!\w)(?:%s)[ \t\u00a0~]*)?'
              % '|'.join(re.escape(w) for w in words))
    return re.compile(_XREF_LEAD + native + _XREF_BODY + tail, re.IGNORECASE)

# reference kind -> (which map holds the number, which label word to use)
_XREF_KINDS = {
    'fig': ('float', 'figure'), 'figure': ('float', 'figure'),
    'tab': ('float', 'table'), 'table': ('float', 'table'),
    'eq': ('label', 'equation'), 'eqn': ('label', 'equation'),
    'equation': ('label', 'equation'),
    'sec': ('label', 'section'), 'section': ('label', 'section'),
    'subsec': ('label', 'section'),
    'app': ('label', 'appendix'), 'appendix': ('label', 'appendix'),
    'alg': ('label', 'algorithm'), 'algorithm': ('label', 'algorithm'),
    'thm': ('label', 'theorem'), 'theorem': ('label', 'theorem'),
    'lem': ('label', 'theorem'), 'lemma': ('label', 'theorem'),
    'def': ('label', 'theorem'), 'definition': ('label', 'theorem'),
    'prop': ('label', 'theorem'), 'proposition': ('label', 'theorem'),
    'cor': ('label', 'theorem'), 'corollary': ('label', 'theorem'),
}

_XREF_WORDS = {'figure': 'Figure', 'table': 'Table', 'equation': 'Equation',
               'section': 'Section', 'appendix': 'Appendix',
               'algorithm': 'Algorithm', 'theorem': 'Theorem'}

# Korean puts the section marker after the number ("4.1절"); everything else
# reads as a prefix in every language here.
_XREF_FORMATS = {'equation': '{label} ({number})'}


def build_bibitem_numbers(md_text):
    r"""{citekey: number} from the inlined \bibitem list, in its own order.

    The optional label is not optional in practice: natbib and plainnat write
    `\bibitem[Adleman 1994]{Adle}`, and requiring the bare form found 0 keys in
    a file holding 75 of them — so all 30 of Shor's citations resolved to
    nothing and `[@Knut]` printed on the page. `_BIBITEM_LABEL_RE` elsewhere in
    this module already accepts the labelled form; this reader had simply never
    been told (K114 is the same shape: learned in one place, not the other).
    """
    keys = re.findall(r'\\bibitem\s*(?:\[[^\]]*\])?\s*\{([^{}]+)\}', md_text)
    return {k.strip(): i + 1 for i, k in enumerate(keys)}


def build_float_numbers(temp_dir):
    """{'fig:compare': 1, 'tab:main result': 2, ...} parsed from flat.tex.

    LaTeX numbers figures and tables in the order their captions appear,
    counting `figure` and `figure*` together, so that is what float_units
    reproduces. A float is not always \\begin{figure}: CafeQ opens with an
    SCfigure from the sidecap package, which LaTeX counts as Figure 1 --
    missing it shifted every later figure down by one, so the caption of
    Figure 2 read "그림 1".

    Returns {} when flat.tex is absent (the calibre backend).
    """
    numbers = {}
    for unit in read_float_units(temp_dir):
        if unit['number'] is None:
            continue
        for label in unit['labels']:
            numbers.setdefault(label, unit['number'])
    return numbers


_SUBENV_RE = re.compile(
    r'\\begin\{(subfigure|subtable)\}(.*?)\\end\{\1\}', re.DOTALL)
_SUBFLOAT_RE = re.compile(r'\\subfloat\s*(?:\[[^\]]*\])?\s*\{')
# `\subfloat[Title]{...}` is lettered; `\subfloat{...}` is not.
_SUBFLOAT_CAPTIONED_RE = re.compile(r'\\subfloat\s*\[[^\]]*\]\s*\{')
_PANEL_CAPTION_RE = re.compile(r'\\caption(?![A-Za-z])')
_SUBREF_RE = re.compile(r'\\subref\s*\{([^{}]+)\}')


def _panel_spans(body):
    """[(start, stop, text)] for each panel in a float, in source order.

    A panel is its own environment. Slicing "from this panel to the next"
    instead would hand the last panel everything after it, including the
    float's own \\caption and \\label.
    """
    spans = []
    for m in _SUBENV_RE.finditer(body):
        spans.append((m.start(), m.end(), m.group(2)))
    for m in _SUBFLOAT_RE.finditer(body):
        close = _balanced_group(body, m.end() - 1)
        if close > 0:
            # Carry the optional argument along so a captioned subfloat can be
            # told from an uncaptioned one downstream.
            spans.append((m.start(), close, body[m.start():close]))
    spans.sort(key=lambda s: s[0])
    return spans


def _panel_is_lettered(scope):
    r"""Does this panel take a letter?

    `subcaption` steps the sub-counter on `\caption`, not on the environment.
    Neural ODE's figure holds four `subfigure`s and the third is a legend with
    no caption, so the paper prints (a) (b) (c) — verified in the source PDF —
    while lettering by position printed `(d)` for the last one. A caption is
    what makes a panel a panel.
    """
    if _PANEL_CAPTION_RE.search(scope):
        return True
    return bool(_SUBFLOAT_CAPTIONED_RE.match(scope.lstrip()))


def _panel_scopes(body):
    """[(start, text)] for each panel in a float, in source order."""
    return [(start, text) for start, _stop, text in _panel_spans(body)]


def build_subfigure_letters(temp_dir):
    """{'fig:corr': 'a', 'fig:adam': 'b', ...} parsed from flat.tex.

    A multi-panel figure letters its panels in source order, and its caption
    refers to them with \\subref. Without this the caption reads
    "(\\subref{fig:corr}) In LLMs ..." and the reader cannot tell which of the
    three plots the sentence is about.
    """
    flat = os.path.join(temp_dir, 'flat.tex')
    if not os.path.exists(flat):
        return {}
    try:
        with open(flat, 'r', encoding='utf-8', errors='replace') as fh:
            tex = strip_tex_comments(fh.read())
    except OSError:
        return {}

    letters = {}
    # The same prefixes `_label_token_re` already accepts. A `wrapfigure` is a
    # float like any other and its panels are lettered like any other's —
    # Neural ODE puts three `subfigure`s inside one, and scanning only for
    # `figure` found no panels at all, so all three `\subref` in its caption
    # printed as LaTeX beside the plots they were pointing at.
    float_re = re.compile(
        r'\\begin\{((?:SC|wrap|sideways|long|floating)?(?:figure|table)\*?)\}'
        r'(.*?)\\end\{\1\}', re.DOTALL)
    for float_match in float_re.finditer(tex):
        index = 0
        for _start, scope in _panel_scopes(float_match.group(2)):
            if not _panel_is_lettered(scope):
                continue          # a legend panel takes no letter, and does
                                  # not advance the one the next panel gets
            for label in re.findall(r'\\label\{([^}]+)\}', scope):
                letters.setdefault(label.strip(), chr(ord('a') + index))
            index += 1
    return letters


_DEFAULT_THEOREM_ENVS = ('theorem', 'lemma', 'definition', 'proposition',
                         'corollary', 'remark', 'assumption')
_COUNTED_STRUCTURAL_ENVS = ('equation', 'align', 'gather', 'multline',
                            'eqnarray', 'algorithm')

# `\newtheorem` carries an optional argument on EITHER side of the title, and
# they mean opposite things: `\newtheorem{lmm}[thrm]{Lemma}` shares the `thrm`
# counter, `\newtheorem{thrm}{Theorem}[section]` scopes it to the section. Only
# the leading one was captured, so `[section]` was not ignored — it could not
# be seen. `read_theorem_environments` said so anyway, and the test that pinned
# the claim passed on a group that was always None.
_NEWTHEOREM_RE = re.compile(
    r'\\newtheorem(\*?)\s*\{([^}]+)\}\s*(?:\[([^\]]*)\])?\s*\{[^}]*\}'
    r'\s*(?:\[([^\]]*)\])?')


def read_theorem_environments(tex):
    """{env: counter_group}, read from the document's own \\newtheorem lines.

    A fixed list of environment names is a guess about the author's taste, and
    Vershynin spends it immediately: he declares `example`, `fact`,
    `observation`, `conjecture`, `remarks` and `definition-notag` too. Five of
    those appear in the body, none were in the list, so the theorem counter
    silently skipped them -- 52 of 59 labels came out with the wrong number and
    225 of 274 body references would have printed one. Nothing errored; the
    numbers were simply, confidently wrong.

    The document already says which environments are numbered and which counter
    each one shares. Read that instead of guessing.

    `\\newtheorem*` takes no number at all. The TRAILING optional argument —
    `\\newtheorem{thrm}{Theorem}[section]` — says the counter is scoped to the
    section; `read_counter_parents` reads it, and `_counter_label` puts the
    section number in front. That used to be waved away as pandoc's business
    (K113), which cost Maynard every one of its 35 theorem references.
    """
    envs = {}
    for m in _NEWTHEOREM_RE.finditer(tex):
        star, name, shared = m.group(1), m.group(2).strip(), m.group(3)
        if star or not name:
            continue
        envs[name] = (shared or name).strip() or name
    return envs


def read_theorem_parents(tex):
    r"""{counter: parent} from the trailing `[within]` of `\newtheorem`.

    Only for a declaration that owns its counter. `\newtheorem{lmm}[thrm]{...}`
    shares `thrm`'s, so the scope belongs to `thrm` and naming `lmm` here would
    reset a counter that does not exist.
    """
    parents = {}
    for m in _NEWTHEOREM_RE.finditer(tex):
        star, name, shared, within = (m.group(1), m.group(2).strip(),
                                      m.group(3), m.group(4))
        if star or not name or shared:
            continue
        within = (within or '').strip()
        if within:
            parents[name] = within
    return parents


_COUNTER_WITHIN_RE = re.compile(
    r'\\(?:counter|number)within\s*\*?\s*\{\s*([A-Za-z@]+)\s*\}'
    r'\s*\{\s*([A-Za-z@]+)\s*\}')
_THE_REDEF_RE = re.compile(
    r'\\def\s*\\the([A-Za-z@]+)\s*\{((?:[^{}]|\{[^{}]*\})*)\}'
    r'|\\renewcommand\s*\*?\s*\{?\s*\\the([A-Za-z@]+)\s*\}?'
    r'\s*\{((?:[^{}]|\{[^{}]*\})*)\}')


def read_counter_parents(tex):
    r"""{counter: parent} for counters printed with another counter in front.

    Shor 1995 prints `(2.1)`, not `(2)`: its preamble resets the equation,
    figure and table counters at each section and redefines `\theequation` to
    carry `\thesection`. Nothing in this pipeline read either signal, so all 48
    equation references and every float reference named a number the paper does
    not print — and the mismatch was nearly accepted as a limit of the
    supported subset.

    It is not one. The number on an equation or a float is stamped by THIS
    pipeline (K46, K62), so it is ours to choose. ~~Only theorem-likes are
    numbered by pandoc and stay out of reach (K113).~~ Nor are they: what
    pandoc wrote is characters in `output.md`, and the build rewrites that line
    anyway (K130). The trailing `[within]` of `\newtheorem` is read here too.

    Nor is this a LaTeX interpreter. One fact is needed per counter — which
    counter, if any, prefixes it — and a document can only say it a few ways.
    Everything else inside `\the...` means "print arabic", which is what we
    already do. Across the whole corpus exactly two papers say anything at all.
    """
    # `[0] or tex` looked like a sensible fallback for a fragment with no
    # \begin{document}, and it also fired when the document begins with one:
    # the empty preamble is falsy, so the whole body got scanned and a counter
    # redefined mid-document was read as a preamble declaration.
    at = tex.find(r'\begin{document}')
    preamble = tex[:at] if at >= 0 else tex
    parents = read_theorem_parents(preamble)
    for m in _COUNTER_WITHIN_RE.finditer(preamble):
        parents[m.group(1)] = m.group(2)
    for m in _THE_REDEF_RE.finditer(preamble):
        name = m.group(1) or m.group(3)
        body = m.group(2) if m.group(1) else m.group(4)
        if not name:
            continue
        ref = re.search(r'\\the([A-Za-z@]+)', body or '')
        # `\def\theequation{\arabic{equation}}` names no parent: it only says
        # how to print, and that is already what happens.
        if ref and ref.group(1) != name:
            parents[name] = ref.group(1)
    return parents


_SETCOUNTER_RE = re.compile(
    r'\\setcounter\s*\{\s*([A-Za-z@]+)\s*\}\s*\{\s*(\d+)\s*\}')


def read_fixed_counter_prefix(tex):
    r"""{counter: value} for a counter set once and never advanced.

    randmat is one chapter of a `book` shipped on its own: it declares
    `\newtheorem{theorem}{Theorem}[chapter]`, writes `\setcounter{chapter}{5}`
    just before the first section, and contains no `\chapter{}` at all. Its
    PDF prints Theorem 5.44, and 258 of its references are dotted while none
    is plain — but with the chapter counter treated as absent the numbers came
    out flat and all 157 disagreed.

    The `5` is in the source, not only in the PDF. Read it, but only when the
    counter is never advanced: a document that really uses `\chapter{}` has a
    prefix that moves, and this must leave that alone rather than pin it.
    """
    fixed = {}
    for m in _SETCOUNTER_RE.finditer(tex):
        name, value = m.group(1), m.group(2)
        if re.search(r'\\%s\b(?!\s*\{\s*\})' % re.escape(name), tex) \
                and re.search(r'\\%s\s*\*?\s*[\[{]' % re.escape(name), tex):
            continue                       # the counter's own command is used
        fixed[name] = value
    return fixed


def _counter_label(counter, n, parents, section_head, fixed=None):
    """`2.1` when the counter is scoped to a section, else `1`.

    `fixed` covers the other way a prefix can be constant: a counter set once
    with `\\setcounter` and never advanced (see `read_fixed_counter_prefix`).
    """
    parent = parents.get(counter)
    if parent == 'section' and section_head:
        return '%s.%d' % (section_head, n)
    if fixed and parent in fixed:
        return '%s.%d' % (fixed[parent], n)
    return str(n)


def _label_token_re(theorem_envs):
    """Compile the scanner for one document's environment vocabulary.

    Longest name first: `definition-notag` must not be read as `definition`
    followed by stray text.
    """
    names = set(_COUNTED_STRUCTURAL_ENVS) | set(theorem_envs)
    alt = '|'.join(re.escape(name) for name in
                   sorted(names, key=lambda n: (-len(n), n)))
    return re.compile(
        r'\\(appendix)(?![a-zA-Z])'
        r'|\\((?:sub)*)section(\*?)\s*\{'
        r'|\\begin\{(' + alt + r')(\*?)\}'
        r'|\\begin\{(?:SC|wrap|sideways|long|floating)?(figure|table)\*?\}'
        r'|\\begin\{(tabular|lstlisting|listing|minted|verbatim'
        r'|thebibliography)\*?\}'
        r'|\\end\{(?:SC|wrap|sideways|long|floating)?(figure|table|tabular'
        r'|lstlisting|listing|minted|verbatim|thebibliography)\*?\}'
        r'|\\label\s*\{([^}]+)\}')


_LABEL_TOKEN_RE = _label_token_re(_DEFAULT_THEOREM_ENVS)

_NUMBERED_MATH_ENVS = ('equation', 'align', 'gather', 'multline', 'eqnarray')


def build_label_index(temp_dir):
    """{'eq:pqe': ('5', 'equation'), 'sec:single': ('4.1', 'section')}.

    The kind matters as much as the number. A label's prefix is the author's
    naming habit, not a fact: CafeQ writes `\\label{lem:norm}` INSIDE an
    `equation`, and reading `lem` as "lemma" printed `정리 3` in a book that
    contains no theorems -- sending the reader to look for something that
    was never there. What the label is attached to is what it names.

    Sections, appendices, equations, algorithms and theorem-likes: everything
    LaTeX numbers except floats, which build_float_numbers already owns.

    Rebuilding the counters is necessary even when the headings show no
    numbers. CafeQ and AlphaQ suppress the number in the heading but their body
    still says "Section 4.1" and "Appendix A.2" -- \\ref returns the counter
    either way.
    """
    flat = os.path.join(temp_dir, 'flat.tex')
    if not os.path.exists(flat):
        return {}
    try:
        with open(flat, 'r', encoding='utf-8', errors='replace') as fh:
            tex = strip_tex_comments(fh.read())
    except OSError:
        return {}

    section = [0, 0, 0]
    in_appendix = False
    counters = {'equation': 0, 'algorithm': 0, 'theorem': 0}
    theorem_envs = read_theorem_environments(tex)
    token_re = _label_token_re(theorem_envs or _DEFAULT_THEOREM_ENVS)
    parents = read_counter_parents(tex)
    fixed = read_fixed_counter_prefix(tex)
    section_head = ''                   # the number a section-scoped counter
                                        # carries in front of its own
    current = None                      # what the next \label would name
    kind = None                         # and what kind of thing that is
    suspended = []                      # section context stacked over floats
    numbers = {}
    for m in token_re.finditer(tex):
        if m.group(1):
            in_appendix = True
            section = [0, 0, 0]
            current, kind = None, None
        elif m.group(2) is not None and m.group(3) is not None:
            if m.group(3) == '*':       # starred sections take no number
                current, kind = None, None
                continue
            depth = m.group(2).count('sub')
            section[depth] += 1
            for deeper in range(depth + 1, 3):
                section[deeper] = 0
            head = (chr(ord('A') + section[0] - 1) if in_appendix
                    else str(section[0]))
            # A `book` shipped as one chapter numbers its sections under it:
            # randmat writes `\setcounter{chapter}{5}` and its own text says
            # "Section 5.4.3", never "Section 4.3". The prefix comes from the
            # source, not from the PDF, and only when no `\chapter{}` moves it.
            if not in_appendix and fixed.get('chapter'):
                head = '%s.%s' % (fixed['chapter'], head)
            current = '.'.join([head] + [str(x) for x in section[1:depth + 1]])
            kind = 'section'
            if depth == 0:
                # A counter scoped to `section` restarts here, and from here on
                # prints this number in front of its own.
                section_head = head
                for counter, parent in parents.items():
                    if parent == 'section':
                        counters[counter] = 0
        elif m.group(4):
            if m.group(5):
                current, kind = None, None
                continue
            env = m.group(4)
            if env in _NUMBERED_MATH_ENVS:
                counters['equation'] += 1
                current = _counter_label('equation', counters['equation'],
                                         parents, section_head, fixed)
                kind = 'equation'
            elif env == 'algorithm':
                counters['algorithm'] += 1
                current, kind = str(counters['algorithm']), 'algorithm'
            else:
                # Environments sharing one counter must share one tally, and
                # an environment declared with its own counter must not touch
                # anyone else's: `definition-notag` restarts at 1 mid-paper,
                # and pandoc prints exactly that.
                group = theorem_envs.get(env, 'theorem')
                counters[group] = counters.get(group, 0) + 1
                # The same call the equation branch makes eight lines up. A
                # counter scoped to the section prints `4.1`, and the reset
                # loop above has already zeroed it at the section head.
                current = _counter_label(group, counters[group],
                                         parents, section_head, fixed)
                kind = 'theorem'
        elif m.group(6) or m.group(7):
            # Entering a float or a verbatim block. A \label inside a float
            # belongs to the float's own counter -- build_float_numbers owns
            # those -- so suspend whatever the last section or equation was.
            suspended.append((current, kind))
            current, kind = None, None
        elif m.group(8):
            # Leaving it. LaTeX scopes a float's \refstepcounter to the float,
            # so a \label placed after \end{figure} still names the enclosing
            # section: AlphaQ's app:hill-derivation is Appendix A.3, and the
            # paper prints exactly that.
            if suspended:
                current, kind = suspended.pop()
        elif m.group(9) and current is not None:
            numbers.setdefault(m.group(9).strip(), (current, kind))
    return numbers


def build_label_numbers(temp_dir):
    """{label: number}, the numbering half of build_label_index."""
    return {key: value[0] for key, value in
            build_label_index(temp_dir).items()}


# A declaration site, not a reference: `**정리 32** (Gaussian).` names the
# theorem being stated. Vershynin labels it `\label{Gaussian}` too, so the key
# alone cannot tell the two apart -- but a declaration always closes an
# emphasis span carrying the number right before the parenthesis, and a
# reference in running prose never does.
_DECL_SITE_TAIL_RE = re.compile(
    r'(?:\*\*|\*)[^*\n]*\d[^*\n]*(?:\*\*|\*)[ \t ]*$')


def _unprefixed_xref_regex(keys, label_words):
    """`(Bai-Yin)` -- a reference whose label carries no kind prefix.

    build_label_index already learned that a label's prefix is the author's
    naming habit rather than a fact. The resolver had not: it only recognised a
    reference by that prefix, so a paper writing `\\label{Bai-Yin}` instead of
    `\\label{thm:bai-yin}` had every one of its cross-references left standing
    as the raw label key. Vershynin has 260 of them, and the reader meets
    `정리 (deviation from 1)` where the paper says Lemma 44.

    Match the keys themselves. They are known exactly, so nothing is guessed.
    """
    if not keys:
        return None
    body = '|'.join(re.escape(k) for k in
                    sorted(keys, key=len, reverse=True))
    lead = ''
    words = sorted({w.strip() for w in label_words if w and w.strip()},
                   key=len, reverse=True)
    if words:
        lead = (r'(?:(?<!\w)(%s)[ \t ~]*)?'
                % '|'.join(re.escape(w) for w in words))
    return re.compile(_XREF_LEAD + lead + r'\(\s*(' + body + r')\s*\)')


def resolve_unprefixed_references(md_text, index, words, formats,
                                  lead_words=None):
    """Resolve `(key)` against the label index. Returns (text, count)."""
    if not index:
        return md_text, 0
    leads = list(words.values()) + list(lead_words or ())
    pattern = _unprefixed_xref_regex(index.keys(), leads)
    if pattern is None:
        return md_text, 0
    hits = [0]

    def sub(m):
        entry = index.get(m.group(2).strip())
        if entry is None:
            return m.group(0)
        number, kind = entry
        lead = m.group(1)
        if lead is None and _DECL_SITE_TAIL_RE.search(md_text[:m.start()]):
            return m.group(0)       # the statement itself, not a pointer to it
        slot = kind if kind in words else 'theorem'
        hits[0] += 1
        if kind == 'theorem' and lead:
            # Keep the word the translator wrote. The index knows the number
            # but not whether this one is a Lemma, a Corollary or a Remark,
            # and "보조정리 28" must not be flattened to "정리 28".
            return '%s %s' % (lead, number)
        template = formats.get(slot, '{label} {number}')
        return template.format(label=words[slot], number=number)

    return pattern.sub(sub, md_text), hits[0]


def resolve_references(md_text, temp_dir, lang_cfg=None):
    """Turn [@key] into [N] and (fig:x) into 'Figure N'.

    Returns (text, stats). Anything that cannot be resolved is left exactly as
    it was, so it stays visible rather than silently becoming a wrong number.
    """
    lang_cfg = lang_cfg or {}
    fig_label = lang_cfg.get('figure_label', 'Figure')
    tab_label = lang_cfg.get('table_label', 'Table')
    stats = {'cites': 0, 'cites_missed': 0, 'xrefs': 0, 'xrefs_missed': 0,
             'subrefs': 0}

    sub_letters = build_subfigure_letters(temp_dir)

    def subref_sub(m):
        letter = sub_letters.get(m.group(1).strip())
        if letter is None:
            return m.group(0)
        stats['subrefs'] += 1
        return letter

    md_text = _SUBREF_RE.sub(subref_sub, md_text)

    bib = build_bibitem_numbers(md_text)

    def cite_sub(m):
        keys = _CITE_KEY_RE.findall(m.group(1))
        # Only rewrite when EVERY key resolves; a half-numbered citation is
        # worse than an untouched one.
        if not keys or any(k not in bib for k in keys):
            stats['cites_missed'] += 1
            return m.group(0)
        stats['cites'] += 1
        return '[' + ', '.join(str(bib[k]) for k in keys) + ']'

    if bib:
        md_text = _CITE_RE.sub(cite_sub, md_text)
    else:
        stats['cites_missed'] = len(_CITE_RE.findall(md_text))

    floats = build_float_numbers(temp_dir)
    index = build_label_index(temp_dir)
    words = {'figure': fig_label, 'table': tab_label}
    for slot, fallback in _XREF_WORDS.items():
        words.setdefault(slot, lang_cfg.get(slot + '_label', fallback))
    formats = dict(_XREF_FORMATS)
    formats.update(lang_cfg.get('ref_formats') or {})

    def xref_sub(m):
        kind = m.group(1).lower()
        source, slot = _XREF_KINDS.get(kind, (None, None))
        if source is None:
            stats['xrefs_missed'] += 1
            return m.group(0)
        rest = m.group(2).strip()
        if source == 'float':
            prefix = 'fig' if slot == 'figure' else 'tab'
            number = floats.get('%s:%s' % (prefix, rest))
        else:
            entry = index.get('%s:%s' % (kind, rest)) or index.get(rest)
            number = entry[0] if entry else None
            # The prefix says what the author called it; the source says what
            # it is. `\cref{lem:norm}` on a label sitting inside an equation
            # printed `정리 3` into a book with no theorems in it.
            #
            # Only where the two are structurally different. An appendix is a
            # section as far as the counter is concerned, so `section` is the
            # generic bucket and not better information: overriding on it
            # turned every `부록 A.10` into `A.10절`.
            if entry and entry[1] in ('equation', 'algorithm') \
                    and entry[1] in words and entry[1] != slot:
                slot = entry[1]
        if number is None:
            stats['xrefs_missed'] += 1
            return m.group(0)
        stats['xrefs'] += 1
        template = formats.get(slot, '{label} {number}')
        out = template.format(label=words[slot], number=number)
        # The translator's own closing word, if the pattern took one. Drop it
        # only where the reference this emits already ends in it; otherwise it
        # belonged to the sentence and goes back untouched.
        closer = m.group(3) if m.re.groups >= 3 else None
        if closer and not out.endswith(closer):
            out += closer
        return out

    # Whatever the template puts around the number is what a translator writes
    # around the placeholder. Korean's closing label stays with the
    # particle-aware pass below, which has carried it across five books.
    lead_affix, trail_affix = template_affixes(formats, words)
    if lang_cfg.get('particle_agreement') is True:
        trail_affix = []
    xref_re = _xref_regex(list(words.values()) + lead_affix, trail_affix)
    if floats or index:
        md_text = xref_re.sub(xref_sub, md_text)
    else:
        stats['xrefs_missed'] = len(xref_re.findall(md_text))

    md_text, keyed = resolve_unprefixed_references(
        md_text, index, words, formats,
        lead_words=lang_cfg.get('theorem_words')
        or _DEFAULT_LANG_CONFIG.get('theorem_words'))
    stats['xrefs'] += keyed

    md_text, bare = resolve_bare_float_labels(md_text, temp_dir, words,
                                              formats)
    stats['xrefs'] += bare
    md_text, stats['doubled'] = drop_doubled_labels(md_text, words, formats,
                                                    lang_cfg)
    md_text, stats['particles'] = fix_particles(md_text, lang_cfg)
    return md_text, stats


# Which Korean particle follows depends on how the PRECEDING SYLLABLE is
# pronounced, and the translator never saw the number: it wrote a particle
# after "(fig:dual_scale)" and this pass then substituted "그림 1" in front of
# it. Twelve of those shipped -- "그림 1를", "표 9은" -- and a Korean reader
# stops at every one, because 1 is read 일 and ends in a consonant.
_DIGIT_HAS_CODA = {'0': True,    # 영
                   '1': True,    # 일
                   '2': False,   # 이
                   '3': True,    # 삼
                   '4': False,   # 사
                   '5': False,   # 오
                   '6': True,    # 육
                   '7': True,    # 칠
                   '8': True,    # 팔
                   '9': False}   # 구
# 으로/로 is the one pair that does not follow the coda rule: a final ㄹ takes
# 로, like 물로 and 서울로. 일, 칠 and 팔 all end in ㄹ, so "8로" is correct
# and "8으로" -- which this pass produced on its first outing -- is not.
_DIGIT_ENDS_IN_RIEUL = {'1', '7', '8'}
# (after a consonant, after a vowel)
_PARTICLE_PAIRS = (('을', '를'), ('은', '는'), ('과', '와'), ('이', '가'),
                   ('으로', '로'))
_NUMBERED_REF_RE = re.compile(
    r'(\d+(?:\.\d+)*)\s*(으로|을|를|은|는|과|와|이|가|로)(?![가-힣])')


# A particle can stand between the doubled label and the next word, so the
# second "절" is not always followed by whitespace. Longest forms first: the
# alternation is tried in order and "에" would otherwise win over "에서는".
_KO_PARTICLES = ('으로', '에서는', '에서', '에는', '에', '의', '을', '를', '은',
                 '는', '과', '와', '이', '가', '로', '도', '만', '부터', '까지')


def resolve_bare_float_labels(md_text, temp_dir, words, formats):
    r"""`(table-mixtral)` -> `표 9`. Returns (text, count).

    A cross-reference is recognised by its prefix -- `(tab:x)`, `(fig:y)`.
    AlphaQ labels one of its tables `\label{table-mixtral}`, with no prefix
    at all, so nothing matched it: the raw label printed to the reader where
    `표 9` belongs, and the pointer to table 9 went with it.

    Two guards, because these pages are full of parenthesised English. Only
    a label attached to a float is eligible — the same paper has a section
    labelled `HT-SR`, which is also how it introduces the acronym. And only
    where a space precedes the bracket, which a gloss never has:
    `Self-Regularization(HT-SR)` has none, `그리고 (table-mixtral)에서` does.
    """
    known = {}
    for unit in read_float_units(temp_dir):
        if unit.get('number') is None:
            continue
        slot = 'figure' if str(unit.get('kind')).startswith('figure') \
            else 'table'
        for label in unit.get('labels') or ():
            known.setdefault(label, (slot, unit['number']))
    if not known:
        return md_text, 0
    pattern = re.compile(r'(?<=\s)\((%s)\)' % '|'.join(
        sorted((re.escape(k) for k in known), key=len, reverse=True)))

    def sub(m):
        slot, number = known[m.group(1)]
        template = formats.get(slot, '{label} {number}')
        return template.format(label=words[slot], number=number)

    return pattern.subn(sub, md_text)


def drop_abbreviated_label(md_text, label):
    r"""`l'éq. Équation (3)` -> `l'Équation (3)`. Returns (text, n).

    The doubling `drop_doubled_labels` already knew about is the SAME word
    twice, which is what happens when the translator writes the label the
    resolver is about to write. A translator who abbreviates instead leaves
    two forms that do not match each other, so nothing fired: French, German
    and Spanish shipped nine of `l'éq. Équation (3)`, `Gl. Gleichung (3)` and
    `la Ec. Ecuación (5)` between them.

    An abbreviation is a prefix of the word it abbreviates, so it is derived
    rather than listed, and a language nobody has run yet is covered. The
    prefix test is also what keeps an ordinary abbreviation in front of a
    reference safe: `cf. Ecuación (3)` and `vs. Tabla 1` are not prefixes of
    what follows them and are left alone.
    """
    lowered = label.lower()
    pattern = re.compile(
        r'(?<![^\W\d_])([^\W\d_]{2,12})\.[ \t]*(?=' + re.escape(label)
        + r'\s*\(?\s*[A-Za-z]?\d)', re.UNICODE)
    count = [0]

    def strip(match):
        if lowered.startswith(match.group(1).lower()):
            count[0] += 1
            return ''
        return match.group(0)

    return pattern.sub(strip, md_text), count[0]


def drop_doubled_labels(md_text, words, formats, lang_cfg=None):
    r"""`4.1절 절과` -> `4.1절과`. Returns (text, n).

    Same wound as the particle fix below. The translator wrote its own '절'
    after the placeholder, never having seen that this pass would substitute
    a reference which already ends in one, and CafeQ shipped "4.1절 절과 4.2
    절 절에서는".

    A prefix label doubles too, and the note that once stood here said it
    could not. The source writes `Figure (Figure_teaser)`, the resolver
    replaces the parenthesised key alone, and the label word the translator
    put in front of it survives: VLA-Adapter shipped `그림 그림 1` and `표 표
    2` twenty times over. The old assumption held only while translators left
    that word in English, where the two forms did not match each other.
    """
    lang_cfg = lang_cfg or {}
    total = 0
    tail = r'(?!\w)'
    if lang_cfg.get('particle_agreement') is True:
        tail = r'(?=(?:%s)?(?!\w))' % '|'.join(_KO_PARTICLES)
    for slot, label in words.items():
        template = formats.get(slot, '{label} {number}')
        head, _sep, after = template.partition('{number}')
        esc = re.escape(label)
        if '{label}' in after:
            # Suffix style: `4.1절 절과` -> `4.1절과`.
            pattern = re.compile(r'(\d[\d.]*' + esc + r')\s*' + esc + tail)
        elif '{label}' in head:
            # Prefix style: `그림 그림 1` -> `그림 1`. Anchored on the number,
            # so a sentence that merely repeats the word is left alone.
            #
            # The number can carry an appendix letter. Requiring a bare digit
            # was the first attempt and it collapsed every body reference
            # while missing every appendix one: sixteen `그림 그림 A1` and
            # `표 표 C1` reached the finished book while the check that had
            # just been written reported zero.
            # The number can also be parenthesised. An equation reference is
            # written `式 (3)`, not `式 3`, so this rule could never see the
            # doubling in one: Chinese printed `那么式 式 (3)` and `针对式 式
            # (5)` with every count reporting zero.
            #
            # Those two are only reachable here, and not at the absorbing
            # stage, because Chinese leaves no space between words: the
            # translator's own 式 sits inside 那么式, and `_xref_regex`
            # deliberately refuses to take a label out of the middle of a
            # word -- the same guard that stops 수식 being split at 식.
            #
            # The residual risk is the mirror of that guard: a Chinese word
            # ending in 式 (模式, 方式, 形式) standing immediately before a
            # reference would lose the label rather than the duplicate. The
            # two shapes are identical in a script with no word boundaries,
            # so this cannot tell them apart; it chooses the visible defect
            # over the invisible one, since a bare `(3)` still reads as an
            # equation reference and `式 式 (3)` reads as nothing.
            pattern = re.compile(esc + r'\s+(' + esc
                                 + r'\s*[（(]?\s*[A-Za-z]?\d)')
            md_text, n = pattern.subn(r'\1', md_text)
            total += n
            md_text, n = drop_abbreviated_label(md_text, label)
            total += n
            continue
        else:
            continue
        md_text, n = pattern.subn(r'\1', md_text)
        total += n
    return md_text, total


def fix_particles(md_text, lang_cfg=None):
    """Make the particle after a resolved number agree with it. (text, n)."""
    if (lang_cfg or {}).get('particle_agreement') is not True:
        return md_text, 0
    count = [0]

    def swap(m):
        digit = m.group(1).rstrip('.').split('.')[-1][-1]
        coda = _DIGIT_HAS_CODA.get(digit)
        if coda is None:
            return m.group(0)
        for after_c, after_v in _PARTICLE_PAIRS:
            if m.group(2) in (after_c, after_v):
                takes_consonant_form = coda
                if after_c == '으로' and digit in _DIGIT_ENDS_IN_RIEUL:
                    takes_consonant_form = False
                want = after_c if takes_consonant_form else after_v
                if want != m.group(2):
                    count[0] += 1
                return m.group(1) + want
        return m.group(0)

    return _NUMBERED_REF_RE.sub(swap, md_text), count[0]


# =============================================================================
# Displayed equation numbers
# =============================================================================
#
# The cross-reference pass resolves "(eq:ilp)" to the number the paper prints,
# so the body says "식 (5)" -- and the equation itself carried no number, which
# left the reader nothing to match it against.
#
# Which display blocks are numbered is not a guess: LaTeX numbers a math
# environment unless it is starred, and `\nonumber` removes a row's number
# inside align. Counted that way, all three papers agree exactly with the
# `(N)` markers printed in their own PDFs (SINQ 7, CafeQ 5, AlphaQ 27).

_DISPLAY_MATH_BLOCK_RE = re.compile(r'\$\$(.*?)\$\$', re.DOTALL)
_MATH_ENV_OPEN_RE = re.compile(
    r'\\begin\{(equation|align|gather|multline|eqnarray|flalign)(\*?)\}')
_ROW_BREAK = '\\\\'


def _numbers_for_block(body):
    """How many numbers LaTeX would print for this display block."""
    total = 0
    for m in _MATH_ENV_OPEN_RE.finditer(body):
        env, star = m.group(1), m.group(2)
        if star:
            continue
        if env in ('align', 'eqnarray', 'flalign'):
            rows = body.count(_ROW_BREAK) + 1
            total += max(0, rows - len(re.findall(r'\\(?:nonumber|notag)', body)))
        else:
            total += 1
    return total


def flat_equation_numbers(temp_dir):
    r"""The number strings the paper issues, in order, or None.

    Only for a paper whose equation counter is scoped to something — Shor 1995
    prints `(2.1)`. Derived from flat.tex because that is where the section
    structure is: the merged markdown has headings but no reliable mapping back
    to the source's own section numbering, and re-deriving it there would be a
    second, disagreeing implementation of the same thing.

    Returns None when the paper numbers flat, so nothing changes for it.
    """
    flat = os.path.join(temp_dir or '', 'flat.tex')
    if not temp_dir or not os.path.isfile(flat):
        return None
    try:
        with open(flat, 'r', encoding='utf-8', errors='replace') as fh:
            tex = strip_tex_comments(fh.read())
    except OSError:
        return None
    if read_counter_parents(tex).get('equation') != 'section':
        return None

    numbers, head, count = [], '', 0
    token = re.compile(r'\\(?:sub)*section(\*?)\s*\{'
                       r'|\\begin\{(' + '|'.join(_NUMBERED_MATH_ENVS) +
                       r')\}')
    depth0 = 0
    for m in token.finditer(tex):
        if m.group(2):
            count += 1
            numbers.append('%s.%d' % (head, count) if head else str(count))
        elif not m.group(1) and m.group(0).count('sub') == 0:
            depth0 += 1
            head, count = str(depth0), 0
    return numbers or None


def equation_numbers(md_text, strings=None):
    """[(start, end, number_or_None)] for every `$$...$$` block, in order.

    None means the block is unnumbered -- a starred environment, or plain
    display math the source never numbered. A block that would take more than
    one number (an align whose rows are each numbered) is left unlabelled but
    still advances the counter, so every later equation keeps the number the
    original gives it.

    `strings` overrides the flat 1..N counting with the numbers the source
    actually prints. It is ignored unless it has exactly one entry per number
    this text issues: a length mismatch means the two views disagree about how
    many numbers exist, and guessing which is right would misnumber the whole
    book silently.
    """
    out, counter = [], 0
    blocks = []
    for m in _DISPLAY_MATH_BLOCK_RE.finditer(md_text):
        wanted = _numbers_for_block(m.group(1))
        blocks.append((m.start(), m.end(), wanted))
        counter += wanted
    if strings is not None and len(strings) != counter:
        strings = None

    counter = 0
    for start, end, wanted in blocks:
        if wanted == 1:
            counter += 1
            label = (strings[counter - 1] if strings else counter)
            out.append((start, end, label))
        else:
            counter += wanted
            out.append((start, end, None))
    return out


def tag_equations_for_markdown(md_text, temp_dir=None):
    """Put `\\qquad(N)` inside each numbered formula. Returns (text, count).

    For the DOCX path only: pandoc builds book.docx straight from the markdown
    and never sees the HTML the other formats are styled through.
    """
    marks = [m for m in equation_numbers(md_text, flat_equation_numbers(
        temp_dir)) if m[2] is not None]
    if not marks:
        return md_text, 0
    pieces, cursor = [], 0
    for start, end, number in marks:
        body = md_text[start + 2:end - 2]
        closing = re.search(r'(\s*\\end\{[a-zA-Z*]+\}\s*)+$', body)
        # `%s`, not `%d`: a section-scoped number is `2.1`.
        tag = '\\qquad(%s)' % number
        if closing:
            body = body[:closing.start()] + tag + body[closing.start():]
        else:
            body = body + tag
        pieces.append(md_text[cursor:start])
        pieces.append('$$' + body + '$$')
        cursor = end
    pieces.append(md_text[cursor:])
    return ''.join(pieces), len(marks)


_BLOCK_MATH_TAG_RE = re.compile(r'<math\b(?=[^>]*\bdisplay="block")')


def tag_equations_in_html(html, md_text, temp_dir=None):
    """Mark each numbered <math display="block"> with its number.

    pandoc emits one block-level <math> per `$$` block, in document order, so
    the mapping is positional. The stylesheet turns the attribute into a
    flush-right label; nothing is added to the text itself, so a copy-paste of
    the formula stays clean.
    """
    numbers = [n for _s, _e, n in
               equation_numbers(md_text, flat_equation_numbers(temp_dir))]
    if not any(n is not None for n in numbers):
        return html, 0

    pieces, cursor, index, tagged = [], 0, 0, 0
    for m in _BLOCK_MATH_TAG_RE.finditer(html):
        number = numbers[index] if index < len(numbers) else None
        index += 1
        pieces.append(html[cursor:m.end()])
        if number is not None:
            pieces.append(' data-eqno="(%s)"' % number)
            tagged += 1
        cursor = m.end()
    pieces.append(html[cursor:])
    return ''.join(pieces), tagged


# =============================================================================
# Bibliography
# =============================================================================

_THEBIB_RE = re.compile(
    r'\\begin\{thebibliography\}.*?\\end\{thebibliography\}', re.DOTALL)
# citeproc renders "Surname, First, and Other. 2024. “Title.” *Venue* 1: 2-3."
# strip_pandoc_divs has already removed the ::: {#refs} wrapper by this point,
# so the shape of the paragraph is all there is to go on.
_CITEPROC_ENTRY_RE = re.compile(
    r'^[A-Z][^\n]{2,60}?\.\s+(?:\d{4}[a-z]?|n\.d\.)\.\s', re.MULTILINE)


def source_has_bib_files(temp_dir):
    """Did the arXiv source ship a .bib? Then citeproc rendered the references.

    arxiv_backend picks citeproc over inlining the .bbl on exactly this test,
    so asking the same question here says whether a reference list already
    exists somewhere in the document.
    """
    root = os.path.join(temp_dir or '', 'arxiv_src')
    if not os.path.isdir(root):
        return False
    for dirpath, _dirs, files in os.walk(root):
        for name in files:
            if name.lower().endswith('.bib'):
                return True
    return False


_LIST_EDGE_HEADING_RE = re.compile(r'(?m)^#{1,6} ')
_LIST_EDGE_PROSE_RE = re.compile(r'[가-힣]{4,}')


def _citeproc_block_start(md_text):
    """Index where the trailing run of citeproc reference paragraphs begins.

    The walk back stops at the EDGE of the list -- a heading, or a paragraph
    of body prose -- and not at an entry that happens to carry a blank line
    inside it. Counting blank lines meant one long preprint entry truncated
    the run: of AlphaQ's 52 references the walk kept one, so the `참고문헌`
    heading was inserted between the last two entries, six pages after the
    list actually began, and the contents page pointed the reader there.
    """
    matches = list(_CITEPROC_ENTRY_RE.finditer(md_text))
    if len(matches) < 5:
        return -1
    start = matches[-1].start()
    for m in reversed(matches[:-1]):
        between = md_text[m.start():start]
        if _LIST_EDGE_HEADING_RE.search(between):
            break
        if len(_LIST_EDGE_PROSE_RE.findall(between)) > 2:
            break
        start = m.start()
    # do not swallow a heading that sits directly above the first entry
    head = md_text.rfind('\n#', 0, start)
    if head != -1 and md_text[head:start].count('\n\n') <= 1:
        return -1
    return start


_THEBIB_ENTRY_RE = re.compile(r'\\bibitem\s*(?:\[[^\]]*\])?\s*\{[^{}]*\}')
# A command separated from its argument by a LINE BREAK. A space there is fine;
# a newline is not, and pandoc fails the whole fragment over it.
_BIB_ARG_GAP_RE = re.compile(r'(\\[A-Za-z]+)[ \t]*\r?\n[ \t]*(?=\{)')


def expand_thebibliography(md_text, label, pandoc=None):
    r"""Render a raw `thebibliography` into text. Returns (md_text, entries).

    Keeping the environment as it was meant handing pandoc raw LaTeX on the
    HTML path, and pandoc drops that without a word -- the same silence that
    ate the algorithm floats. CafeQ's 61 references reached output.md and
    then stopped: the book carries 19 in-text citations and no list for them
    to point at, and no check counted a reference.
    """
    m = _THEBIB_RE.search(md_text)
    if not m:
        return md_text, 0
    inner = re.sub(r'\\begin\{thebibliography\}\s*(?:\{[^{}]*\})?', '',
                   m.group(0))
    inner = re.sub(r'\\end\{thebibliography\}', '', inner)
    first = _THEBIB_ENTRY_RE.search(inner)
    if not first:
        return md_text, 0
    # Everything before the first \bibitem is the .bbl's own preamble.
    entries = [e.strip() for e
               in _THEBIB_ENTRY_RE.split(inner[first.start():]) if e.strip()]
    if not entries:
        return md_text, 0
    # `\newblock` only asks for space between the parts of an entry, but
    # pandoc has no reader for it and takes the group that follows as its
    # argument. A title written `{FrameQuant}: Flexible low-bit ...` -- the
    # braces protect the capitals -- then printed as `: Flexible low-bit`,
    # with the name of the method missing.
    entries = [re.sub(r'\\newblock\b\s*', ' ', e).strip() for e in entries]
    # A `.bbl` escapes a literal punctuation mark by wrapping it in math:
    # `QuIP$\#$`. texmath has nothing to do with `\#`, so the dollars reach
    # the page and the reader sees the markup instead of the character.
    entries = [re.sub(r'\$\\([#%&_${}])\$', r'\1', e) for e in entries]
    # And the same character after the conversion has escaped the raw block:
    # `QuIP$\#$` arrives as `Qu{IP}\${\textbackslash}\#\$`, four pieces of
    # markup standing in for one `#`, and all four print.
    entries = [re.sub(r'\\\$\{\\textbackslash\}\\([#%&_$])\\\$', r'\1', e)
               for e in entries]
    # A `.bbl` wraps its lines, so a command and its argument can end up on
    # separate lines: `\href\n  {url} {text}`. A SPACE there converts; a
    # NEWLINE does not, and pandoc fails the fragment it is in. Two of BERT's
    # 56 entries were written that way.
    entries = [_BIB_ARG_GAP_RE.sub(r'\1', e) for e in entries]

    rendered = []
    pandoc = pandoc or resolve_pandoc()
    if pandoc:
        with tempfile.TemporaryDirectory(prefix='tb-bib-') as work:
            out = _latex_fragment_to_markdown('\n\n'.join(entries), pandoc,
                                              work, 'bib.tex')
            if out:
                rendered = [p.strip() for p in out.split('\n\n') if p.strip()]
            if len(rendered) != len(entries):
                # One entry pandoc cannot read fails the WHOLE fragment, and
                # the fallback below then prints all 56 as raw LaTeX -- the
                # reader loses a formatted reference list over one bad line.
                # Convert them one at a time so a bad entry costs only itself.
                rendered = []
                for i, entry in enumerate(entries):
                    one = _latex_fragment_to_markdown(entry, pandoc, work,
                                                      'bib%04d.tex' % i)
                    rendered.append(' '.join((one or entry).split()))
    if not rendered:                       # no pandoc: the text, unformatted
        rendered = [' '.join(e.split()) for e in entries]

    # Number them. `build_bibitem_numbers` numbers the in-text citations 1..N
    # from this same `\bibitem` list in this same order, and the list itself
    # was rendered as bare paragraphs with no labels at all, so every book
    # this pipeline has produced carries citations reading [1] to [9] over a
    # list of nine unlabelled paragraphs. The reader cannot resolve any of
    # them. It reached the English pass-through edition too, which is what
    # says it was never a translation defect; six reading passes found it and
    # no check did, because nothing had ever compared the two.
    #
    # `[n]` and not `n.`: the citations print `[1]` and a multi-key one
    # prints `[1, 3]`, so the list has to answer in the same notation.
    rendered = ['[%d] %s' % (i + 1, text) for i, text in enumerate(rendered)]

    return (md_text[:m.start()] + '\n\n# %s\n\n' % label
            + '\n\n'.join(rendered) + '\n\n' + md_text[m.end():],
            len(entries))


def resolve_bibliography(md_text, temp_dir, lang_cfg=None):
    """One reference list, with a heading. Returns (text, stats)."""
    lang_cfg = lang_cfg or {}
    label = lang_cfg.get('references_label', 'References')
    stats = {'dropped_duplicate': 0, 'heading_added': 0, 'inlined_rendered': 0}

    has_citeproc = _citeproc_block_start(md_text) != -1
    if _THEBIB_RE.search(md_text):
        if source_has_bib_files(temp_dir) and has_citeproc:
            # The source inlined its own .bbl and also shipped the .bib that
            # citeproc read. Both lists are in the document; keep the rendered
            # one, which carries no LaTeX and no \providecommand preamble.
            md_text, n = _THEBIB_RE.subn('', md_text)
            stats['dropped_duplicate'] = n
        else:
            md_text, stats['inlined_rendered'] = expand_thebibliography(
                md_text, label)

    start = _citeproc_block_start(md_text)
    if start != -1:
        before = md_text[:start].rstrip()
        heading = '\n\n# %s\n\n' % label
        md_text = before + heading + md_text[start:]
        stats['heading_added'] = 1
    return md_text, stats


# =============================================================================
# Raw LaTeX tables
# =============================================================================
#
# The arXiv backend keeps the paper's tables as raw LaTeX
# (`\resizebox{..}{..}{\begin{tabular}...\end{tabular}}`, often inside a
# `\begin{table*}` float) rather than converting them to markdown. pandoc's
# markdown reader parses that as a raw LaTeX block, and a raw block only
# survives into its OWN output format -- so on the HTML path every one of those
# tables was silently DROPPED. Not rendered as literal text, which would at
# least be visible: simply gone, taking the paper's results section with it.
#
# pandoc CAN read a bare tabular when told the input is LaTeX, so each one is
# converted on its own and spliced back as raw HTML (`raw_html` is in
# PANDOC_FROM, so it passes straight through).
#
# The float wrapper has to be consumed too. pandoc drops a `\begin{table*}`
# float wholesale -- converting one yields an empty document -- and worse, if
# the `\begin{table*}` marker is left in place it starts a raw LaTeX block that
# swallows the HTML table injected inside it. So the whole float is replaced,
# and its `\caption{...}` is rendered separately into a real <caption>.

_TABULAR_BEGIN = r'\begin{tabular}'
_TABULAR_END = r'\end{tabular}'
# `tabular*` is the width-setting variant, and searching for the plain spelling
# misses it completely — the string `\begin{tabular}` is not a substring of
# `\begin{tabular*}`. BERT's GLUE results table is written that way: the float
# was found by no scan, converted by nothing, and reached the page as nothing.
_TABULAR_BEGIN_RE = re.compile(r'\\begin\{tabular\*?\}')
_TABULAR_END_RE = re.compile(r'\\end\{tabular\*?\}')
# Wrappers that take the tabular as their last braced argument.
_TABLE_WRAPPERS = (r'\resizebox', r'\scalebox', r'\colorbox', r'\fbox', r'\makebox')
# Float environments a tabular is commonly parked in.
_FLOAT_ENVS = ('table*', 'table', 'figure*', 'figure')


def _balanced_group(text, open_at):
    """Index just past the `{...}` group that starts at open_at, or -1."""
    depth = 0
    for i in range(open_at, len(text)):
        if text[i] == '{':
            depth += 1
        elif text[i] == '}':
            depth -= 1
            if depth == 0:
                return i + 1
    return -1


def _widen_to_wrapper(text, start, stop):
    """Widen a tabular span over a `\\resizebox{..}{..}{ ... }` wrapper.

    Counts the braces left open between the wrapper command and the tabular,
    then consumes that many closing braces after it.
    """
    # A wrapper the author commented out is not wrapping anything. The
    # Transformer's parsing table carries `%\resizebox{1.0}{` on the line above
    # its tabular: one unbalanced brace, on a line TeX never reads. Counted, it
    # consumed a closing brace belonging to the float; `_widen_to_float` then
    # found no float, the table reached pandoc as raw LaTeX, and pandoc dropped
    # it — the whole constituency-parsing table, gone without a word.
    head = -1
    for cmd in _TABLE_WRAPPERS:
        found = start
        while found > 0:
            found = text.rfind(cmd, max(0, start - 400), found)
            if found < 0:
                break
            if not _in_latex_comment(text, found):
                break
        if found > head:
            head = found
    if head < 0:
        return start, stop
    between = _COMMENT_LINE_RE.sub('', text[head:start])
    depth = between.count('{') - between.count('}')
    if depth <= 0:
        return start, stop
    tail = stop
    while depth > 0 and tail < len(text):
        if text[tail] == '}':
            depth -= 1
        elif text[tail] == '{':
            depth += 1
        tail += 1
    return head, tail


def _widen_to_float(text, start, stop):
    r"""Widen to an enclosing float environment. Returns (start, stop, env).

    The `\begin` found behind the tabular only encloses it if it has not
    already closed. Without that test the search walks back into the
    PREVIOUS float, then forward to the NEXT float's `\end`, and returns a
    span covering everything in between -- 22,000 characters in SINQ's case.
    The expander replaces that span with one rendered table and advances its
    cursor past it, so the prose inside is skipped: 316 Korean words in SINQ
    and 194 in AlphaQ, with every table still counted and every check quiet.
    """
    for env in _FLOAT_ENVS:
        opening, closing = '\\begin{%s}' % env, '\\end{%s}' % env
        begin = text.rfind(opening, max(0, start - 4000), start)
        if begin < 0:
            continue
        if text.find(closing, begin, start) >= 0:
            continue                     # that float closed before this table
        end = text.find(closing, stop)
        if end < 0:
            continue
        return begin, end + len(closing), env
    return start, stop, None


def _extract_caption(float_text, tabular_at=None):
    """The LaTeX inside the `\\caption{...}` that belongs to this tabular.

    Taking the float's FIRST caption is wrong whenever a float holds more
    than one: AlphaQ puts two minipages, each with its own caption and its own
    tabular, inside one `table*`. Both rendered tables carried the first
    caption and the second one's text -- "Component ablation on OLMoE-1B-7B"
    -- was nowhere in the book.

    A caption precedes its tabular inside the minipage, so the nearest one
    ABOVE wins; a float that captions below still resolves by falling forward.
    """
    starts = [m.start() for m in _CAPTION_CMD_RE.finditer(float_text)
              if not _in_latex_comment(float_text, m.start())]
    if not starts:
        return None
    if tabular_at is None:
        at = starts[0]
    else:
        above = [s for s in starts if s < tabular_at]
        at = above[-1] if above else starts[0]
    brace = float_text.find('{', at)
    if brace < 0:
        return None
    end = _balanced_group(float_text, brace)
    return float_text[brace + 1:end - 1] if end > 0 else None


_TABLENOTES_RE = re.compile(
    r'\\begin\{tablenotes\}(?:\s*\[[^\]]*\])?(.*?)\\end\{tablenotes\}',
    re.DOTALL)
_NOTE_ITEM_RE = re.compile(r'\\item\s*(?:\[([^\]]*)\])?\s*')


def extract_table_notes(float_text):
    r"""The `threeparttable` note under a table, as LaTeX pandoc can read.

    pandoc has no reader for `tablenotes` and drops the environment whole.
    SINQ has four, and each one defines the dagger its rows carry: with the
    note gone, four tables printed a marker that nothing on the page
    explained. Nothing counted it either -- the rows were complete, the
    numbers were complete, and the sentence saying which baselines were
    re-run rather than quoted was not there at all.
    """
    m = _TABLENOTES_RE.search(float_text)
    if not m:
        return None
    parts = []
    for chunk in re.split(r'(?=\\item\b)', m.group(1)):
        chunk = chunk.strip()
        item = _NOTE_ITEM_RE.match(chunk) if chunk else None
        if not item:
            continue
        marker, body = (item.group(1) or '').strip(), chunk[item.end():].strip()
        if not body:
            continue
        parts.append('\\textsuperscript{%s} %s' % (marker, body)
                     if marker else body)
    return '\n\n'.join(parts) or None


def _matching_tabular_end(text, start):
    r"""Index just past the `\end{tabular}` that closes the one at `start`.

    A cell holding a multi-line header is written as a `tabular` of its own --
    `\begin{tabular}[c]{@{}c@{}}Only RGB\\ Input\end{tabular}` -- and taking
    the FIRST `\end{tabular}` cuts the outer table off inside that cell.
    Three of DeeR-VLA's tables came out as fragments pandoc could not read,
    a fourth table went missing, and the paper's eleven tables were counted
    as fourteen.
    """
    depth, i = 0, start
    while i < len(text):
        open_m = _TABULAR_BEGIN_RE.search(text, i)
        close_m = _TABULAR_END_RE.search(text, i)
        if close_m is None:
            return -1
        if open_m is not None and open_m.start() < close_m.start():
            depth += 1
            i = open_m.end()
            continue
        depth -= 1
        i = close_m.end()
        if depth == 0:
            return i
    return -1


def find_raw_latex_tables(text):
    """Locate every raw LaTeX tabular, with its wrapper, float and caption."""
    out, i = [], 0
    while True:
        open_m = _TABULAR_BEGIN_RE.search(text, i)
        if open_m is None:
            break
        a = open_m.start()
        b = _matching_tabular_end(text, a)
        if b < 0:
            break
        i = b
        wa, wb = _widen_to_wrapper(text, a, b)
        fa, fb, env = _widen_to_float(text, wa, wb)
        out.append({
            'bare': text[a:b],
            'start': fa,
            'stop': fb,
            'float': env,
            'caption': _extract_caption(text[fa:fb], a - fa) if env else None,
            'notes': extract_table_notes(text[fa:fb]),
        })
    # One float can hold two tabulars, and the note under it belongs to the
    # float, not to each of them. Attached to both it prints twice.
    seen = {}
    for entry in out:
        seen[(entry['start'], entry['stop'])] = entry
    for entry in out:
        if seen.get((entry['start'], entry['stop'])) is not entry:
            entry['notes'] = None
    return out


def count_raw_latex_tables(md_text):
    """How many raw LaTeX tables the merged markdown carries."""
    return len(find_raw_latex_tables(md_text))


# A pandoc table caption is a line of its own opening with `: `.
_MD_TABLE_CAPTION_RE = re.compile(r'(?m)^([ \t]*:[ \t]+)(?=\S)')
# So is a pandoc DEFINITION LIST item, which is why the shape alone cannot
# tell them apart. A caption abuts its table; a definition abuts its term.
_MD_TABLE_ROW_RE = re.compile(r'^\s*(?:\|.*\||\+[-=+:]{2,}\+)\s*$')


def markdown_table_captions(md_text):
    r"""[(line start, offset just past the `: `)] for real table captions.

    Matching `^: ` alone counts every definition list item as a caption.
    VLA-Adapter carries fourteen of them, holding its Question, Key Finding
    and Conclusion items, and not one markdown table. Ten of its fifteen
    table numbers landed on that prose and ten real tables were left with no
    number at all, so the book printed `표 1` over a question and nothing
    over the table a reader was sent to.

    A caption is adjacent to its table, above it or below it, blank lines
    aside. A definition list item is adjacent to its term. That is the one
    difference the text carries, so it is what this asks about.
    """
    lines = md_text.split('\n')
    starts, pos = [], 0
    for line in lines:
        starts.append(pos)
        pos += len(line) + 1

    def is_row(i):
        return 0 <= i < len(lines) and bool(_MD_TABLE_ROW_RE.match(lines[i]))

    def nearest(i, step):
        i += step
        while 0 <= i < len(lines) and not lines[i].strip():
            i += step
        return i

    out = []
    for i, line in enumerate(lines):
        m = _MD_TABLE_CAPTION_RE.match(line)
        if not m:
            continue
        if is_row(nearest(i, -1)) or is_row(nearest(i, 1)):
            out.append((starts[i], starts[i] + m.end()))
    return out


_ALREADY_NUMBERED_RE = re.compile(
    r'\s*(?:\\textbf\{|\*\*)?\s*[^\s\\*{}]{1,12}\s*\d+\s*'
    r'(?:\(\s*Table\s*\d+\s*\))?\s*(?:\}|\*\*)')


def number_table_captions(md_text, temp_dir, lang_cfg=None):
    """Put "표 5 (Table 5)" in front of every table caption. (text, count).

    Figures carried their number and tables did not, so the body said "표 5에서
    보듯이" and the caption above the table said nothing a reader could match
    it against -- the same gap the equations had before they were numbered.

    Done here, on the merged markdown, because this is the last point where
    both kinds of table are still visible as text: the raw LaTeX ones keep
    their `\\caption{}` and pandoc's own tables keep their `: ` line. After
    this the two output paths render captions separately and would each need
    their own version of this.

    Only a captioned table takes a number, which is exactly LaTeX's rule.
    """
    lang_cfg = lang_cfg or {}
    label = lang_cfg.get('table_label', 'Table')
    numbers = [u['number'] for u in read_float_units(temp_dir)
               if u['kind'] == 'table' and u['number'] is not None]
    if not numbers:
        return md_text, 0

    def badge(number):
        # `%s`, not `%d`: a float counter scoped to the section carries a
        # number like `3.1`. This was the fourth site to follow from that and
        # the one that got missed — the equation taggers and the link regexes
        # were changed together, and the caption badge was not, so the build
        # died on `%d format: a real number is required, not str`.
        text = '%s %s' % (label, number)
        if label != 'Table':
            text += ' (Table %s)' % number
        return text

    # Every caption in the document, in the order a reader meets them.
    #
    # Iterate CAPTIONS, not tables: AlphaQ puts two minipages with their own
    # \caption inside one table* float, which LaTeX numbers as two tables
    # (K47). Walking tables gave both of them the float's first caption and
    # stamped two badges onto it.
    inside = sorted({(t['start'], t['stop'])
                     for t in find_raw_latex_tables(md_text)})
    events = []
    for start, stop in inside:
        for m in _CAPTION_CMD_RE.finditer(md_text, start, stop):
            # An author who keeps an old caption commented out above the live
            # one leaves two \caption commands in the float. DeeR-VLA does,
            # and counting the dead one gave its first table two numbers,
            # pushed every later table one on, and left the last with none.
            if _in_latex_comment(md_text, m.start()):
                continue
            brace = md_text.find('{', m.end() - 1)
            if 0 <= brace < stop:
                events.append((brace + 1, 'latex'))
    for line_start, at in markdown_table_captions(md_text):
        if any(a <= line_start < b for a, b in inside):
            continue                      # a stray `: ` within a raw float
        events.append((at, 'markdown'))
    events = sorted(set(events))

    pieces, cursor, used = [], 0, 0
    for at, kind in events:
        if used >= len(numbers):
            break
        pieces.append(md_text[cursor:at])
        text = badge(numbers[used])
        cursor = at
        used += 1
        # Already numbered: this text has been through here before. Skipping
        # keeps the pass idempotent, so a repair that feeds the numbered
        # markdown back does not give every caption a second badge.
        if _ALREADY_NUMBERED_RE.match(md_text, at):
            continue
        pieces.append('\\textbf{%s} ' % text if kind == 'latex'
                      else '**%s** ' % text)
    pieces.append(md_text[cursor:])
    return ''.join(pieces), used


def check_badge_placement(md_text, lang_cfg=None):
    r"""Did every table badge land on a table? Returns (ok, detail).

    Counting is what let this ship. `number_table_captions` issued fifteen
    numbers, wrote fifteen badges and reported fifteen, and ten of them were
    sitting on prose: pandoc's definition list opens with `: ` and so does a
    table caption, so the Question and Key Finding items took the numbers
    while ten real tables got none. Every count agreed. The page printed
    `표 1` over a question, and the sentence that said "see 표 1" pointed at
    it.

    So this asks where each badge IS, not how many were written. A badge
    belongs to a `\caption{}` inside a raw float, or to a `: ` line that
    abuts a markdown table. Anywhere else it is on prose.
    """
    lang_cfg = lang_cfg or {}
    label = lang_cfg.get('table_label', 'Table')
    inside = sorted({(t['start'], t['stop'])
                     for t in find_raw_latex_tables(md_text)})
    captions = {at for _start, at in markdown_table_captions(md_text)}

    # `표 1 (Table 1)` and `표 B1 (Table B1)`. A length cap was the first
    # attempt and it silently matched only the body tables: the appendix
    # badges are two characters longer, so the check validated 8 of 15 and
    # reported every one of them fine. A check that sees half the population
    # is worse than none, so the shape is spelled out instead of bounded.
    badge = re.compile(
        r'(?:\\textbf\{|\*\*)\s*' + re.escape(label)
        + r'\s+[A-Za-z]?[\d.]+(?:\s*\(Table\s+[A-Za-z]?[\d.]+\))?\s*'
          r'(?:\}|\*\*)')
    stray = []
    total = 0
    for m in badge.finditer(md_text):
        total += 1
        at = m.start()
        if any(a <= at < b for a, b in inside):
            continue                       # inside a raw float: a caption
        if any(abs(at - c) <= 2 for c in captions):
            continue                       # on a real markdown table caption
        line_start = md_text.rfind('\n', 0, at) + 1
        line_end = md_text.find('\n', at)
        stray.append(' '.join(
            md_text[line_start:line_end if line_end > 0 else at + 60].split()))

    if not stray:
        return True, '%d table badge(s), every one on a table caption' % total
    return False, ('%d of %d table badge(s) sit on prose, not on a table:\n  %s'
                   % (len(stray), total,
                      '\n  '.join(s[:96] for s in stray[:6])))


# pifont's tick and cross, which pandoc has no reader for: it drops the
# command and emits NOTHING, so a column of them comes out blank. VLA-Adapter
# lost all twelve marks in its table 7, whose entire content is which
# condition each method uses; the six success rates were left attached to
# nothing and two rows became indistinguishable. Every count agreed, because
# a mark is not a value and no probe was counting marks.
#
# Only the codes whose glyph is settled are mapped. An unknown \ding is left
# exactly as written, so it prints and can be seen, rather than being guessed
# at and quietly turned into the wrong symbol.
_DING_GLYPHS = {
    '51': '\u2713',      # check mark
    '52': '\u2714',      # heavy check mark
    '53': '\u2715',      # multiplication x
    '54': '\u2716',      # heavy multiplication x
    '55': '\u2717',      # ballot x
    '56': '\u2718',      # heavy ballot x
}
_DING_RE = re.compile(r'\\ding\s*\{\s*(\d+)\s*\}')


def substitute_dings(latex):
    """`\\ding{51}` -> the character. Returns (text, replaced, unknown)."""
    unknown = []

    def swap(m):
        glyph = _DING_GLYPHS.get(m.group(1))
        if glyph is None:
            unknown.append(m.group(1))
            return m.group(0)
        return glyph

    out, total = _DING_RE.subn(swap, latex)
    return out, total - len(unknown), sorted(set(unknown))


def _latex_fragment_to_html(latex, pandoc, work, name, inline=False,
                            math_mode='mathml'):
    """Render one LaTeX fragment to HTML. Returns '' when pandoc cannot.

    --mathml matters: without it pandoc emits `<span class="math inline">` with
    raw TeX inside, and a bare backslash there escapes the following '<' when
    the markdown reader re-reads the block, leaving a literal `</span>` printed
    in the cell. --wrap=none keeps pandoc from breaking lines inside tags.
    """
    latex, _swapped, _unknown = substitute_dings(latex)
    path = os.path.join(work, name)
    with open(path, 'w', encoding='utf-8', newline='') as fh:
        fh.write(latex + '\n')
    cmd = [pandoc, '-f', 'latex', '-t', 'html', '--wrap=none', path]
    if math_mode == 'mathml':
        cmd.insert(-1, '--mathml')
    try:
        result = subprocess.run(cmd,
                                capture_output=True, text=True,
                                encoding='utf-8', errors='replace', timeout=120)
    except (OSError, subprocess.SubprocessError):
        return ''
    if result.returncode != 0:
        return ''
    html = (result.stdout or '').strip()
    if inline:
        html = re.sub(r'^<p>|</p>$', '', html).strip()
    return html


# Plain `-t markdown` writes a table however it likes. For anything wider or
# more spanned than a few columns it chooses a SIMPLE table -- columns marked
# by character position, no `|` anywhere -- which _is_markdown_table does not
# recognise, so nine of AlphaQ's twelve tables were dropped to plain text in
# the Word file while the HTML had all twelve. It also wraps the table in a
# `::: table*` div that prints literally. Pipe tables and nothing else.
# Grid is allowed; simple and multiline are not. The rule that matters is
# that the table pandoc writes must be one `_is_markdown_table` can see and
# pandoc can read back, and a simple table satisfies neither: it marks columns
# by character position with no `|` anywhere, which is how nine of AlphaQ's
# twelve tables fell to plain text in the Word file while the HTML had all
# twelve. A grid table opens with `+---+`, which that check has always
# recognised.
#
# Turning grid off as well went further than the reason required, and it cost
# four tables per book: a table needing more than a pipe table can say -- a
# multi-line cell, a spanned column -- came back as prose and was left as raw
# LaTeX, so DeeR-VLA's DOCX carried 7 of its 11 tables. The hazard recorded
# for grid tables in `grid_tables_to_pipe` is a different one: it belongs to
# markdown a translator edits, where a widened CJK cell no longer lines up.
# Nothing edits this markdown between pandoc writing it and pandoc reading it.
_FRAGMENT_WRITER = ('markdown-simple_tables-multiline_tables+grid_tables'
                    '+pipe_tables-fenced_divs-native_divs-raw_html')


def _latex_fragment_to_markdown(latex, pandoc, work, name):
    """Render one LaTeX fragment to markdown. '' when pandoc cannot."""
    path = os.path.join(work, name)
    with open(path, 'w', encoding='utf-8', newline='') as fh:
        fh.write(latex + '\n')
    try:
        result = subprocess.run(
            [pandoc, '-f', 'latex', '-t', _FRAGMENT_WRITER, '--wrap=none',
             path],
            capture_output=True, text=True,
            encoding='utf-8', errors='replace', timeout=120)
    except (OSError, subprocess.SubprocessError):
        return ''
    return (result.stdout or '').strip() if result.returncode == 0 else ''


def _is_markdown_table(text):
    """Did pandoc actually produce a table, or just prose?"""
    return '|' in text or '+--' in text or '+==' in text


# A border may carry en/em dashes: a chunk that went through a smart-quotes
# pass has them, and a border this pattern does not recognise gets read as a
# content row -- which cut CafeQ's widest table in half.
_GRID_BORDER_RE = re.compile(r'^\+[-=:+‐-―]{3,}\+[ \t]*$')
# `|          +--------+--------+` -- the continuation border of a cell that
# spans rows. A pipe table has no way to say that, so a table containing one
# is left as it is rather than half-converted.
_INTERIOR_BORDER_RE = re.compile(r'\+[-=:‐-―]{3,}')


def grid_tables_to_pipe(md_text):
    """Rewrite any grid table as a pipe table. Returns (text, count).

    A grid table marks its columns by CHARACTER POSITION and pandoc lays one
    out by DISPLAY width, counting a Hangul syllable as two columns. Translate
    a cell and pad it to the same character count -- the obvious thing to do,
    and what a sub-agent does -- and the `|` no longer meet the `+`, so pandoc
    abandons the table and prints one cell per line with the pipes still in
    it. CafeQ shipped three that way.

    The format cannot be turned off at ingest: it is the only one pandoc can
    use for a spanning multi-deck header, and without it pandoc writes the
    literal text `[TABLE]` instead and the table is gone. So the conversion
    happens here, after translation, when nothing can drift any further.

    Cells are read from the `|` separators, never from column positions --
    the positions are exactly what has gone wrong by this point. A spanning
    cell keeps its text and the columns it covered are left empty.
    """
    lines = md_text.split('\n')
    out, i, count = [], 0, 0
    while i < len(lines):
        if not _GRID_BORDER_RE.match(lines[i]):
            out.append(lines[i])
            i += 1
            continue
        # The whole block first, so a table this cannot convert is copied
        # verbatim in one piece. Bailing out line by line let the scan
        # re-enter halfway down and eat the interior borders of a table it
        # had just decided to leave alone.
        stop = i
        while stop < len(lines) and lines[stop].strip():
            stop += 1
        ncols = len(re.findall(r'[-=:‐-―]{3,}', lines[i]))
        rows, clean = [], True
        for line in lines[i + 1:stop]:
            if _GRID_BORDER_RE.match(line):
                continue
            if '|' not in line or _INTERIOR_BORDER_RE.search(line):
                # Either not a row at all, or a row-spanning cell's
                # continuation border (`|      +-----+-----+`), a shape a
                # pipe table cannot express.
                clean = False
                break
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            rows.append((cells + [''] * ncols)[:ncols])
        j = stop
        if not clean or len(rows) < 2 or ncols < 2:
            out.extend(lines[i:stop])
            i = stop
            continue
        out.append('| ' + ' | '.join(rows[0]) + ' |')
        out.append('|' + '|'.join(['---'] * ncols) + '|')
        for row in rows[1:]:
            out.append('| ' + ' | '.join(row) + ' |')
        count += 1
        i = j
    return '\n'.join(out), count


_HEADER_END_RE = re.compile(r'\\(?:midrule|hline)\b')
_ROW_BREAK_RE = re.compile(r'\\\\')
_TR_RE = re.compile(r'<tr[^>]*>.*?</tr>', re.DOTALL)


def header_row_count(latex):
    """How many rows of this tabular are its header, per the LaTeX itself.

    pandoc finds a header by looking for a rule, and the answer it gives
    depends on which rules a paper happens to use: SINQ's main results table
    and three of AlphaQ's produced no <thead> at all, so the header rule never
    drew, the header never repeated across a page break, and not one cell was
    a <th>. Nine columns of numbers sat under nothing.

    Counting is not a heuristic: the header is whatever precedes the first
    \\midrule or \\hline that is not the top rule.
    """
    body = _header_body(latex)
    end = _first_body_rule(body)
    if end is None:
        return 0
    rows = len(_ROW_BREAK_RE.findall(body[:end]))
    total = len(_ROW_BREAK_RE.findall(body))
    return rows if 0 < rows < total else 0


def _header_body(latex):
    """The tabular's rows, with comments and the top rule out of the way."""
    body = latex[latex.find('}', latex.find('{')) + 1:]
    return re.sub(r'(?m)%.*$', '', body).replace('\\toprule', '', 1)


def _first_body_rule(body):
    """Where the header ends, skipping a rule that sits above every row.

    A booktabs table opens with \\toprule, which is removed above. An
    \\hline-ruled table opens with \\hline instead, and taking that one as the
    header's end would say the table has no header at all.
    """
    for m in _HEADER_END_RE.finditer(body):
        if _ROW_BREAK_RE.search(body[:m.start()]):
            return m.start()
    return None


_TBODY_RE = re.compile(r'<tbody[^>]*>(.*?)</tbody>', re.DOTALL)


_SOFT_RULE_RE = re.compile(r'\\addlinespace\b|\\\\\s*\[\s*\d')


def body_rule_rows(latex):
    """{row index: 'hard'|'soft'} for body rows that open a group.

    A paper marks its row groups two ways and pandoc renders neither. SINQ
    separates bit-widths with a rule; AlphaQ's Table 1 nests them -- a
    `\\midrule` between models and an `\\addlinespace` between the bit budgets
    inside each model -- so nine rows of numbers ran together with only the
    span label in the margin to break them up.

    Space is not reproducible here (the cells are one row tall either way), so
    the softer boundary becomes a lighter rule: solid between models, hairline
    between bit groups. The hierarchy survives, which is the part that was
    lost.
    """
    body = _header_body(latex)
    end = _first_body_rule(body)
    if end is None:
        return {}
    after = body[end:]
    after = after[_HEADER_END_RE.match(after).end():] \
        if _HEADER_END_RE.match(after) else after
    marked = {}
    for index, piece in enumerate(_ROW_BREAK_RE.split(after)[:-1]):
        if not index:
            continue
        if _HEADER_END_RE.search(piece):
            marked[index] = 'hard'
        elif _SOFT_RULE_RE.search(piece):
            marked[index] = 'soft'
    return marked


def mark_body_rules(html, latex):
    """Give each row that opens a group the class the print sheet draws."""
    marked = body_rule_rows(latex)
    if not marked:
        return html
    body = _TBODY_RE.search(html)
    if not body:
        return html
    rows = _TR_RE.findall(body.group(1))
    out = []
    for i, row in enumerate(rows):
        kind = marked.get(i)
        if kind:
            css = 'rule-above' if kind == 'hard' else 'rule-above-soft'
            row = re.sub(r'<tr\b', '<tr class="%s"' % css, row, count=1)
        out.append(row)
    return (html[:body.start()] + '<tbody>\n' + '\n'.join(out) + '\n</tbody>'
            + html[body.end():])


_ROWCOLOR_RE = re.compile(r'\\rowcolor\s*(?:\[[^\]]*\])?\s*\{[^{}]*\}')


def shaded_body_rows(latex):
    r"""({row index}, row count) for the body rows the paper shades.

    `\rowcolor[rgb]{ .900, .900, .900}` is how a results table says which
    rows are the authors' own. pandoc drops it, so VLA-Adapter's five
    `(Ours)` rows sat in the book looking like every competitor's. Indexed
    the way `body_rule_rows` indexes, because the same `<tbody>` rows are
    what both of them mark.

    The count comes back too: shading the wrong row would credit somebody
    else's numbers to the authors, so the caller can refuse when the source
    and the rendered table disagree about how many rows there are.
    """
    body = _header_body(latex)
    end = _first_body_rule(body)
    if end is None:
        return set(), 0
    after = body[end:]
    after = after[_HEADER_END_RE.match(after).end():] \
        if _HEADER_END_RE.match(after) else after
    pieces = _ROW_BREAK_RE.split(after)[:-1]
    return ({i for i, piece in enumerate(pieces) if _ROWCOLOR_RE.search(piece)},
            len(pieces))


def _add_row_class(row, name):
    """Add a class to a `<tr>`, keeping any it already carries."""
    m = re.match(r'<tr\b([^>]*)>', row)
    if not m:
        return row
    attrs = m.group(1)
    if 'class="' in attrs:
        attrs = re.sub(r'class="([^"]*)"',
                       lambda x: 'class="%s %s"' % (x.group(1), name),
                       attrs, count=1)
    else:
        attrs += ' class="%s"' % name
    return '<tr' + attrs + '>' + row[m.end():]


def mark_shaded_rows(html, latex):
    """Give the paper's own rows back the shading it marked them with.

    Returns (html, count). Runs before `split_row_groups`, which turns the
    single `<tbody>` into several and would leave the indices meaning
    something else.
    """
    marked, total = shaded_body_rows(latex)
    if not marked:
        return html, 0
    body = _TBODY_RE.search(html)
    if not body:
        return html, 0
    rows = _TR_RE.findall(body.group(1))
    if len(rows) != total:
        # Refuse rather than guess. A band on the wrong row is a claim about
        # whose result is whose, and it is not one this can make on a count
        # it already knows to be wrong.
        return html, 0
    out = [_add_row_class(row, 'row-shaded') if i in marked else row
           for i, row in enumerate(rows)]
    return (html[:body.start()] + '<tbody>\n' + '\n'.join(out) + '\n</tbody>'
            + html[body.end():]), len(marked)


def labelled_group_starts(latex):
    r"""Body row indices where a group carrying a label begins.

    `\multirow{10}{*}{Mixtral-8x7B}` labels ten rows and prints once, so the
    other nine carry no model name -- fine on one page, not fine across two.
    AlphaQ's table 1 broke inside such a group and the next page opened with
    `PMQ 7.42 ...` under an empty model and an empty bit budget: every row
    present, and nothing on the page saying what they are of.

    The `\multirow` itself is long gone by the time the book is built --
    conversion unwraps it to the plain label in the group's first row. What
    survives is the rule the paper drew between groups, which is the same
    boundary read a different way.
    """
    return sorted(i for i, kind in body_rule_rows(latex).items()
                  if kind == 'hard')


def split_row_groups(html, latex):
    """One <tbody> per labelled group, so a break lands between them.

    Returns (html, groups). The browser is told not to break inside a group
    and moves the page boundary to the next one; where a group is taller
    than the page it breaks anyway, which is all anything could do.
    """
    starts = labelled_group_starts(latex)
    if not starts:
        return html, 0
    body = _TBODY_RE.search(html)
    if not body:
        return html, 0
    rows = _TR_RE.findall(body.group(1))
    bounds = [s for s in starts if 0 < s < len(rows)]
    if not bounds:
        return html, 0
    bounds = [0] + bounds
    chunks = []
    for i, start in enumerate(bounds):
        stop = bounds[i + 1] if i + 1 < len(bounds) else len(rows)
        chunks.append('<tbody class="rowgroup">\n'
                      + '\n'.join(rows[start:stop]) + '\n</tbody>')
    return (html[:body.start()] + '\n'.join(chunks) + html[body.end():],
            len(chunks))


def promote_header_rows(html, latex):
    """Wrap the header rows in <thead> and make their cells <th>.

    Rebuilds the whole body region rather than splicing tags in: pandoc has
    already wrapped every row in one <tbody>, and opening a <thead> in front
    of it would leave that wrapper unbalanced.
    """
    if '<thead' in html:
        return html
    wanted = header_row_count(latex)
    if not wanted:
        return html
    body = _TBODY_RE.search(html)
    if not body:
        return html
    rows = _TR_RE.findall(body.group(1))
    if wanted >= len(rows):
        return html
    head = '\n'.join(re.sub(r'<(/?)td\b', r'<\1th', r) for r in rows[:wanted])
    rest = '\n'.join(rows[wanted:])
    return (html[:body.start()]
            + '<thead>\n%s\n</thead>\n<tbody>\n%s\n</tbody>' % (head, rest)
            + html[body.end():])


_SYMBOL_MATH_RE = re.compile(
    r'<math\b[^>]*>\s*<semantics>\s*<m[iox]>(&#?\w+;|[^<>&])</m[iox]>\s*'
    r'(?:<annotation\b[^>]*>.*?</annotation>\s*)?</semantics>\s*</math>',
    re.DOTALL)


def simplify_symbol_math(html):
    """A formula that is one symbol is a character. Returns (html, n).

    Chromium repeats a `<thead>` on every page its table runs onto, and out
    of a two-row header it drops the inline `<math>` when it does. AlphaQ's
    table 1 said `WikiText2 ↓` and `정확도 ↑` on the first page and neither on
    the second, so the continuation never said which direction was better --
    and every count still balanced, because the header was there.

    `<mo>↓</mo>` inside `<semantics>` draws exactly what the bare character
    draws. The bare character also survives the repeat.

    Only the header. A body cell is never repeated, so it never hits this.

    The character keeps a `math` class rather than standing bare, because
    the math check counts formulas asked for against formulas delivered and
    a symbol dropped out of that total is a hole the check would stop being
    able to see.
    """
    total = [0]

    def in_head(m):
        head, n = _SYMBOL_MATH_RE.subn(
            lambda s: '<span class="math-symbol">%s</span>' % s.group(1),
            m.group(0))
        total[0] += n
        return head

    return re.sub(r'(?s)<thead\b.*?</thead>', in_head, html), total[0]


_BIBITEM_LABEL_RE = re.compile(r'\\bibitem\s*\[([^\]]*)\]\s*\{([^{}]*)\}')
_LABEL_YEAR_RE = re.compile(r'^(.*?)\(\s*([^()]*?)\s*\)')
_FRAGMENT_CITE_RE = re.compile(
    r'\\cite([a-zA-Z]*)\s*(?:\[[^\]]*\])*\s*\{([^{}]*)\}')
# `\text{PQE}` is amsmath's, and outside math pandoc drops it with its body.
# CafeQ's table 7 printed no header at all over the column it is sorted by.
# `\textrm` says the same thing in both modes.
_TEXT_MACRO_RE = re.compile(r'\\text(?=\s*\{)')


_COLOR_DECL_RE = re.compile(r'\\color\s*(?:\[[^\]]*\])?\s*\{([^{}]*)\}')


def _cell_end(tex, pos):
    r"""Where a bare `\color` stops: the end of its table cell."""
    depth = 0
    i = pos
    while i < len(tex):
        ch = tex[i]
        if ch == '{':
            depth += 1
        elif ch == '}':
            if depth == 0:
                return i
            depth -= 1
        elif depth == 0:
            if ch == '&':
                return i
            if ch == '\\' and tex[i:i + 2] == '\\\\':
                return i
        i += 1
    return len(tex)


def rewrite_color_declarations(tex):
    r"""`{\color{red} 17.14}` -> `{\textcolor{red}{17.14}}`. (text, count).

    The declaration form colours everything to the end of its group and
    pandoc's reader drops it; the command form survives, which is why the
    word `빨간색` in SINQ's captions prints red while the values it points at
    do not. Twelve marked results in tables 3 and 4 came out black, and both
    captions tell the reader to look for red — so the caption promises
    something the page cannot deliver, and no count noticed because every
    number was present.
    """
    out, cursor, count = [], 0, 0
    for m in _COLOR_DECL_RE.finditer(tex):
        if m.start() < cursor:
            continue
        # `{\color{red} X}` scopes to its brace; a bare `\color{red} X` in a
        # cell scopes to the `&`. SINQ writes both, and a third: the
        # declaration inside `\textbf{...}`. The braces are always left in
        # place, or `\textbf` loses the argument it was holding.
        opening = tex.rfind('{', cursor, m.start())
        if opening >= 0 and not tex[opening + 1:m.start()].strip():
            close = _balanced_group(tex, opening)
            stop = close - 1 if close > 0 else -1
        else:
            stop = _cell_end(tex, m.end())
        if stop < 0:
            continue
        body = tex[m.end():stop].strip()
        if not body:
            continue
        out.append(tex[cursor:m.start()])
        out.append('\\textcolor{%s}{%s}' % (m.group(1), body))
        cursor = stop
        count += 1
    out.append(tex[cursor:])
    return ''.join(out), count


def build_citation_labels(md_text):
    r"""{key: 'Hendrycks et al. 2021a'} from the inlined natbib bibliography.

    A citation inside a raw table never reaches the resolver: the float is
    kept verbatim from flat.tex, so `\citep{...}` arrives at pandoc, which
    has no bibliography here and drops the call together with its key. CafeQ's
    table 6 exists to say which benchmark came from which paper, and all
    sixteen sources were deleted out of it.
    """
    labels = {}
    for m in _BIBITEM_LABEL_RE.finditer(md_text):
        split = _LABEL_YEAR_RE.match(m.group(1))
        if not split:
            continue
        authors = split.group(1).replace('~', ' ').strip()
        year = re.sub(r'[{}]', '', split.group(2)).strip()
        if authors and year:
            labels[m.group(2).strip()] = '%s %s' % (authors, year)
    return labels


_BIB_ENTRY_RE = re.compile(r'@[A-Za-z]+\s*\{\s*([^,\s]+)\s*,(.*?)(?=\n@|\Z)',
                           re.DOTALL)
_BIB_FIELD_RE = re.compile(r'(?im)^\s*(author|year|date)\s*=\s*'
                           r'(\{(?:[^{}]|\{[^{}]*\})*\}|"[^"]*"|\d+)')


def _bib_surname(chunk):
    """The family name out of one BibTeX author entry."""
    chunk = re.sub(r'[{}\\]', '', chunk).strip()
    if ',' in chunk:                       # `Last, First`
        return chunk.split(',')[0].strip()
    parts = chunk.split()                  # `First Last`
    return parts[-1] if parts else ''


def build_citation_labels_from_bib(temp_dir):
    r"""{key: 'Authors Year'} from the `.bib` the paper shipped.

    `build_citation_labels` reads an INLINED `\bibitem` list, which is one of
    the two shapes a paper arrives in. VLA-Adapter ships a `.bib` and lets
    citeproc render it, so that map came back empty, `resolve_fragment_
    citations` had nothing to resolve with, and pandoc dropped all 51 `\citep`
    calls inside its tables: a 22-baseline comparison in which no number could
    be traced to the paper it came from. The mechanism was there and called;
    it was built from the one bibliography this paper does not use.

    Returns {} when there is no tarball or no `.bib`, so the inlined path is
    unaffected.
    """
    src = os.path.join(temp_dir or '', 'arxiv_src')
    if not temp_dir or not os.path.isdir(src):
        return {}
    labels = {}
    for root, _dirs, names in os.walk(src):
        for name in sorted(names):
            if not name.endswith('.bib'):
                continue
            try:
                with open(os.path.join(root, name), encoding='utf-8',
                          errors='replace') as fh:
                    text = fh.read()
            except OSError:
                continue
            for key, body in _BIB_ENTRY_RE.findall(text):
                fields = {}
                for field, raw in _BIB_FIELD_RE.findall(body):
                    fields[field.lower()] = raw.strip('{}" ')
                year = fields.get('year') or ''
                if not year:
                    year = (re.search(r'\b(1[89]|20)\d{2}\b',
                                      fields.get('date', '')) or [''])
                    year = year.group(0) if hasattr(year, 'group') else ''
                names_ = [a for a in re.split(r'\s+and\s+',
                                              fields.get('author', ''))
                          if a.strip()]
                surnames = [_bib_surname(a) for a in names_]
                surnames = [s for s in surnames if s]
                if not surnames or not year:
                    continue
                if len(surnames) == 1:
                    who = surnames[0]
                elif len(surnames) == 2:
                    who = '%s and %s' % (surnames[0], surnames[1])
                else:
                    who = '%s et al.' % surnames[0]
                labels.setdefault(key.strip(), '%s %s' % (who, year))
    return labels


_FRAGMENT_REF_RE = re.compile(r'\\ref\s*\{([^{}]+)\}')


def resolve_fragment_references(tex, numbers):
    r"""`\ref{TableD1}` inside a raw table -> `D1`. (text, done, missed).

    A protected float never meets `resolve_references`, so its `\ref` calls
    reach pandoc, which prints the key. VLA-Adapter's captions carried twenty
    of them, `[TableD1]` and `[AppendixG]` among others, each a pointer the
    reader cannot follow.

    Only the NUMBER is substituted, never the word. All twenty already had a
    Korean label in front of them, written by the translator, so supplying
    another would print it twice, and choosing between 부록 and 절 for a
    section key would be a guess this does not need to make.

    The number comes from the index, not from the key, because the two
    disagree: this paper labels its appendix H `AppendixG` and LaTeX prints
    H. A key that resolves to nothing is left exactly as written, so it stays
    visible rather than becoming a confident wrong letter.
    """
    missed = []

    def swap(m):
        value = numbers.get(m.group(1).strip())
        if value is None:
            missed.append(m.group(1).strip())
            return m.group(0)
        return str(value)

    out, total = _FRAGMENT_REF_RE.subn(swap, tex)
    return out, total - len(missed), sorted(set(missed))


def fragment_reference_numbers(temp_dir):
    r"""{label: printed number} for anything a `\ref` inside a float names.

    Empty without a temp dir. `expand_raw_latex_tables` is called with none
    in four existing tests and by any caller that only has markdown, and
    `build_label_index` joins the path unguarded.
    """
    if not temp_dir or not os.path.isdir(temp_dir):
        return {}
    numbers = {}
    for label, number in build_float_numbers(temp_dir).items():
        numbers.setdefault(label, str(number))
    for key, entry in build_label_index(temp_dir).items():
        # Keyed by the label exactly as written. The colon in `eq:pqe` is
        # part of the name, not a kind prefix -- an earlier version here
        # also registered the tail, which buys nothing (`\ref` always writes
        # the full label) and silently collides `eq:main` with `tab:main`.
        numbers.setdefault(key, str(entry[0]))
    return numbers


def resolve_fragment_citations(tex, labels):
    """Render a raw table's `\\citep{key}` the way the body renders it.

    Returns (text, count). A key with no entry leaves the whole call alone:
    a citation that is visibly unresolved can be fixed, and one silently
    attributed to the wrong paper cannot be noticed at all.
    """
    if not labels:
        return tex, 0
    count = [0]

    def sub(m):
        keys = [k.strip() for k in m.group(2).split(',') if k.strip()]
        shown = [labels[k] for k in keys if k in labels]
        if not keys or len(shown) != len(keys):
            return m.group(0)
        count[0] += 1
        joined = '; '.join(shown)
        if m.group(1) in ('t', 'author', 'alt'):
            head, _sep, year = joined.rpartition(' ')
            return '%s (%s)' % (head, year) if head else joined
        return '(%s)' % joined

    return _FRAGMENT_CITE_RE.sub(sub, tex), count[0]


_CELL_WRAPPER_HEAD_RE = re.compile(
    r'\\(?:makecell|thead)\s*(?:\[[^\]\n]*\])?\s*(?=\{)')
# The row break INSIDE such a cell. It separates two lines of one cell, not
# two table rows, so it becomes a space rather than a break.
_CELL_ROW_BREAK_RE = re.compile(r'\s*\\\\\s*')


def unwrap_table_cell_wrappers(md_text):
    r"""`\makecell{a\\b}` -> `a b`. Returns (text, count).

    makecell's job is a line break inside one cell, and pandoc has no reader
    for it — so the command AND its argument are dropped and the cell arrives
    EMPTY. Measured: a cell holding `\makecell{one\\two}` renders as `<td></td>`
    with nothing said, the same silent swallow as K110's wrappers but one cell
    at a time. PaLM has 92 of them and `\thead` behaves identically.

    Joining the lines with a space keeps every word; only the line break is
    lost, which is typesetting rather than content.
    """
    count = 0
    while True:
        m = _CELL_WRAPPER_HEAD_RE.search(md_text)
        if not m:
            return md_text, count
        open_at = md_text.index('{', m.end() - 1)
        close = _balanced_group(md_text, open_at)
        if close < 0:
            return md_text, count
        inner = md_text[open_at + 1:close - 1]
        md_text = (md_text[:m.start()]
                   + _CELL_ROW_BREAK_RE.sub(' ', inner).strip()
                   + md_text[close:])
        count += 1


_TABBING_RE = re.compile(r'\\begin\{tabbing\}(.*?)\\end\{tabbing\}', re.S)
# The tab-stop template line: escaped spaces and `\=` marks ending in `\kill`.
# It positions the columns and prints nothing at all.
_TABBING_KILL_RE = re.compile(r'^[^\n]*\\kill[ \t]*$\n?', re.M)
_TABBING_STOP_RE = re.compile(r'\\[=>]')
# A code fence renders nothing, so `$\,-$` prints as five characters where the
# paper prints a minus sign. Shor's three listings hold nineteen of these, and
# the whole set is eleven fragments using two commands — `\,` and his own
# `\mod{n}`. The delimiters and the spacing come off; anything else is left
# standing rather than guessed at, so a macro the fence cannot render stays
# visible instead of quietly becoming something wrong.
_TABBING_MATH_RE = re.compile(r'\$([^$\n]*)\$')
# No letter boundary: `\,` is a control SYMBOL, complete in two characters, so
# the next character is never part of its name. Requiring a non-letter after it
# left `result_{\,i}` standing in the third listing.
_MATH_SPACING_RE = re.compile(r'\\[,;:!]')


def _tabbing_math(m):
    return _MATH_SPACING_RE.sub('', m.group(1))
_TABBING_FONT_RE = re.compile(r'\{\s*\\(?:it|rm|bf|tt|sf|sc)\s+([^{}]*)\}')


_ONE_ARG_DEF_RE = re.compile(
    r'\\(?:new|renew|provide)command\s*\{?\s*\\([A-Za-z]+)\s*\}?\s*\[1\]\s*\{')


def read_one_argument_macros(temp_dir):
    r"""{name: body-with-#1} for the paper's own single-argument shorthand.

    `read_math_macros` refuses these on purpose — expanding a macro properly
    needs a real expander and guessing corrupts formulas that render today
    (K121). This is a much smaller claim, used only inside a `tabbing` fence,
    where nothing renders at all and the alternative is printing `\mod{n}` at
    the reader. One argument, one body, and anything carrying a conditional is
    left alone.
    """
    macros = {}
    flat = os.path.join(temp_dir or '', 'flat.tex')
    if not os.path.isfile(flat):
        return macros
    try:
        with open(flat, 'r', encoding='utf-8', errors='replace') as fh:
            tex = strip_tex_comments(fh.read())
    except OSError:
        return macros
    tex = tex.split(r'\begin{document}')[0]
    for m in _ONE_ARG_DEF_RE.finditer(tex):
        close = _balanced_group(tex, m.end() - 1)
        if close < 0:
            continue
        body = tex[m.end():close - 1]
        if body.count('#1') != 1 or '#2' in body or r'\if' in body:
            continue
        macros[m.group(1)] = body
    return macros


def _expand_one_arg(text, macros):
    r"""`\mod{n}` -> ` (mod n)` for the macros this paper defines."""
    for name, body in macros.items():
        head = re.compile(r'\\%s(?![A-Za-z])\s*\{' % re.escape(name))
        while True:
            m = head.search(text)
            if not m:
                break
            close = _balanced_group(text, m.end() - 1)
            if close < 0:
                break
            arg = text[m.end():close - 1]
            text = text[:m.start()] + body.replace('#1', arg) + text[close:]
    return text


def unwrap_tabbing(md_text, macros=None):
    r"""Turn a `tabbing` environment into a code block. Returns (text, count).

    `tabbing` is how a 1990s paper sets aligned pseudocode, and pandoc has no
    reader for it — so on the HTML path the whole environment is dropped
    without a word, exactly as K110's wrappers were. Shor's three algorithm
    listings vanished that way, and only the raw-block fidelity count saw it.

    A code fence is the honest target: the environment's whole purpose here is
    preformatted alignment, which is what a fence preserves. The tab stops
    become indentation, the `\kill` template line goes (it prints nothing),
    and the old font switches around variable names are unwrapped.
    """
    count = [0]

    def convert(m):
        body = _TABBING_KILL_RE.sub('', m.group(1))
        # Before the font unwrap, deliberately: `\mod{n}` expands to
        # `{\rm \ (mod\ }n)`, and it is that unwrap which turns it into text.
        if macros:
            body = _expand_one_arg(body, macros)
        body = _TABBING_FONT_RE.sub(r'\1', body)
        body = _TABBING_STOP_RE.sub('    ', body)
        body = _TABBING_MATH_RE.sub(_tabbing_math, body)
        # `\\` ends a row; `\ ` is an escaped space holding indentation.
        # The source usually breaks the line after `\\` as well, so replacing
        # it with a newline doubles every gap and the listing prints
        # double-spaced. Take the newline that follows with it.
        body = re.sub(r'\\\\[ \t]*\n?', '\n', body).replace('\\ ', ' ')
        lines = [line.rstrip() for line in body.split('\n')]
        while lines and not lines[0].strip():
            lines.pop(0)
        while lines and not lines[-1].strip():
            lines.pop()
        if not lines:
            return ''
        count[0] += 1
        return '\n```\n%s\n```\n' % '\n'.join(lines)

    return _TABBING_RE.sub(convert, md_text), count[0]


_AT_SPACING_RE = re.compile(r'@\{[^{}]*\}')
# Brace-aware: a spec like `{|c|@{}}` contains braces, and a flat
# `\{[^{}]*\}` cannot match it. A flat pattern reported these specs as plain
# `{c}` and sent three hypotheses down the wrong path before the real shape
# turned up.
_MULTICOLUMN_SPEC_RE = re.compile(
    r'(\\multicolumn\s*\{[^{}]*\}\s*\{)((?:[^{}]|\{[^{}]*\})*)(\})')


def drop_multicolumn_spacing(tex):
    r"""Remove `@{}` from `\multicolumn` column specs. Returns (text, count).

    `@{}` suppresses inter-column padding and says nothing about content, but
    pandoc refuses a `\multicolumn{3}{|c|@{}}` outright — measured: the same
    table converts with `{|c|}` and not with `{|c|@{}}`, while `@{}` in the
    TABULAR spec is read without complaint. Shor's two truth-table floats were
    dropped whole on the HTML path over it, and pandoc's markdown reader drops
    a raw block without a word (K57), so only the table-fidelity count noticed.
    """
    return _MULTICOLUMN_SPEC_RE.subn(
        lambda m: m.group(1) + _AT_SPACING_RE.sub('', m.group(2))
        + m.group(3), tex)


def expand_raw_latex_tables(md_text, pandoc=None, math_mode='mathml',
                            output='html', temp_dir=None):
    """Convert raw LaTeX tables into real tables.

    `output='html'` gives HTML tables carrying the sizing classes the print
    sheet needs. `output='markdown'` gives markdown tables instead, which is
    what the DOCX path requires: raw HTML survives ONLY the HTML path, and
    injecting `<table>` left book.docx with zero tables and no complaint from
    any check. pandoc cannot express every LaTeX table as markdown, so the two
    are produced separately rather than one being derived from the other.

    Returns (new_text, converted, failed). A table pandoc cannot read is left
    exactly as it was, so check_table_fidelity reports the shortfall instead of
    the build quietly shipping a book with a hole in it.
    """
    tables = find_raw_latex_tables(md_text)
    if not tables:
        return md_text, 0, 0

    pandoc = pandoc or resolve_pandoc()
    if not pandoc:
        print("WARNING: %d raw LaTeX table(s) need pandoc to be expanded; "
              "they will be missing from the output" % len(tables))
        return md_text, 0, len(tables)

    converted = failed = 0
    cite_labels = build_citation_labels(md_text)
    if not cite_labels:
        # No inlined `\bibitem` list. The paper shipped a `.bib` and let
        # citeproc render it, so the keys live there instead (K152).
        cite_labels = build_citation_labels_from_bib(temp_dir)
    ref_numbers = fragment_reference_numbers(temp_dir)
    refs_done, refs_missed = 0, []
    shaded_rows = 0
    pieces, cursor = [], 0
    with tempfile.TemporaryDirectory(prefix='tb-tex-') as work:
        for n, t in enumerate(tables):
            # A band label is written `\multirow{4}{*}{\rotatebox{90}{...}}`
            # and reaches here whole, because a protected table float is kept
            # verbatim all the way from flat.tex. pandoc drops both calls with
            # their bodies, so nine of SINQ's tables rendered their group
            # column empty: the same four method rows twice, with nothing
            # saying which block was 3-bit and which was 4-bit.
            t['bare'], _labels = arxiv_backend.unwrap_rotatebox(t['bare'])
            t['bare'], _notes = arxiv_backend.unwrap_table_notes(t['bare'])
            t['bare'], _at = drop_multicolumn_spacing(t['bare'])
            # A protected float never met the citation resolver or the
            # leftover-command pass, so both have to happen here or the
            # reader loses the source beside a benchmark and the header
            # above a column.
            for part in ('bare', 'caption', 'notes'):
                if not t.get(part):
                    continue
                fixed, _n = resolve_fragment_citations(t[part], cite_labels)
                fixed, done, missed = resolve_fragment_references(
                    fixed, ref_numbers)
                refs_done += done
                refs_missed.extend(missed)
                fixed, _c = rewrite_color_declarations(fixed)
                t[part] = _TEXT_MACRO_RE.sub(r'\\textrm', fixed)
            if output == 'markdown':
                body = _latex_fragment_to_markdown(t['bare'], pandoc, work,
                                                   'table%d.tex' % n)
                if not _is_markdown_table(body):
                    failed += 1
                    continue
                if t['caption']:
                    cap = _latex_fragment_to_markdown(t['caption'], pandoc, work,
                                                      'cap%d.tex' % n)
                    if cap:
                        body = '**%s**\n\n%s' % (cap.replace('\n', ' ').strip(), body)
                if t.get('notes'):
                    note = _latex_fragment_to_markdown(t['notes'], pandoc, work,
                                                       'note%d.tex' % n)
                    if note:
                        body = '%s\n\n%s' % (body, note.strip())
                pieces.append(md_text[cursor:t['start']])
                pieces.append('\n\n' + body + '\n\n')
                cursor = t['stop']
                converted += 1
                continue

            html = _latex_fragment_to_html(t['bare'], pandoc, work,
                                           'table%d.tex' % n, math_mode=math_mode)
            if '<table' not in html:
                failed += 1
                continue
            html = promote_header_rows(html, t['bare'])
            html = mark_body_rules(html, t['bare'])
            html, _shaded = mark_shaded_rows(html, t['bare'])
            shaded_rows += _shaded
            html, _groups = split_row_groups(html, t['bare'])
            html, _symbols = simplify_symbol_math(html)
            # A paper table can carry ten columns; at body size that wraps every
            # header mid-word. Tag it so the print sheet can step the size down.
            rows = re.findall(r'<tr>.*?</tr>', html, re.DOTALL)
            columns = max((r.count('<td') + r.count('<th')) for r in rows) if rows else 0
            if columns > 6:
                html = html.replace('<table>', '<table class="cols-many">', 1)
            elif columns > 4:
                html = html.replace('<table>', '<table class="cols-wide">', 1)
            if t['caption']:
                cap = _latex_fragment_to_html(t['caption'], pandoc, work,
                                              'cap%d.tex' % n, inline=True,
                                              math_mode=math_mode)
                if cap:
                    html = re.sub(r'<table[^>]*>',
                                  lambda m: m.group(0) + '\n<caption>%s</caption>' % cap,
                                  html, count=1)
            if t.get('notes'):
                note = _latex_fragment_to_html(t['notes'], pandoc, work,
                                               'note%d.tex' % n,
                                               math_mode=math_mode)
                if note:
                    html += '\n<div class="table-notes">%s</div>' % note.strip()
            # --mathml keeps the original TeX in <annotation>, so the block
            # is full of backslashes -- and the markdown reader still applies
            # backslash escapes inside a raw HTML block. That is what ate the
            # '<' in <mi>\</mi> and printed a literal </mi> in the cell.
            # Numeric entities are inert to markdown and decode back to a
            # backslash in the HTML parser, so the TeX annotation stays intact.
            html = html.replace('\\', '&#92;')
            pieces.append(md_text[cursor:t['start']])
            # Blank lines both sides so the markdown reader sees one raw HTML
            # block rather than an inline run inside a paragraph.
            pieces.append('\n\n' + html + '\n\n')
            cursor = t['stop']
            converted += 1
    pieces.append(md_text[cursor:])
    if refs_done or refs_missed:
        print("Raw LaTeX tables: %d cross-reference(s) in captions resolved"
              % refs_done)
    if refs_missed:
        # Named, not swallowed: a key nobody can resolve prints as itself and
        # the reader meets `[TableD1]` where a number belongs.
        print("  %d key(s) had no number and will print raw: %s"
              % (len(set(refs_missed)), ', '.join(sorted(set(refs_missed)))))
    if shaded_rows:
        print("Raw LaTeX tables: %d row(s) the paper shades marked as its own"
              % shaded_rows)
    return ''.join(pieces), converted, failed


def convert_with_pandoc(md_file, html_file, title, lang_attr, math_mode='mathml'):
    """Convert markdown to HTML using pandoc.

    math_mode='mathml' emits native MathML with no JavaScript and no external
    requests, so equations render offline and survive into EPUB/PDF. KaTeX and
    MathJax modes are deliberately not offered: pandoc points them at a CDN,
    which is dead in an offline reader and inert inside DOCX.
    """
    pandoc = resolve_pandoc()
    if not pandoc:
        return False

    cmd = [
        pandoc, md_file, '-o', html_file,
        '--standalone',
        '--metadata', f'title={title}',
        '--metadata', f'lang={lang_attr}',
        '--from', pandoc_from(lang_attr),
        '--to', 'html5',
        '--wrap=preserve',
    ]
    if math_mode == 'mathml':
        cmd.append('--mathml')

    try:
        result = subprocess.run(
            cmd, capture_output=True, text=True,
            encoding='utf-8', errors='replace', check=True
        )
        if result.stderr and result.stderr.strip():
            print(f"pandoc warnings:\n{result.stderr.strip()[:1500]}")
        print(f"Converted with pandoc ({math_mode} math)")
        return True
    except subprocess.CalledProcessError as e:
        print(f"Pandoc conversion failed (rc={e.returncode}): {(e.stderr or '')[:1500]}")
        return False


# A math span: $$...$$ (may span lines) or $...$ (single line, non-empty).
_MATH_SPAN_RE = re.compile(r'(?<!\\)(\$\$.+?\$\$|\$[^$\n]+?\$)', re.DOTALL)


def _protect_math(text):
    """Stash math spans behind sentinels so markdown emphasis rules cannot
    mangle them. LaTeX subscripts (`x_1 ... y_2`) are otherwise eaten by the
    `_italic_` rule, which silently corrupts every formula."""
    stash = []

    def take(m):
        stash.append(m.group(1))
        return f'\x00MATH{len(stash) - 1}\x00'

    return _MATH_SPAN_RE.sub(take, text), stash


def _restore_math_as_tex(html, stash):
    """Re-insert stashed math as marked-up TeX. Tier 1 has no TeX->MathML
    engine, so the formula stays legible and machine-recoverable rather than
    being silently mangled."""
    for i, raw in enumerate(stash):
        span = f'<span class="math tex-fallback">{_html_lib.escape(raw)}</span>'
        html = html.replace(f'\x00MATH{i}\x00', span)
    return html


def convert_with_python_markdown(md_file, html_file, title):
    """Convert markdown to HTML using python-markdown (fallback 1)"""
    if not MARKDOWN_AVAILABLE:
        return False

    try:
        with open(md_file, 'r', encoding='utf-8') as f:
            md_content = f.read()

        md_content, math_stash = _protect_math(md_content)

        # 'codehilite' is dropped: it silently requires Pygments. 'attr_list'
        # and 'md_in_html' keep pandoc-style attributes and raw HTML blocks
        # (e.g. <figure>) intact.
        extensions = ['toc', 'tables', 'fenced_code', 'attr_list',
                      'md_in_html', 'sane_lists', 'nl2br']
        md = markdown.Markdown(extensions=extensions)
        html_content = md.convert(md_content)
        html_content = _restore_math_as_tex(html_content, math_stash)
        if math_stash:
            print(f"WARNING: tier-1 converter used — {len(math_stash)} formula(s) "
                  f"emitted as raw TeX, not rendered math.")

        full_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="UTF-8">
<title>{title}</title>
</head>
<body>
{html_content}
</body>
</html>"""

        with open(html_file, 'w', encoding='utf-8') as f:
            f.write(full_html)

        print("Converted with python-markdown (fallback)")
        return True
    except Exception as e:
        print(f"python-markdown conversion failed: {e}")
        return False


def convert_with_basic_regex(md_file, html_file, title):
    """Convert markdown to HTML using basic regex (fallback 2)"""
    try:
        with open(md_file, 'r', encoding='utf-8') as f:
            md_content = f.read()

        # Even in the degraded path, keep math out of reach of the emphasis
        # regexes below — `_` is pervasive in LaTeX subscripts.
        html_content, math_stash = _protect_math(md_content)

        # Headers
        html_content = re.sub(r'^#### (.*?)$', r'<h4>\1</h4>', html_content, flags=re.MULTILINE)
        html_content = re.sub(r'^### (.*?)$', r'<h3>\1</h3>', html_content, flags=re.MULTILINE)
        html_content = re.sub(r'^## (.*?)$', r'<h2>\1</h2>', html_content, flags=re.MULTILINE)
        html_content = re.sub(r'^# (.*?)$', r'<h1>\1</h1>', html_content, flags=re.MULTILINE)

        # Bold and italic
        html_content = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', html_content)
        html_content = re.sub(r'\*(.*?)\*', r'<em>\1</em>', html_content)
        # Narrowed: only underscores that delimit a word, so identifiers like
        # PL_Alpha_Hill and aten::copy_ survive intact.
        html_content = re.sub(r'(?<![\w\\])_([^_\n]+)_(?![\w])', r'<em>\1</em>', html_content)

        # Images — escape alt and src so quotes in alt text don't break the tag
        def _md_img_to_html(m):
            alt = _html_lib.escape(m.group(1), quote=True)
            src = _html_lib.escape(m.group(2), quote=True)
            return f'<img src="{src}" alt="{alt}">'
        html_content = re.sub(r'!\[([^\]]*)\]\(([^)]*)\)', _md_img_to_html, html_content)

        # Links
        html_content = re.sub(r'\[([^\]]*)\]\(([^)]*)\)', r'<a href="\2">\1</a>', html_content)

        # Lists and paragraphs
        lines = html_content.split('\n')
        result_lines = []
        in_list = False

        for line in lines:
            stripped = line.strip()
            if stripped.startswith('- '):
                if not in_list:
                    result_lines.append('<ul>')
                    in_list = 'ul'
                item = stripped[2:]
                result_lines.append(f'<li>{item}</li>')
            elif re.match(r'^\d+\. ', stripped):
                if not in_list:
                    result_lines.append('<ol>')
                    in_list = 'ol'
                item = re.sub(r'^\d+\. ', '', stripped)
                result_lines.append(f'<li>{item}</li>')
            else:
                if in_list:
                    result_lines.append(f'</{in_list}>')
                    in_list = False
                if stripped and not stripped.startswith('<'):
                    result_lines.append(f'<p>{line}</p>')
                else:
                    result_lines.append(line)

        if in_list:
            result_lines.append(f'</{in_list}>')

        html_content = '\n'.join(result_lines)

        # Page separators
        html_content = re.sub(r'<p>---</p>', '<div class="page-separator"></div>', html_content)

        html_content = _restore_math_as_tex(html_content, math_stash)

        full_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="UTF-8">
<title>{title}</title>
</head>
<body>
{html_content}
</body>
</html>"""

        with open(html_file, 'w', encoding='utf-8') as f:
            f.write(full_html)

        print("Converted with basic regex (fallback 2)")
        return True
    except Exception as e:
        print(f"Basic regex conversion failed: {e}")
        return False


_TEMPLATE_TOKEN_RE = re.compile(
    r'\$(body|title|title_page|lang|body_font|toc_label'
    r'|page_size|page_margin|print_font_size|print_line_height'
    r'|h1_break_before|h1_page_break_before)\$')


def build_title_page(title, source=''):
    r"""The page a paper opens with, or '' when there is nothing to put on it.

    The book opened on its table of contents: no title, no authors, no
    affiliations. VLA-Adapter names sixteen people in its title block and the
    book credited one of them, in a `<meta>` tag no reader sees.

    Carries no byline of its own on purpose. `apply_template_to_html` already
    puts one after the first `</h1>`, which is now this heading -- a second
    one here printed both, the short metadata form above the full list. The
    names reach it through the `byline` argument instead.

    Returns '' without a title, so a source that never had one does not gain
    a blank leaf.
    """
    if not (title or '').strip():
        return ''
    parts = ['<section class="title-page">',
             '<h1 class="title-page-title">%s</h1>'
             % _html_lib.escape(title.strip())]
    if (source or '').strip():
        parts.append('<p class="title-page-source">%s</p>'
                     % _html_lib.escape(source.strip()))
    parts.append('</section>')
    return '\n'.join(parts)


def apply_template_to_html(html_content, template_file, output_file, title, lang_cfg,
                           author=None, print_cfg=None, title_page='',
                           byline=None):
    """Apply a template to HTML content with language-aware substitutions.

    `author` is the metadata form and goes in the `<meta>` tag. `byline` is
    what the reader sees under the title; it defaults to `author` and is
    given the full list when one is known, because the metadata form is a
    catalogue entry -- "Yihao Wang et al." -- and a title page that credits
    one of sixteen people is not a title page.
    """
    if not template_file or not os.path.exists(template_file):
        print(f"Warning: Template {template_file} not found")
        return False

    try:
        with open(template_file, 'r', encoding='utf-8') as f:
            template_content = f.read()

        # Normalize the body placeholder so one substitution pass handles all
        # template shapes.
        if '$body$' not in template_content:
            if '{{content}}' in template_content:
                template_content = template_content.replace('{{content}}', '$body$')
            elif '</body>' in template_content:
                template_content = template_content.replace('</body>', '$body$\n</body>')
            else:
                template_content = template_content + '$body$'

        values = {
            'body': html_content,
            'title': title,
            'title_page': title_page,
            'lang': lang_cfg['lang_attr'],
            'body_font': lang_cfg['font_family'],
            'toc_label': lang_cfg['toc_label'],
        }
        # Every alternative in _TEMPLATE_TOKEN_RE must have a value here, or a
        # template that happens to use one raises KeyError inside the sub().
        # template.html ignores these; that costs nothing.
        values.update(layout.template_values(print_cfg or layout.get_print_profile()))

        # ONE pass over the TEMPLATE only. Substituted values are never
        # rescanned, so `$...$` math in the body — or a `$` in the title —
        # can never be reinterpreted as a template token.
        full_html = _TEMPLATE_TOKEN_RE.sub(lambda m: values[m.group(1)], template_content)

        # Inject author meta tag into <head> so calibre_html_publish.py can extract it
        if author:
            # escape(): a quote in the author name would otherwise close the
            # attribute early. Callable replacement: a backslash or `\g<1>` in
            # the name would otherwise be read as a regex escape and corrupt
            # (or crash) the substitution.
            author_meta = f'<meta name="author" content="{_html_lib.escape(author, quote=True)}">'
            if '<head>' in full_html or '<head ' in full_html:
                full_html = re.sub(
                    r'(<head[^>]*>)',
                    lambda m: m.group(1) + '\n    ' + author_meta,
                    full_html,
                    count=1,
                    flags=re.IGNORECASE
                )
            # And on the page. The names reached the metadata and stopped
            # there, so all three books opened with a bare title and no byline
            # anywhere -- a paper that does not say who wrote it. The title is
            # the body's first <h1>; the contents page is added after this.
            shown = (byline or author).replace(';', ',')
            byline_html = '<p class="byline">%s</p>' % _html_lib.escape(
                re.sub(r'\s{2,}', ' ', shown).strip())
            # Searched from <body>, not from the top of the file. The old
            # `re.sub(r'</h1>', ..., count=1)` matched the whole document,
            # and a CSS comment in <head> that merely NAMED the tag took the
            # byline: sixteen authors went into the stylesheet, where they
            # rendered as nothing at all and the title page came out bare.
            body_at = re.search(r'<body[^>]*>', full_html, re.IGNORECASE)
            start = body_at.end() if body_at else 0
            heading = re.search(r'</h1>', full_html[start:])
            if heading:
                at = start + heading.end()
                full_html = (full_html[:at] + '\n' + byline_html
                             + full_html[at:])

        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(full_html)
        return True
    except Exception as e:
        print(f"Error applying template: {e}")
        return False


_MD_TABLE_DELIM_RE = re.compile(
    r'^\s*\|?(?:\s*:?-{2,}:?\s*\|)+\s*:?-{2,}:?\s*\|?\s*$'
)


def count_md_tables(md_text):
    """Count pipe tables by a delimiter row preceded by a header row."""
    lines = md_text.splitlines()
    total = 0
    for i, line in enumerate(lines):
        if i and _MD_TABLE_DELIM_RE.match(line) and lines[i - 1].count('|') >= 2:
            total += 1
    return total


def check_table_fidelity(md_file, html_text, strict=True):
    """Fail if markdown tables did not become real <table> elements.

    The regex fallback has no table support at all, so it renders every row as
    `<p>| a | b |</p>`. That is invisible in a build log but obvious in the
    PDF, which is exactly the kind of silent degradation this check exists to
    turn into a loud error."""
    md_text = Path(md_file).read_text(encoding='utf-8')
    md_tables = count_md_tables(md_text)
    # Raw LaTeX tables count too. Leaving them out is what let five arXiv
    # result tables disappear while this check printed "OK".
    tex_tables = count_raw_latex_tables(md_text)
    want = md_tables + tex_tables
    if want == 0:
        return True

    got = len(re.findall(r'<table\b', html_text, re.IGNORECASE))
    detail = f"{md_tables} markdown"
    if tex_tables:
        detail += f" + {tex_tables} raw-LaTeX"
    if got >= want:
        print(f"Table check: {got} <table> for {want} table(s) ({detail}) — OK")
        return True

    stray = len(re.findall(r'<p>\s*\|', html_text))
    print(f"ERROR: table fidelity check failed — output.md has {want} "
          f"table(s) ({detail}) but the HTML has only {got} <table> element(s)"
          + (f" ({stray} paragraph(s) start with a literal '|', i.e. tables "
             f"were rendered as plain text)" if stray else ""))
    print("  cause: the markdown->HTML converter does not support tables "
          "(regex fallback), the table syntax in output.md is malformed, or a "
          "raw LaTeX tabular could not be converted (pandoc's markdown reader "
          "DROPS raw LaTeX blocks on the HTML path, it does not warn).")
    return not strict


_MATH_ELEMENT_RE = re.compile(r'<math\b.*?</math>', re.DOTALL)
_EMPHASIS_TAG_RE = re.compile(r'</?(?:em|i|strong|b)\b[^>]*>')


def find_spliced_math(html_text):
    """Formulas with markdown emphasis run through the middle of them.

    Every other check here counts: spans in, spans out. This one is the
    reason that is not enough. The count was right -- 225 formulas asked for,
    225 delivered -- while five of them had an `<em>` opened inside one and
    closed inside another, because a literal `*` in the MathML paired with
    the next `*` further down the table. The reader got `45.6^{}` where the
    paper prints `45.6*`, and no total anywhere disagreed.
    """
    return [m for m in _MATH_ELEMENT_RE.findall(html_text)
            if _EMPHASIS_TAG_RE.search(m)]


# MathML keeps the formula's ORIGINAL TeX in an annotation element. When a
# display nests math inside a text argument — Maynard's
# `\text{for infinitely many $n$ all of $n+h_1$, ...}` — that annotation
# contains real `$...$` pairs, and a leftover count taken over the whole page
# reads them back as unrendered math. Maynard reported four while the formula
# was on the page, correctly typeset, every glyph present.
#
# The cost is not the noise. A genuine leak produces the same number and looks
# identical, so the check could no longer tell "a formula printed as source"
# from "a formula rendered and described". Excluding the annotation restores
# the distinction; nothing that renders is counted.
_MATHML_ANNOTATION_RE = re.compile(r'<annotation\b[^>]*>.*?</annotation>',
                                   re.DOTALL)


def check_math_fidelity(md_file, html_text):
    """Warn/fail if math in output.md did not render in the HTML."""
    md_text = Path(md_file).read_text(encoding='utf-8')
    want = len(_MATH_SPAN_RE.findall(md_text))
    if want == 0:
        print("Math check: no $...$ math in output.md — skipped")
        return True

    got = (len(re.findall(r'<math\b', html_text))
           + len(re.findall(r'class="[^"]*\bmath\b', html_text)))
    leftover = len(_MATH_SPAN_RE.findall(
        _MATHML_ANNOTATION_RE.sub(' ', html_text)))
    print(f"Math check: {want} TeX span(s) in md -> {got} rendered, "
          f"{leftover} raw $ span(s) remaining")
    if got == 0:
        print("ERROR: math present in output.md but nothing rendered in the HTML.")
        print("  cause: the converter dropped math, or +tex_math_dollars is missing.")
        return False
    if leftover:
        print(f"WARNING: {leftover} TeX span(s) reached the HTML unrendered")

    spliced = find_spliced_math(html_text)
    if spliced:
        print(f"ERROR: {len(spliced)} formula(s) have markdown emphasis "
              f"spliced through them")
        for blob in spliced[:3]:
            ann = re.search(r'<annotation[^>]*>(.*?)</annotation>', blob,
                            re.DOTALL)
            print("  %s" % ' '.join((ann.group(1) if ann else blob).split())[:90])
        print("  cause: markdown was parsed inside the raw HTML tables, so a "
              "literal * or _ in one formula paired with the next one and the "
              "character it consumed is gone from the page.")
        return False
    return True


def check_image_refs_resolve(temp_dir):
    """Verify every image reference in output.md points at a real file."""
    md_path = os.path.join(temp_dir, 'output.md')
    if not os.path.exists(md_path):
        return True
    md = Path(md_path).read_text(encoding='utf-8')
    html_srcs, md_srcs, _ = _scan_image_refs(md)
    refs = set(html_srcs) | set(md_srcs)

    missing, external = [], []
    for ref in sorted(refs):
        if ref.startswith(('http://', 'https://', 'data:', '//')):
            external.append(ref)
            continue
        clean = ref.split('#')[0].split('?')[0]
        if not os.path.isfile(os.path.normpath(os.path.join(temp_dir, clean))):
            missing.append(ref)

    print(f"Image check: {len(refs)} unique ref(s), {len(missing)} missing, "
          f"{len(external)} external")
    if external:
        print(f"  WARNING: external image refs will not render offline: {external[:5]}")
    if missing:
        print("ERROR: image references in output.md do not resolve under the temp dir:")
        for m in missing[:20]:
            print(f"  - {m}")
        print("  fix: re-run convert.py to re-extract images, or correct the paths "
              "in output.md / the offending output_chunk*.md")
        return False
    return True


_TABLE_EL_RE = re.compile(r'<table(?:\s[^>]*)?>.*?</table>', re.DOTALL)
_TABULAR_IN_FLOAT_RE = re.compile(
    r'\\begin\{(tabular\*?|tabularx|longtable|array)\}.*?\\end\{\1\}',
    re.DOTALL)


def table_structures(temp_dir):
    """[{'header': n, 'rules': {...}}] per table, in the order they appear.

    The rules a paper draws live in flat.tex and nowhere else. Whether a
    given table reaches the page as raw LaTeX or as a pandoc table is decided
    by what pandoc happened to be able to convert -- half of AlphaQ's tables
    changed sides once the ingest stopped fighting it -- and a converted one
    arrives as pipe syntax, which carries no rule at all. Reading the
    structure from the source and applying it by table ORDER covers both, and
    the order is the same one the caption numbering uses, which source_probe
    already checks against the original PDF.
    """
    flat = os.path.join(temp_dir or '', 'flat.tex')
    if not temp_dir or not os.path.exists(flat):
        return []
    try:
        with open(flat, 'r', encoding='utf-8', errors='replace') as fh:
            tex = strip_tex_comments(fh.read())
    except OSError:
        return []
    out = []
    for unit in float_units(tex):
        if unit['kind'] != 'table' or unit['number'] is None:
            continue
        found = _TABULAR_IN_FLOAT_RE.search(tex, unit['start'], unit['stop'])
        if not found:
            out.append({'header': 0, 'rules': {}})
            continue
        tab = found.group(0)
        out.append({'header': header_row_count(tab),
                    'rules': body_rule_rows(tab)})
    return out


def apply_table_structure(html, temp_dir):
    """Give every table its header and its group rules. (html, count)."""
    plans = table_structures(temp_dir)
    if not plans:
        return html, 0
    # Plans are matched to tables BY POSITION, which is only meaningful while
    # both counts agree. If one table failed to render, every table after it
    # would quietly receive the previous one's header depth and rule rows --
    # a wrong answer that looks exactly like a right one.
    found = len(_TABLE_EL_RE.findall(html))
    if found != len(plans):
        sys.stderr.write(
            'warning: %d table(s) in the source, %d in the HTML; header rows '
            'and group rules are matched by position, so this book needs a '
            'look before its tables are trusted\n' % (len(plans), found))
    index, applied = [0], [0]

    def fix(m):
        table = m.group(0)
        i = index[0]
        index[0] += 1
        if i >= len(plans):
            return table
        plan = plans[i]
        before = table
        if plan['header']:
            table = _promote_html_header(table, plan['header'])
        if plan['rules']:
            table = _mark_html_rules(table, plan['rules'])
        if table != before:
            applied[0] += 1
        return table

    return _TABLE_EL_RE.sub(fix, html), applied[0]


def _promote_html_header(table, wanted):
    if '<thead' in table:
        return table
    body = _TBODY_RE.search(table)
    if not body:
        return table
    rows = _TR_RE.findall(body.group(1))
    if not 0 < wanted < len(rows):
        return table
    head = '\n'.join(re.sub(r'<(/?)td\b', r'<\1th', r) for r in rows[:wanted])
    rest = '\n'.join(rows[wanted:])
    return (table[:body.start()]
            + '<thead>\n%s\n</thead>\n<tbody>\n%s\n</tbody>' % (head, rest)
            + table[body.end():])


def _mark_html_rules(table, rules):
    body = _TBODY_RE.search(table)
    if not body:
        return table
    rows = _TR_RE.findall(body.group(1))
    out = []
    for i, row in enumerate(rows):
        kind = rules.get(i)
        if kind and 'rule-above' not in row:
            css = 'rule-above' if kind == 'hard' else 'rule-above-soft'
            row = re.sub(r'<tr\b', '<tr class="%s"' % css, row, count=1)
        out.append(row)
    return (table[:body.start()] + '<tbody>\n' + '\n'.join(out) + '\n</tbody>'
            + table[body.end():])


_SUMMARY_ROW_RE = re.compile(
    r'^\s*(?:<[^>]+>|\s)*(?:평균|합계|전체|총계|平均|合计|合計|总计|總計'
      r'|全体|全體|Average|Avg\.?|Total|Mean'
    r'|Overall)\b', re.IGNORECASE)
_CELL_ONE_RE = re.compile(r'<t[hd][^>]*>(.*?)</t[hd]>', re.DOTALL)


def rule_off_summary_rows(html):
    """Rule off the Average/Total block at the foot of a long table.

    A paper ends a long results table with summary rows and separates them
    with a `\\midrule`. Tables pandoc could convert take the markdown path,
    where pipe syntax has no way to carry a rule, so the boundary is simply
    gone -- CafeQ's twenty-row per-task table ran its four averages straight
    on from its sixteen benchmarks.

    Narrow on purpose: only a table that has no group rule already, only a
    summary row in the last third, only the first such row. Measured across
    three papers it fires once, exactly where the source has its rule.
    """
    def fix(m):
        table = m.group(0)
        if 'rule-above' in table:
            return table
        body = _TBODY_RE.search(table)
        if not body:
            return table
        rows = _TR_RE.findall(body.group(1))
        if len(rows) < 6:
            return table
        for i, row in enumerate(rows):
            if i < len(rows) * 2 / 3:
                continue
            cell = _CELL_ONE_RE.search(row)
            if not cell or not _SUMMARY_ROW_RE.match(cell.group(1)):
                continue
            rows[i] = re.sub(r'<tr\b', '<tr class="rule-above"', row, count=1)
            inner = '\n'.join(rows)
            return (table[:body.start()] + '<tbody>\n' + inner + '\n</tbody>'
                    + table[body.end():])
        return table

    return _TABLE_EL_RE.sub(fix, html)


def finish_table_rules(html_file, temp_dir):
    """Give every table the rules the paper drew, whichever path it took.

    Source first, guess second: a table whose structure is still readable in
    flat.tex gets exactly that, and only one with nothing left to read falls
    through to the summary-row heuristic.
    """
    try:
        with open(html_file, 'r', encoding='utf-8') as fh:
            content = fh.read()
    except OSError as exc:
        print(f"Error reading HTML for table rules: {exc}")
        return
    content, applied = apply_table_structure(content, temp_dir)
    content = rule_off_summary_rows(content)
    try:
        with open(html_file, 'w', encoding='utf-8') as fh:
            fh.write(content)
    except OSError as exc:
        print(f"Error writing table rules: {exc}")
        return
    if applied:
        print(f"Tables: {applied} given the header and group rules the source "
              f"draws")



# Elements whose text must never be rewritten: markup that is not prose, and
# the captions themselves -- "그림 1 (Fig. 1)" inside a figcaption would
# otherwise become a link to the figure it already labels.
_XREF_SKIP = frozenset((
    'style', 'script', 'math', 'code', 'pre', 'a', 'figcaption', 'caption',
    'head', 'title', 'nav', 'textarea',
))
_XREF_TAG_RE = re.compile(r'<!--.*?-->|<(/?)([a-zA-Z][\w:-]*)([^>]*)>', re.DOTALL)
_FIGURE_BLOCK_RE = re.compile(r'<figure\b[^>]*>.*?</figure>', re.DOTALL)
# `(\d+)` alone cannot see a section-scoped number, and the failure is silent
# in the K80 way: the anchor is simply never created, "식 (2.2)" links to an id
# nothing carries, and every cross-reference in the book points at nothing
# without one error or warning. Dots are safe in an id — only `href="#..."`
# selects these, never CSS or JS.
_MATH_EQNO_RE = re.compile(
    r'<math\b(?![^>]*\bid=)([^>]*\bdata-eqno="\((\d+(?:\.\d+)*)\)")')
_HAS_ID_RE = re.compile(r'\bid\s*=')
# An author-year citation: opens on a capital, closes on a year. Bracketed
# citations are deliberately not matched -- across these three books the
# pattern found two, and both were ordinary prose.
_CITE_TEXT_RE = re.compile(r'\((?=[A-Z])[^()]{2,200}?(?:19|20)\d{2}[a-z]?\)')


def _walk_text_nodes(html, transform):
    """Apply `transform` to prose only, never to markup or to a caption."""
    out, pos, open_skips = [], 0, {}

    def prose():
        return not any(open_skips.values())

    for match in _XREF_TAG_RE.finditer(html):
        chunk = html[pos:match.start()]
        if chunk:
            out.append(transform(chunk) if prose() else chunk)
        out.append(match.group(0))
        pos = match.end()
        name = (match.group(2) or '').lower()
        if name in _XREF_SKIP:
            if match.group(1):
                open_skips[name] = max(0, open_skips.get(name, 0) - 1)
            elif not (match.group(3) or '').rstrip().endswith('/'):
                open_skips[name] = open_skips.get(name, 0) + 1
    tail = html[pos:]
    if tail:
        out.append(transform(tail) if prose() else tail)
    return ''.join(out)


def _add_id(tag_text, anchor):
    """Put an id on an opening tag that has none."""
    if _HAS_ID_RE.search(tag_text):
        return tag_text
    return tag_text[:tag_text.index('>')].rstrip() \
        + ' id="%s"' % anchor + tag_text[tag_text.index('>'):]


_EM_RE = re.compile(r'(?s)<(em|i)((?:\s[^>]*)?)>(.*?)</\1>')
# Hangul, kana and Han. None of the three has a real italic in the faces this
# pipeline can rely on, and Chromium answers a request for one by synthesising
# an oblique -- which for Chinese it then emits as a Type3 object, one per
# glyph. Was Hangul only, and Chinese emphasis went on producing Type3 long
# after the Korean case was solved.
_CJK_TEXT_RE = re.compile(r'[가-힣぀-ヿ一-鿿]')


def mark_cjk_emphasis(html):
    r"""Tag emphasis that actually contains CJK. Returns (html, marked).

    CJK has no italic, so the print sheet renders such emphasis bold instead
    of letting Chromium synthesise an oblique. That rule was written as
    `:lang(ko) em` -- and the root element is `lang="ko"`, so it matched
    EVERY <em> in the book. `\textit{16.67}` and `\textit{Wiki2}` printed
    bold, inside tables whose caption says the best result is the bold one:
    ten of SINQ's tables showed their FP16 baseline row as the winner. Which
    is why this looks at what the element CONTAINS and never at the document
    language.
    """
    marked = [0]

    def sub(m):
        tag, attrs, body = m.group(1), m.group(2), m.group(3)
        text = re.sub(r'<[^>]+>', '', body)
        if not _CJK_TEXT_RE.search(text):
            return m.group(0)
        marked[0] += 1
        if re.search(r'\bclass="', attrs):
            attrs = re.sub(r'\bclass="([^"]*)"', r'class="\1 cjk"', attrs, 1)
        else:
            attrs += ' class="cjk"'
        return '<%s%s>%s</%s>' % (tag, attrs, body, tag)

    return _EM_RE.sub(sub, html), marked[0]


def anchor_reference_targets(html, lang_cfg=None):
    """Give each figure, table and numbered equation an id. (html, targets)."""
    lang_cfg = lang_cfg or {}
    fig_label = lang_cfg.get('figure_label', 'Figure')
    tab_label = lang_cfg.get('table_label', 'Table')
    targets = set()

    def number_in(block, label, opener):
        # A window after the opening tag, NOT the whole element: one AlphaQ
        # caption is 24 KB of inline MathML, and requiring the closing tag
        # made three figures match nothing at all. The label is always at the
        # very front of a caption.
        start = re.search(r'<%s\b[^>]*>' % opener, block)
        if not start:
            return None
        window = re.sub(r'<[^>]+>', ' ', block[start.end():start.end() + 400])
        hit = re.search(re.escape(label) + r'\s*(\d+)', window)
        return hit.group(1) if hit else None

    def anchored(block, anchor, inner):
        """Put `anchor` on the element, or on its caption if it has an id.

        A float that carried a `\\label{}` in the source arrives with that
        label AS its id, and an element can hold only one. `_add_id` then
        returns the tag untouched, no `tab-N` is created, and every reference
        rewritten to `#tab-N` points at nothing — a dead in-page link, which
        neither errors nor prints. Six of CafeQ's eight tables were in that
        state in both builds. The caption is inside the element and is where
        a reader following the link wants to land anyway.
        """
        targets.add(anchor)
        head = block[:block.index('>') + 1]
        rest = block[block.index('>') + 1:]
        if not _HAS_ID_RE.search(head):
            return _add_id(head, anchor) + rest
        cap = re.search(r'<%s\b[^>]*>' % inner, rest)
        if cap and not _HAS_ID_RE.search(cap.group(0)):
            return head + rest[:cap.start()] + _add_id(cap.group(0), anchor) \
                + rest[cap.end():]
        # Both already spoken for: keep the label id and let the reference
        # resolve to it, rather than inventing a target that is not there.
        existing = re.search(r'\bid="([^"]+)"', head)
        if existing:
            targets.discard(anchor)
            targets.add(existing.group(1))
        return block

    def fig_sub(match):
        block = match.group(0)
        number = number_in(block, fig_label, 'figcaption')
        if number is None:
            return block
        return anchored(block, 'fig-%s' % number, 'figcaption')

    def tab_sub(match):
        block = match.group(0)
        number = number_in(block, tab_label, 'caption')
        if number is None:
            return block
        return anchored(block, 'tab-%s' % number, 'caption')

    def eq_sub(match):
        targets.add('eq-%s' % match.group(2))
        return '<math id="eq-%s"%s' % (match.group(2), match.group(1))

    html = _FIGURE_BLOCK_RE.sub(fig_sub, html)
    html = _TABLE_EL_RE.sub(tab_sub, html)
    html = _MATH_EQNO_RE.sub(eq_sub, html)
    return html, targets


def link_cross_references(html, lang_cfg=None):
    """Colour every reference; link the ones whose target is unambiguous.

    Takes and returns the body HTML rather than a file, because it has to run
    AFTER the equations are numbered, and by then the body is a string on its
    way into both templates.
    """
    lang_cfg = lang_cfg or {}
    html, targets = anchor_reference_targets(html, lang_cfg)

    fig = re.escape(lang_cfg.get('figure_label', 'Figure'))
    tab = re.escape(lang_cfg.get('table_label', 'Table'))
    eqn = re.escape(lang_cfg.get('equation_label', 'Equation'))
    app = re.escape(lang_cfg.get('appendix_label', 'Appendix'))
    stats = {'linked': 0, 'coloured': 0}

    linkable = [
        # Dotted, for the same reason as _MATH_EQNO_RE above: a paper that
        # numbers per section prints `그림 3.1`, and matching only `3` would
        # link the reference to an anchor that does not exist.
        (re.compile(fig + r'\s*(\d+(?:\.\d+)*)'), 'fig'),
        (re.compile(tab + r'\s*(\d+(?:\.\d+)*)'), 'tab'),
        (re.compile(eqn + r'\s*\((\d+(?:\.\d+)*)\)'), 'eq'),
    ]
    plain = [
        re.compile(app + r'\s*[A-Z](?:\.\d+)*'),
        _CITE_TEXT_RE,
    ]

    def link_sub(prefix):
        def sub(match):
            anchor = '%s-%s' % (prefix, match.group(1))
            if anchor not in targets:
                return match.group(0)
            stats['linked'] += 1
            return '<a class="xref" href="#%s">%s</a>' % (anchor, match.group(0))
        return sub

    def colour_sub(match):
        stats['coloured'] += 1
        return '<span class="xref">%s</span>' % match.group(0)

    def transform(text):
        if '<' in text or '>' in text:
            return text                # never rewrite a fragment of markup
        for pattern in plain:
            text = pattern.sub(colour_sub, text)
        for pattern, prefix in linkable:
            text = pattern.sub(link_sub(prefix), text)
        return text

    return _walk_text_nodes(html, transform), stats

def process_html_separators(html_file):
    """Process page separators in HTML"""
    try:
        with open(html_file, 'r', encoding='utf-8') as f:
            content = f.read()

        content = re.sub(r'<hr\s*/?>', '<div class="page-separator"></div>', content)
        content = re.sub(r'<p>\s*---\s*</p>', '<div class="page-separator"></div>', content)

        with open(html_file, 'w', encoding='utf-8') as f:
            f.write(content)
    except Exception as e:
        print(f"Error processing separators: {e}")


_CAPTION_BODY_RE = re.compile(r'\\caption\s*(?:\[[^\]]*\])?\s*\{')


def source_captions(temp_dir):
    r"""Every `\caption{}` body in `flat.tex`, whitespace collapsed.

    `flat.tex` is the flattened LaTeX the whole book was built from, and it is
    the one file the table agents never touch: the sidecars they edit are
    copies. So it is the only pristine record of what a caption said before
    anybody translated it, which is what makes a caption still identical to
    its source detectable in ANY target language, not just the ones written
    in another script.
    """
    path = os.path.join(temp_dir or '', 'flat.tex')
    if not temp_dir or not os.path.isfile(path):
        return set()
    try:
        with open(path, encoding='utf-8', errors='replace') as handle:
            text = handle.read()
    except (IOError, OSError):
        return set()
    out = set()
    for match in _CAPTION_BODY_RE.finditer(text):
        start = match.end() - 1
        depth, i = 1, start + 1
        while i < len(text) and depth:
            depth += (text[i] == '{') - (text[i] == '}')
            i += 1
        body = ' '.join(text[start + 1:i - 1].split())
        if len(body) >= 12:
            out.add(body)
    return out


_CAPTION_WORD_RE = re.compile(r'[^\W\d_]+', re.UNICODE)
_CAPTION_COMMENT_RE = re.compile(r'(?<!\\)%.*')
_CAPTION_KEYED_RE = re.compile(
    r'\\(?:cite[a-z]*|label|ref|eqref|nameref)\s*\{[^{}]*\}')
_CAPTION_MATH_RE = re.compile(r'\$[^$]*\$')
_CAPTION_COMMAND_RE = re.compile(r'\\[A-Za-z]+')

# Measured over 73 correctly translated captions from six papers and eight
# books: the longest run any of them shares with its source is THREE, and the
# shortest run in a caption nobody translated is SIX. The threshold sits in
# that gap. Raising it costs short untranslated fragments; lowering it below
# four would have fired on VLA-Adapter's model-name captions.
_UNTRANSLATED_RUN = 4


def caption_prose(text):
    r"""A caption with its LaTeX removed, so only words a reader sees remain.

    Every part taken out here is identical in the source and in a correct
    translation, by design, and each one was measured making a correctly
    translated caption look untranslated:

      * `% ...` comments. SINQ's captions keep the paper's own commented-out
        English wording below the Korean, and the reader never sees it. That
        alone scored a run of 28.
      * `\citep{OpenVLA-2024}`, `\label{}`, `\ref{}`: the key must not change.
      * `$...$` maths.
      * command names. `\textbf`, `\texttt`, `\mbox` are structure, not words.
    """
    text = _CAPTION_COMMENT_RE.sub(' ', text)
    text = _CAPTION_KEYED_RE.sub(' ', text)
    text = _CAPTION_MATH_RE.sub(' ', text)
    text = _CAPTION_COMMAND_RE.sub(' ', text)
    return text.replace('{', ' ').replace('}', ' ')


def caption_words(text):
    """Words of the caption's prose, original case kept: capitalisation is
    what separates a model name from a function word."""
    return _CAPTION_WORD_RE.findall(caption_prose(text))


def _is_prose_word(token):
    r"""Lowercase and alphabetic: the shape a function word has.

    What a correct translation shares with its source is names -- OpenVLA,
    DeepSeek-V2-Lite, Qwen1.5-MoE, CALVIN, LIBERO-Long -- and every one of
    them is capitalised or carries a digit. Untranslated prose brings
    lowercase words along and cannot avoid them. Frequency was tried first
    and is worse: the commonest short words of an ML paper include MoE, so a
    run of model names passes a frequency test.
    """
    return len(token) >= 2 and token.isalpha() and token.islower()


def longest_source_run(caption, originals):
    r"""The longest run of PROSE words this caption shares with its source.

    A translator borrows single words from the source: an acronym, a dataset
    name, a cognate. Text nobody translated arrives as a contiguous run. The
    length of the longest run tells the two apart in any target language,
    because the run is made of source words whatever the target is written
    in. A run of names does not count, or a Korean caption naming four models
    in a row would be reported as untranslated.
    """
    have = caption_words(caption)
    if not have:
        return 0
    best = 0
    for original in originals:
        want = [w.lower() for w in caption_words(original)]
        row = [0] * (len(want) + 1)
        for i in range(1, len(have) + 1):
            prev = 0
            for j in range(1, len(want) + 1):
                keep = row[j]
                row[j] = prev + 1 if have[i - 1].lower() == want[j - 1] else 0
                if row[j] > best and any(_is_prose_word(w)
                                         for w in have[i - row[j]:i]):
                    best = row[j]
                prev = keep
    return best


def translation_is_passthrough(temp_dir):
    r"""Did this run copy its chunks through instead of translating them?

    An English edition of an English paper is the pipeline's honest answer
    when there is nothing to translate, and its captions are identical to
    `flat.tex` because that is CORRECT, not because a step was skipped. No
    caption check can tell those apart, so it must not try.

    Asking `lang == 'en'` would be the wrong question twice over: it assumes
    every source paper is English, so a French paper rendered into English
    would silently lose its caption check, and an English paper rendered into
    English is recognisable without guessing at either language. Ask the
    artefact instead. If the translated chunks are the source chunks, nothing
    was translated anywhere, and the captions are not evidence of anything.
    """
    if not temp_dir or not os.path.isdir(temp_dir):
        return False
    names = sorted(n for n in os.listdir(temp_dir)
                   if re.match(r'^chunk\d+\.md$', n))
    checked = 0
    for name in names:
        output = os.path.join(temp_dir, 'output_' + name)
        if not os.path.isfile(output):
            continue
        try:
            with open(os.path.join(temp_dir, name), encoding='utf-8',
                      errors='replace') as fh:
                source = fh.read()
            with open(output, encoding='utf-8', errors='replace') as fh:
                translated = fh.read()
        except (IOError, OSError):
            return False
        if ' '.join(source.split()) != ' '.join(translated.split()):
            return False
        checked += 1
        if checked >= 3:
            break
    return checked > 0


def untranslated_captions(md_text, lang, temp_dir=None):
    r"""Table captions still in the source language. [] when unmeasurable.

    Step 4.6 of SKILL.md translates the words inside table floats, because a
    float is protected behind a `⟦T####⟧` placeholder and no translator ever
    sees its `\caption{}`. That step is prose, and prose gets skipped: it was
    skipped for three editions of one paper in a single session, twice after
    being raised.

    Nothing noticed, and the step's own text says why -- the book comes out
    with its tables in the source language and EVERY existing check passes,
    because they count tables, images and values and those are all correct.
    A green run actively confirms the wrong conclusion.

    So the build asks the artefact instead of trusting that the step ran,
    three ways, because no one of them covers every way the step can be half
    done:

      * the caption is not in the target's script at all. Decisive for ko,
        ja and zh, and blind to fr, de and es.
      * the caption is still word for word what `flat.tex` says. Works in any
        target language, and is the only thing that catches a caption of two
        or three words nobody touched.
      * the caption carries a run of four or more consecutive source words.
        This is what catches a caption translated HALFWAY, which the other
        two both wave through: one target-script character satisfies the
        first, and a half-translated caption is not identical, so it
        satisfies the second.

    A run that copied its chunks through instead of translating them, which
    is the honest rendering of an English paper into English, cannot be
    judged by the last two: a caption identical to the source is correct
    there and indistinguishable from a skipped step. That is recognised from
    the chunks rather than from the language name (K68).
    """
    base = (lang or '').split('-')[0]
    try:
        import verify_chunk
        ranges = verify_chunk._SCRIPT_RANGES.get(base)
    except Exception:                                     # noqa: BLE001
        ranges = None
    # A run that copied its chunks through translated nothing anywhere, so a
    # caption that is not in the target's script is not evidence that step
    # 4.6 was skipped: there was no step 4.6 to skip. The abstention was
    # first applied only to `originals`, which left the SCRIPT test running,
    # and that broke the dry run of SKILL.md 2.5 -- a build with nothing
    # translated is exactly what that step is for, and the gate refused it
    # for every paper holding a table. Found by putting a new paper through
    # the pipeline, not by any check.
    if translation_is_passthrough(temp_dir):
        return []
    originals = source_captions(temp_dir)
    if not ranges and not originals:
        return []

    out = []
    for table in find_raw_latex_tables(md_text):
        caption = (table.get('caption') or '').strip()
        # A very short caption is a label, not prose, and says nothing about
        # whether anybody translated it.
        if len(caption) < 12:
            continue
        flat = ' '.join(caption.split())
        wrong_script = ranges and not any(
            verify_chunk._in_target_script(ch, ranges) for ch in caption)
        half_done = (originals
                     and longest_source_run(caption, originals)
                     >= _UNTRANSLATED_RUN)
        if wrong_script or flat in originals or half_done:
            out.append(flat[:70])
    return out


def convert_md_to_html(temp_dir, title, lang_cfg, author=None,
                       allow_degraded=False, math_mode='mathml', force=False,
                       print_cfg=None):
    """Convert output.md to HTML with templates"""
    print("=== Converting markdown to HTML ===")

    md_file = os.path.join(temp_dir, 'output.md')
    if not os.path.exists(md_file):
        print("Error: output.md not found.")
        return False

    with open(md_file, encoding='utf-8', errors='replace') as fh:
        _merged = fh.read()
    stale_captions = untranslated_captions(_merged,
                                           lang_cfg.get('lang_attr', ''),
                                           temp_dir)
    if stale_captions:
        print("ERROR: %d table caption(s) are still in the source language."
              % len(stale_captions))
        for line in stale_captions[:4]:
            print("  - %s" % line)
        print("  A table float sits behind a placeholder, so no translator "
              "saw its \\caption{}. Every other check passes: the tables, "
              "the values and the counts are all correct.")
        print("  SKILL.md step 4.6 translates them. Start with:")
        print("      python tests/format_probe.py \"%s\" --lang %s"
              % (temp_dir, (lang_cfg.get('lang_attr') or '').split('-')[0]))
        raise SystemExit(1)

    book_doc_file = os.path.join(temp_dir, 'book_doc.html')

    # Skip HTML generation if book_doc.html exists and is newer than output.md
    if os.path.exists(book_doc_file) and not force:
        if os.path.getmtime(book_doc_file) > os.path.getmtime(md_file):
            if _check_generated_html_sanity(book_doc_file):
                print("Skipping HTML generation - book_doc.html is up to date")
                return True
            print("Stale book_doc.html failed image sanity — regenerating")
            os.remove(book_doc_file)
        else:
            print("Re-generating HTML - output.md is newer")

    temp_html_file = os.path.join(temp_dir, 'output.html')

    # Raw LaTeX tables must become HTML before ANY markdown->HTML converter
    # sees them: pandoc's markdown reader parses them as raw LaTeX blocks, and
    # a raw block only survives into its own output format, so on the HTML path
    # they vanish without a warning. output.md itself is left untouched -- it
    # stays the faithful merged translation -- and the converters read the
    # expanded copy instead.
    source_md = md_file
    md_text = Path(md_file).read_text(encoding='utf-8')

    # Citations and cross-references first: they are plain-text substitutions
    # and must happen before any table becomes HTML.
    md_text, unwrapped = unwrap_prose_environments(md_text)
    if unwrapped:
        print(f"Prose environments: {unwrapped} rewritten as markdown; pandoc "
              f"takes each as one raw block and drops it whole")

    md_text, cells = unwrap_table_cell_wrappers(md_text)
    if cells:
        print("Tables: %d \\makecell/\\thead cell(s) unwrapped; pandoc drops "
              "the command WITH its text and leaves the cell empty" % cells)

    md_text, tabbed = unwrap_tabbing(
        md_text, read_one_argument_macros(temp_dir))
    if tabbed:
        print("Pseudocode: %d tabbing environment(s) rewritten as code blocks; "
              "pandoc has no reader for tabbing and drops it silently" % tabbed)

    md_text, leftover_stats = normalize_latex_leftovers(md_text)
    if any(leftover_stats.values()):
        print("LaTeX leftovers: %d drop-cap(s) restored, %d stray command(s) "
              "removed, %d escaped bracket(s) un-mathed"
              % (leftover_stats['parstart'], leftover_stats['dropped'],
                 leftover_stats['brackets']))
    if leftover_stats.get('index_terms'):
        # A reduction, not a repair: the paper has an index and this book will
        # not. Said out loud so it is not one more silent loss (K110).
        print("Index: %d \\index term(s) dropped — the original builds an "
              "index, this book has none" % leftover_stats['index_terms'])

    md_text, nested_hits = split_nested_math_text(md_text)
    if nested_hits:
        print("Math: %d nested $...$ inside a text-mode argument split out; "
              "every $-pairing scanner downstream mis-closes on those"
              % nested_hits)

    md_text, math_refs = resolve_math_references(md_text, temp_dir)
    if math_refs:
        print("Math: %d \\ref(s) inside formulas resolved; texmath has no "
              "reader for one and refuses the whole display" % math_refs)

    md_text, macro_hits = expand_math_macros(md_text, temp_dir)
    if macro_hits:
        print("Math: %d formula(s) had the paper's own shorthand expanded "
              "(\\< -> \\langle); texmath knows only standard commands"
              % macro_hits)

    md_text, text_fonts = rewrite_text_fonts_in_math(md_text)
    if text_fonts:
        print(f"Math: {text_fonts} text-mode font switch(es) inside formulas "
              f"rewritten (\\textsc -> \\mathrm); texmath refuses the whole "
              f"formula over one")

    md_text, math_stats = normalize_math_commands(md_text)
    if math_stats['fonts']:
        print("Math: %d legacy font switch(es) rewritten (\\rm -> \\mathrm)"
              % math_stats['fonts'])
    if math_stats['accents']:
        print("Math: %d accent argument(s) braced "
              "(\\widetilde\\mathbf{X} -> \\widetilde{\\mathbf{X}}); texmath "
              "drops the whole formula over one"
              % math_stats['accents'])

    md_text, unboxed = unwrap_text_boxed_math_fonts(md_text)
    if unboxed:
        print("Math: %d math font(s) unwrapped from a text box "
              "(\\text{\\mathtt{x}} -> \\mathtt{x}); texmath refuses the "
              "nesting" % unboxed)

    md_text, ref_stats = resolve_references(md_text, temp_dir, lang_cfg)
    if ref_stats.get('subrefs'):
        print("Sub-figure references: %d resolved to panel letters"
              % ref_stats['subrefs'])

    # Restore the original's section numbering, so a reader can map a
    # translated heading back to the paper they are holding.
    md_text, sec_stats = number_sections(md_text, temp_dir)
    if sec_stats['numbered']:
        pdf = sec_stats.get('pdf') or {}
        detail = ''
        if pdf:
            bits = ['%d numbered' % pdf.get('matched', 0)]
            if pdf.get('unnumbered'):
                bits.append('%d unnumbered in the original' % pdf['unnumbered'])
            if pdf.get('wrapped'):
                bits.append('%d matched across a line break' % pdf['wrapped'])
            if pdf.get('missing'):
                bits.append('%d not found' % pdf['missing'])
            detail = ' (%s)' % ', '.join(bits)
        print(f"Sections: {sec_stats['numbered']} heading(s) keyed to the "
              f"original{detail}")
    elif sec_stats['skipped_reason']:
        print(f"Sections: not numbered — {sec_stats['skipped_reason']}")

    # Immediately after, because the two halves have to agree: the reference
    # pass above now resolves a theorem label to the number the paper prints,
    # and a declaration line still reading `정리 1` under prose saying
    # `정리 1.1` is worse than either alone (K130).
    md_text, thm_stats = number_theorem_statements(md_text, temp_dir, lang_cfg)
    if thm_stats['numbered'] or thm_stats['unnumbered']:
        bits = ['%d renumbered' % thm_stats['numbered']]
        if thm_stats['unnumbered']:
            bits.append('%d left unnumbered, as the paper prints them'
                        % thm_stats['unnumbered'])
        print('Theorem statements: %s' % ', '.join(bits))
    elif thm_stats['skipped_reason']:
        print('Theorem statements: left as they are — %s'
              % thm_stats['skipped_reason'])

    # After numbering, deliberately: the heading this adds is not in flat.tex,
    # and inserting it first made the ladder come out one too long.
    md_text, bib_stats = resolve_bibliography(md_text, temp_dir, lang_cfg)
    if bib_stats['dropped_duplicate']:
        print("References: dropped the inlined \\thebibliography — the source "
              "also ships a .bib and citeproc already rendered the list")
    if bib_stats['heading_added']:
        print("References: heading added over the rendered reference list")

    md_text, orphan_notes = rescue_orphan_footnotes(md_text)
    if orphan_notes:
        print(f"Footnotes: {orphan_notes} note(s) nothing referenced moved to "
              f"the front matter, where pandoc would have dropped them")

    # Captions become real <figure>/<figcaption> so they stop reading as body.
    md_text, trimmed = apply_graphics_trim(md_text, temp_dir)
    if trimmed:
        print(f"Figures: {trimmed} cropped as the source asked")
    md_text, fig_count = format_figure_blocks(md_text, lang_cfg, temp_dir)
    if fig_count:
        print(f"Figures: {fig_count} caption(s) formatted")

    md_text, grid_count = grid_tables_to_pipe(md_text)
    if grid_count:
        print(f"Tables: {grid_count} grid table(s) rewritten as pipe tables "
              f"(a grid one collapses once its cells change width)")

    md_text, tab_count = number_table_captions(md_text, temp_dir, lang_cfg)
    if tab_count:
        print(f"Tables: {tab_count} caption(s) numbered")
    # Where they landed, not how many were written. The count agreed while
    # ten of fifteen sat on prose, so the count is not the check (K151).
    placed_ok, placed_detail = check_badge_placement(md_text, lang_cfg)
    if not placed_ok:
        print("BLOCKING: %s" % placed_detail)
        print("  A table number on prose sends every sentence that cites it "
              "at the wrong thing, and no count can see it.")
        raise SystemExit(1)
    if tab_count:
        print("Tables: %s" % placed_detail)

    if any(ref_stats.values()):
        print("References: %d citation(s), %d cross-reference(s) resolved"
              % (ref_stats['cites'], ref_stats['xrefs'])
              + ("; %d citation(s) and %d cross-reference(s) had no target"
                 % (ref_stats['cites_missed'], ref_stats['xrefs_missed'])
                 if (ref_stats['cites_missed'] or ref_stats['xrefs_missed']) else ""))

    # Two prepared copies, because the two output paths cannot share one.
    #
    #   prepared.md      -> DOCX. Tables as MARKDOWN: pandoc drops raw HTML
    #                       entirely when writing DOCX, and injecting <table>
    #                       here left book.docx with zero tables and no check
    #                       complaining. Not every LaTeX table survives the
    #                       trip, so the shortfall is reported.
    #   pandoc_input.md  -> HTML/PDF/EPUB. Tables as HTML, which keeps the
    #                       column-count classes the print sheet sizes on.
    #
    # output.md itself is never rewritten; it stays the faithful translation.
    #
    # Algorithm floats are expanded FIRST, and into markdown rather than HTML,
    # so both prepared copies carry them: raw HTML survives only the HTML path,
    # which is the trap the two table copies exist to avoid. Everything the
    # float contains -- `$...$` included -- then travels the ordinary route.
    md_text, algo_ok, algo_bad = algorithm_float.expand_algorithm_floats(
        md_text, lang=lang_cfg.get('lang_attr', 'en'))
    if algo_ok or algo_bad:
        note = f"Algorithm floats: {algo_ok} converted to markdown"
        if algo_bad:
            note += f", {algo_bad} FAILED (they will be missing)"
        print(note)

    # The prompt tells every sub-agent to spell a term out on ITS first use,
    # because each one sees a single chunk and that is the only first use it
    # can know about. A term running through ten chunks therefore arrives
    # glossed ten times. Here the whole book is in one string, so the second
    # and later copies come out. output.md keeps them: it is the faithful
    # record of what each sub-agent wrote.
    md_text, glosses_dropped = glossary.dedupe_glosses(md_text)
    if glosses_dropped:
        print(f"First-use glosses: {glosses_dropped} repeat(s) removed")

    docx_md, docx_ok, docx_bad = expand_raw_latex_tables(
        md_text, math_mode=math_mode, output='markdown', temp_dir=temp_dir)
    # pandoc builds book.docx straight from this file and never sees the HTML
    # the other formats are styled through, so the equation number has to go
    # inside the formula here.
    docx_md, docx_tagged = tag_equations_for_markdown(docx_md, temp_dir)
    prepared = os.path.join(temp_dir, 'prepared.md')
    with open(prepared, 'w', encoding='utf-8', newline='') as fh:
        fh.write(docx_md)
    if docx_ok or docx_bad:
        note = f"Raw LaTeX tables (DOCX): {docx_ok} converted to markdown"
        if docx_bad:
            note += (f", {docx_bad} could not be expressed as a markdown table "
                     f"and stay as text")
        print(note)

    expanded_md, tables_ok, tables_bad = expand_raw_latex_tables(
        md_text, math_mode=math_mode, output='html', temp_dir=temp_dir)
    if tables_ok or tables_bad:
        note = f"Raw LaTeX tables: {tables_ok} converted to HTML"
        if tables_bad:
            note += f", {tables_bad} FAILED (they will be missing)"
        print(note)
    source_md = os.path.join(temp_dir, 'pandoc_input.md') if tables_ok else prepared
    if tables_ok:
        with open(source_md, 'w', encoding='utf-8', newline='') as fh:
            fh.write(expanded_md)

    # Tier 0 pandoc -> tier 1 python-markdown -> (only on request) regex
    success = False
    used = None
    pandoc = resolve_pandoc()
    if pandoc:
        success = convert_with_pandoc(source_md, temp_html_file, title,
                                      lang_cfg['lang_attr'], math_mode=math_mode)
        if success:
            used = 'pandoc'
    else:
        print("WARNING: pandoc not found (checked pypandoc, PATH, and standard "
              "install dirs)")

    if not success and MARKDOWN_AVAILABLE:
        if convert_with_python_markdown(source_md, temp_html_file, title):
            success, used = True, 'python-markdown'

    if not success:
        if not allow_degraded:
            print("\nERROR: no high-fidelity markdown converter available.")
            print(f"  pandoc:          {'found' if pandoc else 'NOT FOUND'}")
            print(f"  python-markdown: {'installed' if MARKDOWN_AVAILABLE else 'NOT INSTALLED'}")
            print("  The regex fallback cannot render tables or math and would")
            print("  silently produce a degraded book. Refusing to continue.")
            print("  fix: install pandoc (https://pandoc.org/installing.html) or")
            print("       run `python -m pip install markdown`,")
            print("  or re-run with --allow-degraded-html to accept table-less output.")
            return False
        print("WARNING: --allow-degraded-html — using the regex fallback. "
              "Tables will be LOST and math may be mangled.")
        if convert_with_basic_regex(source_md, temp_html_file, title):
            success, used = True, 'basic-regex(DEGRADED)'

    if not success:
        print("Error: All markdown-to-HTML converters failed")
        return False
    print(f"HTML converter used: {used}")

    process_html_separators(temp_html_file)
    finish_table_rules(temp_html_file, temp_dir)

    # Extract body content
    try:
        with open(temp_html_file, 'r', encoding='utf-8') as f:
            html_content = f.read()
    except Exception as e:
        print(f"Error reading HTML file: {e}")
        return False

    body_match = re.search(r'<body[^>]*>(.*?)</body>', html_content, re.DOTALL | re.IGNORECASE)
    body_content = body_match.group(1).strip() if body_match else html_content

    # Mark the numbered equations. The stylesheet sets the label flush right,
    # so nothing is added to the formula text itself and a copy-paste of the
    # equation stays clean.
    body_content, eq_tagged = tag_equations_in_html(body_content, md_text,
                                                    temp_dir)
    if eq_tagged or docx_tagged:
        print("Equations: %d numbered (%d in the DOCX copy)"
              % (eq_tagged, docx_tagged))

    body_content, cjk_em = mark_cjk_emphasis(body_content)
    if cjk_em:
        print("Emphasis: %d CJK run(s) marked for the bold substitute "
              "(Latin emphasis keeps real italics)" % cjk_em)

    # After the equations are numbered, so "식 (7)" has something to point at.
    body_content, xrefs = link_cross_references(body_content, lang_cfg)
    if xrefs['linked'] or xrefs['coloured']:
        print("References: %d linked, %d coloured"
              % (xrefs['linked'], xrefs['coloured']))

    if not check_table_fidelity(md_file, body_content, strict=not allow_degraded):
        return False
    if not check_math_fidelity(md_file, body_content):
        return False

    # Every check above counts something: tables, images, equations, captions.
    # A float in an environment nobody wrote a handler for is not counted by
    # any of them, so it can be deleted between output.md and the HTML while
    # the build reports success. This one asks the opposite question -- does
    # each raw-LaTeX block still have prose on the page? -- and needs no list
    # of environments to be kept up to date.
    # Fingerprint the text we HANDED to pandoc, not the text we started with.
    # `md_text` has already had its algorithm floats expanded and its repeated
    # glosses removed, so a caption reading `다운스트림(downstream)` before the
    # dedupe and `다운스트림` after it no longer looks like a lost table. A
    # float the expander could not convert is still `\begin{...}` here, so
    # nothing that matters escapes the check.
    lost = algorithm_float.check_latex_float_fidelity(md_text, body_content)
    # A picture DRAWN IN CODE has no image file anywhere in the source, and no
    # stage of this pipeline can produce one — the absence is a limitation, not
    # a defect, and failing the build over it only stops the other 99% of the
    # book from being made. Say what the reader loses and carry on.
    drawn = [row for row in lost if row[0].rstrip('*') in _CODE_DRAWN_FIGURES]
    lost = [row for row in lost if row not in drawn]
    if drawn:
        print("WARNING: %d figure(s) are drawn in TikZ/PGF code with no image "
              "file in the source, so they cannot be rendered and are absent "
              "from the book:" % len(drawn))
        for env, phrase in drawn:
            print("  \\begin{%s} ... %s" % (env, phrase))
    if lost:
        print("ERROR: %d raw LaTeX block(s) reached the markdown and left no "
              "trace in the HTML. pandoc drops raw LaTeX on the HTML path "
              "without warning, so this content is missing from the book:"
              % len(lost))
        for env, phrase in lost:
            print("  \\begin{%s} ... %s" % (env, phrase))
        if not allow_degraded:
            return False

    # Generate book_doc.html with ebook template.
    # The ebook template gets font_family_ebook (double-quoted CSS, which is
    # what Calibre's parser wants); until now that config key was never read.
    template_ebook = os.path.join(SCRIPT_DIR, 'template_ebook.html')
    book_doc_file = os.path.join(temp_dir, 'book_doc.html')
    ebook_cfg = {**lang_cfg,
                 'font_family': lang_cfg.get('font_family_ebook', lang_cfg['font_family'])}
    # The names have been in config.txt under `creator=` since conversion,
    # with nowhere to go: the book opened on its table of contents and
    # credited one of sixteen authors, in a <meta> tag no reader sees.
    book_cfg = load_config(temp_dir) or {}
    arxiv_id = (book_cfg.get('arxiv_id') or '').strip()
    title_page = build_title_page(
        title, source='arXiv:%s' % arxiv_id if arxiv_id else '')
    full_byline = (book_cfg.get('creator') or '').strip() or author
    apply_template_to_html(body_content, template_ebook, book_doc_file, title,
                           ebook_cfg, author, print_cfg=print_cfg,
                           title_page=title_page, byline=full_byline)

    # Generate book.html with web template
    template_web = os.path.join(SCRIPT_DIR, 'template.html')
    book_file = os.path.join(temp_dir, 'book.html')
    apply_template_to_html(body_content, template_web, book_file, title, lang_cfg, author)

    if not _check_generated_html_sanity(book_doc_file):
        return False
    if not _check_generated_html_sanity(book_file):
        return False

    print(f"Generated: output.html, book_doc.html, book.html")
    return True


# =============================================================================
# Step 6: Add TOC
# =============================================================================

def generate_heading_id(text, existing_ids):
    """Generate unique ID for heading"""
    base_id = re.sub(r'[^\w\s-]', '', text.lower())
    base_id = re.sub(r'[-\s]+', '-', base_id)
    base_id = base_id.strip('-')

    if not base_id:
        base_id = 'heading'

    heading_id = base_id
    counter = 1
    while heading_id in existing_ids:
        heading_id = f"{base_id}-{counter}"
        counter += 1

    return heading_id


def generate_simple_toc_html(toc_data):
    """Generate simple HTML for table of contents"""
    if not toc_data:
        return ""

    toc_html = '<ul>\n'
    current_level = 1

    for item in toc_data:
        level = item['level']
        text = item['text']
        heading_id = item['id']

        if level > current_level:
            while current_level < level:
                toc_html += '<li><ul>\n'
                current_level += 1
        elif level < current_level:
            while current_level > level:
                toc_html += '</ul></li>\n'
                current_level -= 1

        toc_html += f'<li><a href="#{heading_id}">{text}</a></li>\n'

    while current_level > 1:
        toc_html += '</ul></li>\n'
        current_level -= 1

    toc_html += '</ul>\n'
    return toc_html


def insert_toc_with_bs4(html_file):
    """Insert TOC using BeautifulSoup"""
    try:
        with open(html_file, 'r', encoding='utf-8') as f:
            html_content = f.read()
    except Exception as e:
        print(f"Error reading HTML file: {e}")
        return False

    soup = BeautifulSoup(html_content, 'html.parser')

    toc_data = []
    existing_ids = []

    for heading in soup.find_all(['h1', 'h2', 'h3', 'h4', 'h5', 'h6']):
        level = int(heading.name[1])
        text = heading.get_text().strip()
        if not text:
            continue

        heading_id = generate_heading_id(text, existing_ids)
        existing_ids.append(heading_id)
        heading['id'] = heading_id
        toc_data.append({'level': level, 'text': text, 'id': heading_id})

    if not toc_data:
        print("No headings found for TOC")
        return False

    toc_html = generate_simple_toc_html(toc_data)

    toc_content_div = soup.find('div', class_='toc-content')
    if toc_content_div:
        toc_content_div.clear()
        toc_soup = BeautifulSoup(toc_html, 'html.parser')
        toc_content_div.append(toc_soup)
        print(f"TOC inserted ({len(toc_data)} headings)")
    else:
        print("Warning: .toc-content div not found, TOC not inserted")
        return False

    try:
        with open(html_file, 'w', encoding='utf-8') as f:
            f.write(str(soup))
        return True
    except Exception as e:
        print(f"Error saving HTML file: {e}")
        return False


def insert_toc_with_regex(html_file):
    """Insert TOC using regex (fallback)"""
    try:
        with open(html_file, 'r', encoding='utf-8') as f:
            html_content = f.read()
    except Exception as e:
        print(f"Error reading HTML file: {e}")
        return False

    # Rewrite in one pass over the matches. Building `<h2>text</h2>` and
    # str.replace()-ing it finds nothing, because pandoc writes
    # `<h2 id="slug">`: every link in the floating TOC pointed at an anchor
    # that was never created, so the whole sidebar was dead on the web page.
    heading_pattern = r'<(h[1-6])([^>]*)>(.*?)</\1>'
    matches = list(re.finditer(heading_pattern, html_content,
                               re.IGNORECASE | re.DOTALL))
    if not matches:
        print("No headings found for TOC")
        return False

    toc_html = '<ul>\n'
    pieces, cursor = [], 0
    for i, m in enumerate(matches):
        tag, attrs, text = m.group(1), m.group(2), m.group(3)
        level = int(tag[1])
        clean_text = re.sub(r'<[^>]+>', '',
                            _ANNOTATION_RE.sub('', text)).strip()
        # Keep the anchor the heading already has -- other links may use it.
        existing = re.search(r'\bid="([^"]+)"', attrs)
        if existing:
            heading_id = existing.group(1)
            pieces.append(html_content[cursor:m.end()])
        else:
            heading_id = f"heading-{i+1}"
            pieces.append(html_content[cursor:m.start()])
            pieces.append(f'<{tag} id="{heading_id}"{attrs}>{text}</{tag}>')
        cursor = m.end()

        toc_html += '  ' * (level - 1)
        toc_html += f'<li><a href="#{heading_id}">{clean_text}</a></li>\n'

    pieces.append(html_content[cursor:])
    html_content = ''.join(pieces)
    toc_html += '</ul>\n'

    toc_content_pattern = r'(<div[^>]*class="toc-content[^"]*"[^>]*>).*?(</div>)'
    if re.search(toc_content_pattern, html_content, re.DOTALL):
        # A callable, not a template: toc_html carries whatever the
        # headings carry, and a heading with math in it ($\ell_2$) puts a
        # backslash into the replacement, where re reads it as an escape.
        html_content = re.sub(
            toc_content_pattern,
            lambda m: m.group(1) + toc_html + m.group(2),
            html_content,
            flags=re.DOTALL
        )
        print(f"TOC inserted ({len(matches)} headings)")
    else:
        print("Warning: .toc-content div not found")
        return False

    try:
        with open(html_file, 'w', encoding='utf-8') as f:
            f.write(html_content)
        return True
    except Exception as e:
        print(f"Error saving HTML file: {e}")
        return False


# The page-number slot in a print TOC entry. chromium_pdf renders once to find
# out which page each heading landed on, substitutes the real numbers, then
# renders again -- Chromium implements no `target-counter()`, so there is no
# way to ask CSS for "the page this link points at".
PRINT_TOC_SENTINEL = '\u00a7\u00a7%d\u00a7\u00a7'


def build_print_toc(html_content, toc_label='Contents', max_level=3):
    """Return (html_with_ids_and_toc, entry_count).

    book.html gets a floating sidebar TOC, which is a screen affordance and is
    hidden in print. book_doc.html -- the source for both EPUB and PDF -- had
    no TOC at all, so every generated book shipped without one.
    """
    headings = list(re.finditer(r'<(h[1-6])([^>]*)>(.*?)</\1>',
                                html_content, re.IGNORECASE | re.DOTALL))
    entries, pieces, cursor, n = [], [], 0, 0
    for m in headings:
        tag, attrs, inner = m.group(1), m.group(2), m.group(3)
        text = re.sub(r'<[^>]+>', '', _ANNOTATION_RE.sub('', inner)).strip()
        text = _html_lib.unescape(text)
        if not text:
            continue
        level = int(tag[1])
        # pandoc renders the document title as <h1 class="title"> inside
        # <header id="title-block-header">. A book's own title is not a TOC
        # entry, and listing it makes the first line point at itself.
        is_title = re.search(r'class="[^"]*\btitle\b', attrs or '') is not None
        n += 1
        heading_id = 'sec-%d' % n
        if 'id=' not in attrs.lower():
            pieces.append(html_content[cursor:m.start()])
            pieces.append('<%s id="%s"%s>%s</%s>' % (tag, heading_id, attrs, inner, tag))
            cursor = m.end()
        else:
            found = re.search(r'id="([^"]+)"', attrs)
            heading_id = found.group(1) if found else heading_id
        if level <= max_level and not is_title:
            entries.append({'level': level, 'text': text, 'id': heading_id})
    pieces.append(html_content[cursor:])
    html_content = ''.join(pieces)

    if not entries:
        return html_content, 0

    rows = ['<nav class="print-toc" role="doc-toc">',
            '<h1 class="print-toc-title">%s</h1>' % _html_lib.escape(toc_label),
            '<ul class="print-toc-list">']
    for i, e in enumerate(entries):
        rows.append(
            '<li class="toc-l%d"><a href="#%s">'
            '<span class="toc-text">%s</span>'
            '<span class="toc-dots"></span>'
            '<span class="toc-page" data-toc="%d">%s</span>'
            '</a></li>' % (e['level'], e['id'], _html_lib.escape(e['text']), i,
                           PRINT_TOC_SENTINEL % i))
    rows.append('</ul></nav>')
    toc_html = '\n'.join(rows)

    # After the title page, when there is one. Pinned to the opening <body>
    # tag the contents came first and the title page second, so the book
    # opened on its own table of contents with the title on the leaf behind
    # it -- which is how the title page looked absent even once it was built.
    title_page = re.search(r'<section class="title-page">.*?</section>',
                           html_content, re.DOTALL)
    if title_page:
        at = title_page.end()
        html_content = html_content[:at] + '\n' + toc_html + html_content[at:]
    elif re.search(r'<body[^>]*>', html_content, re.IGNORECASE):
        html_content = re.sub(r'(<body[^>]*>)', lambda mm: mm.group(1) + '\n' + toc_html,
                              html_content, count=1, flags=re.IGNORECASE)
    else:
        html_content = toc_html + html_content
    return html_content, len(entries)


def add_print_toc_to_ebook(temp_dir, toc_label='Contents'):
    """Insert the print/EPUB table of contents into book_doc.html."""
    path = os.path.join(temp_dir, 'book_doc.html')
    if not os.path.exists(path):
        print("Warning: book_doc.html not found, skipping print TOC")
        return False
    with open(path, 'r', encoding='utf-8') as fh:
        content = fh.read()
    if 'class="print-toc"' in content:
        return True
    content, count = build_print_toc(content, toc_label)
    if not count:
        print("No headings found for the print TOC")
        return False
    with open(path, 'w', encoding='utf-8') as fh:
        fh.write(content)
    print(f"Print TOC inserted into book_doc.html ({count} entries)")
    return True


def add_toc(temp_dir, lang_cfg=None):
    """Add the sidebar TOC to book.html and a print TOC to book_doc.html"""
    print("=== Adding Table of Contents ===")

    book_file = os.path.join(temp_dir, 'book.html')
    if not os.path.exists(book_file):
        print("Warning: book.html not found, skipping TOC")
        return False

    if BS4_AVAILABLE:
        ok = insert_toc_with_bs4(book_file)
    else:
        ok = insert_toc_with_regex(book_file)

    # book_doc.html feeds BOTH the EPUB and the PDF and used to get no TOC at
    # all, while this step still printed "TOC inserted" -- which read as though
    # every format had one.
    label = (lang_cfg or {}).get('toc_label', 'Contents')
    add_print_toc_to_ebook(temp_dir, label)
    return ok


# =============================================================================
# Step 7: Generate DOCX/EPUB/PDF with error transparency
# =============================================================================

def generate_format(html_file, temp_dir, output_ext, lang_attr, cover=None,
                    pdf_engine='chromium', print_cfg=None):
    """Generate a specific format.

    PDF goes to headless Chromium by default, because Calibre honours
    neither @page nor the print stylesheet and was silently producing
    US Letter pages with 25.4mm margins. Everything else still goes to
    calibre_html_publish.py.
    """
    output_file = os.path.join(temp_dir, f"book{output_ext}")
    cover = cover if output_ext == '.epub' else None
    if cover and not os.path.isfile(cover):
        print(f"Cover image not found: {cover}")
        return None

    if os.path.exists(output_file):
        output_mtime = os.path.getmtime(output_file)

        # Check if source HTML is newer
        html_newer = os.path.getmtime(html_file) > output_mtime

        # Check if any image asset is newer (Calibre embeds these)
        images_newer = False
        images_dir = os.path.join(temp_dir, 'images')
        if os.path.isdir(images_dir):
            for img in os.listdir(images_dir):
                img_path = os.path.join(images_dir, img)
                if os.path.isfile(img_path) and os.path.getmtime(img_path) > output_mtime:
                    images_newer = True
                    break

        cover_newer = bool(cover and os.path.getmtime(cover) > output_mtime)

        if not html_newer and not images_newer and not cover_newer:
            file_size = os.path.getsize(output_file)
            print(f"Skipping {output_ext} - already exists and up to date ({file_size:,} bytes)")
            return output_file
        else:
            reasons = []
            if html_newer:
                reasons.append("source HTML changed")
            if images_newer:
                reasons.append("image assets changed")
            if cover_newer:
                reasons.append("cover image changed")
            print(f"Rebuilding {output_ext} - {', '.join(reasons)}")

    if output_ext == '.pdf' and pdf_engine == 'chromium':
        # Rendered in place: book_doc.html sits next to images/, so a
        # relative <img src="images/x.png"> resolves with no copying,
        # no work.html and no temp directory to clean up.
        def _render(src, out):
            return chromium_pdf.html_to_pdf(src, out, lang=lang_attr,
                                            profile=print_cfg)

        # A formula as wide as the text column prints under its own number,
        # and no stylesheet can prevent it: MathML compresses its spacing to
        # the box rather than shrinking into it. The size that clears the
        # number cannot be known without laying the page out, so this renders,
        # measures what it rendered, and renders again if anything collided.
        # A book with no collision costs exactly one render, as before.
        ok, collided = equation_fit.fit_equations(html_file, output_file,
                                                  _render, print_cfg)
        if collided:
            print("ERROR: equation %s still prints under its own number at "
                  "the smallest size tried." % ', '.join(collided))
            print("  The formula fills the text column; nothing downstream "
                  "can separate them, and the number is unreadable on the "
                  "page while every content check passes.")
            raise SystemExit(1)
        return output_file if ok and os.path.exists(output_file) else None

    publish_script = os.path.join(SCRIPT_DIR, "calibre_html_publish.py")
    if not os.path.exists(publish_script):
        print(f"calibre_html_publish.py not found at: {publish_script}")
        return None

    try:
        # sys.executable, not "python3": on Windows only `python`/`py` resolve,
        # so a hardcoded python3 fails every format from cmd/PowerShell.
        cmd = [sys.executable, publish_script, html_file, "-o", output_file,
               "--lang", lang_attr]
        if cover:
            cmd.extend(["--cover", cover])
        env = {**os.environ, 'PYTHONIOENCODING': 'utf-8'}
        result = subprocess.run(cmd, check=True, capture_output=True, text=True,
                                encoding='utf-8', errors='replace', env=env)

        if os.path.exists(output_file):
            file_size = os.path.getsize(output_file)
            return output_file
        else:
            print(f"Failed to generate {output_ext}")
            if result.stdout:
                print(f"  stdout: {result.stdout[-500:]}")
            return None
    except subprocess.CalledProcessError as e:
        print(f"Failed to generate {output_ext}")
        if e.stdout:
            print(f"  stdout: {e.stdout[-500:]}")
        if e.stderr:
            print(f"  stderr: {e.stderr[-500:]}")
        return None
    except Exception as e:
        print(f"Error generating {output_ext}: {e}")
        return None


def build_reference_docx(temp_dir, lang_cfg, print_cfg):
    """Produce a pandoc --reference-doc matching the print profile.

    Without one, pandoc uses its built-in default: Calibri 11pt on US Letter
    with 1-inch margins and no East Asian font mapping at all, so a Korean
    translation opened in Word came out in whatever fallback face Word chose.

    Starts from pandoc's own default so every style pandoc writes into
    (Heading N, Source Code, Table Caption, ...) definitely exists, then
    overrides page setup and fonts. Returns the path, or None if unavailable.
    """
    try:
        import docx
        from docx.shared import Pt, Mm
        from docx.oxml.ns import qn
    except ImportError:
        print("  (python-docx not installed — DOCX keeps pandoc's default styling)")
        return None

    pandoc = resolve_pandoc()
    if not pandoc:
        return None

    ref_path = os.path.join(temp_dir, 'reference.docx')
    try:
        with open(ref_path, 'wb') as fh:
            result = subprocess.run(
                [pandoc, '--print-default-data-file', 'reference.docx'],
                stdout=fh, stderr=subprocess.PIPE, timeout=120)
        if result.returncode != 0 or not os.path.getsize(ref_path):
            return None

        latin, cjk = layout.docx_fonts(lang_cfg)
        body_pt = print_cfg.get('base_font_size_pt', 11.5)
        line_height = print_cfg.get('line_height', 1.5)
        width_mm, height_mm = layout.page_size_mm(print_cfg)

        doc = docx.Document(ref_path)

        section = doc.sections[0]
        section.page_width = Mm(width_mm)
        section.page_height = Mm(height_mm)
        section.left_margin = Mm(print_cfg.get('margin_left_mm', 18))
        section.right_margin = Mm(print_cfg.get('margin_right_mm', 18))
        section.top_margin = Mm(print_cfg.get('margin_top_mm', 18))
        section.bottom_margin = Mm(print_cfg.get('margin_bottom_mm', 22))

        def set_fonts(style, latin_face, cjk_face, size_pt=None, bold=None):
            font = style.font
            font.name = latin_face
            if size_pt is not None:
                font.size = Pt(size_pt)
            if bold is not None:
                font.bold = bold
            rpr = style.element.get_or_add_rPr()
            rfonts = rpr.get_or_add_rFonts()
            for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
                rfonts.set(qn(attr), latin_face)
            # The one Word needs for Hangul/CJK runs, and the one pandoc's
            # default reference doc never sets.
            rfonts.set(qn('w:eastAsia'), cjk_face)

        names = {s.name for s in doc.styles}

        if 'Normal' in names:
            normal = doc.styles['Normal']
            set_fonts(normal, latin, cjk, body_pt)
            normal.paragraph_format.line_spacing = line_height
            normal.paragraph_format.space_after = Pt(0)

        # 제목 고딕 / 본문 명조: the same pairing the print sheet uses.
        heading_latin, heading_cjk = layout.docx_heading_fonts(lang_cfg)
        ladder = {'Title': body_pt * 1.9, 'Heading 1': body_pt * 1.75,
                  'Heading 2': body_pt * 1.45, 'Heading 3': body_pt * 1.22,
                  'Heading 4': body_pt * 1.05, 'Heading 5': body_pt,
                  'Heading 6': body_pt}
        for style_name, size in ladder.items():
            if style_name in names:
                set_fonts(doc.styles[style_name], heading_latin, heading_cjk,
                          round(size, 1), bold=True)

        for mono_style in ('Source Code', 'Verbatim Char'):
            if mono_style in names:
                set_fonts(doc.styles[mono_style], 'Consolas', 'DotumChe',
                          round(body_pt * 0.85, 1))

        doc.save(ref_path)
        return ref_path
    except Exception as e:
        print(f"  (could not build a DOCX reference doc: {e})")
        try:
            os.remove(ref_path)
        except OSError:
            pass
        return None



_W_TBL_RE = re.compile(r'<w:tbl>.*?</w:tbl>', re.DOTALL)
_W_TR_RE = re.compile(r'<w:tr\b[^>]*>.*?</w:tr>', re.DOTALL)
_W_TRPR_RE = re.compile(r'<w:trPr>(.*?)</w:trPr>', re.DOTALL)
# Everything the schema puts BEFORE w:tblHeader inside w:trPr. Word rejects
# the file outright if a child turns up out of order.
_W_BEFORE_HEADER_RE = re.compile(
    r'(?:<w:cnfStyle\b[^>]*/?>|<w:divId\b[^>]*/?>|<w:gridBefore\b[^>]*/?>'
    r'|<w:gridAfter\b[^>]*/?>|<w:wBefore\b[^>]*/?>|<w:wAfter\b[^>]*/?>'
    r'|<w:cantSplit\b[^>]*/?>|<w:trHeight\b[^>]*/?>)*')


def _mark_row_as_header(row):
    """Add <w:tblHeader/> to one <w:tr>, in schema order."""
    if 'w:tblHeader' in row:
        return row, False
    body = _W_TRPR_RE.search(row)
    if body:
        head = _W_BEFORE_HEADER_RE.match(body.group(1))
        at = body.start(1) + head.end()
        return row[:at] + '<w:tblHeader/>' + row[at:], True
    at = row.index('>') + 1
    return row[:at] + '<w:trPr><w:tblHeader/></w:trPr>' + row[at:], True


def mark_docx_header_rows(docx_path, temp_dir):
    """Repeat each table's header rows on every page in Word. (marked, tables).

    pandoc marks one header row or none, so a multi-deck header -- the case
    the reader most needs repeated -- was never repeated at all.
    """
    plans = table_structures(temp_dir)
    if not plans or not os.path.isfile(docx_path):
        return 0, 0
    try:
        with zipfile.ZipFile(docx_path) as zf:
            names = zf.namelist()
            blobs = {name: zf.read(name) for name in names}
    except (zipfile.BadZipFile, OSError):
        return 0, 0
    if 'word/document.xml' not in blobs:
        return 0, 0
    doc = blobs['word/document.xml'].decode('utf-8')

    tables = _W_TBL_RE.findall(doc)
    if len(tables) != len(plans):
        # Matched by position, so a disagreement means the mapping cannot be
        # trusted. Leaving Word as pandoc left it is the safe answer.
        sys.stderr.write(
            'warning: %d table(s) in the source, %d in the DOCX; Word header '
            'rows left as pandoc set them\n' % (len(plans), len(tables)))
        return 0, len(tables)

    marked, index = [0], [0]

    def fix(match):
        table = match.group(0)
        want = plans[index[0]].get('header') or 0
        index[0] += 1
        if want < 1:
            return table
        rows, count = [], [0]

        def row_fix(row_match):
            if count[0] >= want:
                return row_match.group(0)
            count[0] += 1
            row, changed = _mark_row_as_header(row_match.group(0))
            if changed:
                marked[0] += 1
            return row

        rows.append(_W_TR_RE.sub(row_fix, table))
        return rows[0]

    doc = _W_TBL_RE.sub(fix, doc)
    if not marked[0]:
        return 0, len(tables)

    blobs['word/document.xml'] = doc.encode('utf-8')
    tmp = docx_path + '.tmp'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as out:
        for name in names:
            out.writestr(name, blobs[name])
    os.replace(tmp, docx_path)
    return marked[0], len(tables)


def generate_docx_with_pandoc(temp_dir, title, author, lang_attr,
                              lang_cfg=None, print_cfg=None):
    """Build book.docx straight from output.md via pandoc.

    Calibre's DOCX writer has no math support whatsoever (its from_html class
    exposes no math handling), so routing DOCX through Calibre silently drops
    every equation. pandoc emits native OMML — real, editable Word equations.
    """
    pandoc = resolve_pandoc()
    if not pandoc:
        return None

    # prepared.md is output.md with citations and cross-references resolved.
    # DOCX is built straight from markdown, so without this it would still show
    # raw [@key] and (fig:x) markers that the HTML path had already fixed.
    md_file = os.path.join(temp_dir, 'prepared.md')
    if not os.path.exists(md_file):
        md_file = os.path.join(temp_dir, 'output.md')
    out_file = os.path.join(temp_dir, 'book.docx')
    if not os.path.exists(md_file):
        return None

    # pandoc runs with cwd=temp_dir so that relative `images/x.png` refs
    # resolve and get embedded into the .docx. That means the paths handed to
    # pandoc must be relative to temp_dir too — passing `temp_dir/output.md`
    # there makes pandoc look for `temp_dir/temp_dir/output.md`, which fails
    # and silently downgrades the whole DOCX to the math-less Calibre path.
    cmd = [
        pandoc, os.path.basename(md_file), '-o', os.path.basename(out_file),
        '--from', pandoc_from(lang_attr), '--to', 'docx',
        '--standalone', '--toc', '--toc-depth=3',
        '--metadata', f'title={title}',
        '--metadata', f'author={author}',
        '--metadata', f'lang={lang_attr}',
        '--resource-path', '.',
        '--wrap=preserve',
    ]

    # Page setup and CJK fonts. Without this pandoc falls back to Calibri 11pt
    # on Letter with 1-inch margins and no eastAsia mapping.
    reference = build_reference_docx(temp_dir, lang_cfg or {},
                                     print_cfg or layout.get_print_profile())
    if reference:
        cmd.extend(['--reference-doc', os.path.basename(reference)])
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, cwd=temp_dir,
                                encoding='utf-8', errors='replace', timeout=600)
        if result.returncode != 0:
            print(f"pandoc docx failed: {(result.stderr or '')[:1200]}")
            return None
        return out_file if os.path.exists(out_file) else None
    except (OSError, subprocess.TimeoutExpired) as e:
        print(f"pandoc docx error: {e}")
        return None


def generate_formats(temp_dir, lang_attr, cover=None, title=None, author=None,
                     docx_engine='pandoc', pdf_engine='chromium', print_cfg=None,
                     lang_cfg=None):
    """Generate DOCX, EPUB, and PDF with result summary"""
    print("=== Generating output formats ===")

    html_file = os.path.join(temp_dir, "book_doc.html")
    if not os.path.exists(html_file):
        html_files = glob.glob(os.path.join(temp_dir, "*.html"))
        if html_files:
            html_file = max(html_files, key=os.path.getmtime)
        else:
            print("No HTML files found for format generation")
            return

    results = {}

    if docx_engine == 'pandoc' and resolve_pandoc():
        docx = generate_docx_with_pandoc(temp_dir, title or 'Translated Book',
                                         author or 'Unknown Author', lang_attr,
                                         lang_cfg=lang_cfg, print_cfg=print_cfg)
        if docx:
            marked, seen = mark_docx_header_rows(docx, temp_dir)
            if marked:
                print(f"  Word header rows repeated: {marked} row(s) "
                      f"across {seen} table(s)")
            results['.docx'] = ('OK', f"{os.path.getsize(docx):,} bytes (pandoc/OMML)")
        else:
            print("pandoc docx failed — falling back to Calibre (math will be lost)")
            docx = generate_format(html_file, temp_dir, '.docx', lang_attr)
            results['.docx'] = (('OK', f"{os.path.getsize(docx):,} bytes")
                                if docx else ('FAILED', ''))
    else:
        docx = generate_format(html_file, temp_dir, '.docx', lang_attr)
        results['.docx'] = (('OK', f"{os.path.getsize(docx):,} bytes")
                            if docx else ('FAILED', ''))

    for ext in ['.epub', '.pdf']:
        result = generate_format(html_file, temp_dir, ext, lang_attr, cover=cover,
                                 pdf_engine=pdf_engine, print_cfg=print_cfg)
        if result:
            file_size = os.path.getsize(result)
            detail = f"{file_size:,} bytes"
            if ext == '.pdf':
                detail += f" ({pdf_engine})"
            results[ext] = ('OK', detail)
        else:
            results[ext] = ('FAILED', '')

    # Print summary table
    print("\nFormat results:")
    has_failures = False
    for ext, (status, detail) in results.items():
        if status == 'OK':
            print(f"  {ext}: {status} ({detail})")
        else:
            print(f"  {ext}: {status}")
            has_failures = True

    return not has_failures


def _validate_export_name(name):
    """Validate an export filename stem. Keep aliases inside temp_dir."""
    if not name or not name.strip():
        raise ValueError("--export-name must not be empty")
    if '\x00' in name or '/' in name or '\\' in name:
        raise ValueError("--export-name must be a filename stem, not a path")
    return name.strip()


def export_named_aliases(temp_dir, export_name):
    """Copy canonical outputs to optional user-facing filenames.

    Canonical artifacts remain untouched. The alias names use export_name as a
    filename stem, with book_doc.html receiving a _doc suffix to avoid colliding
    with the web HTML alias.
    """
    stem = _validate_export_name(export_name)
    mappings = {
        "book.html": f"{stem}.html",
        "book_doc.html": f"{stem}_doc.html",
        "book.docx": f"{stem}.docx",
        "book.epub": f"{stem}.epub",
        "book.pdf": f"{stem}.pdf",
    }
    copied = []
    for src_name, dst_name in mappings.items():
        src = os.path.join(temp_dir, src_name)
        if not os.path.exists(src):
            continue
        dst = os.path.join(temp_dir, dst_name)
        if os.path.abspath(src) == os.path.abspath(dst):
            continue
        shutil.copy2(src, dst)
        copied.append(dst_name)
    return copied


# =============================================================================
# Main
# =============================================================================

def main():
    parser = argparse.ArgumentParser(description='Merge translated pages and build final outputs')
    parser.add_argument('--temp-dir', required=True, help='Temp directory path')
    parser.add_argument('--title', default=None, help='Translated book title (override config)')
    parser.add_argument('--author', default=None, help='Author name (override config)')
    parser.add_argument('--lang', default=None, help='Output language code (override config)')
    parser.add_argument('--cover', default=None, help='Cover image path for EPUB output')
    parser.add_argument('--export-name', default=None, help='Optional filename stem for exported alias copies')
    parser.add_argument('--cleanup', action='store_true', help='Remove intermediate artifacts after successful build')
    parser.add_argument('--build-only', action='store_true',
                        help='Skip the merge step and build from the existing output.md')
    parser.add_argument('--force-html', action='store_true',
                        help='Always regenerate HTML even if book_doc.html looks fresh')
    parser.add_argument('--allow-degraded-html', action='store_true',
                        help='Permit the regex fallback, which loses tables and mangles math')
    parser.add_argument('--math', choices=['mathml', 'none'], default='mathml',
                        help='Math rendering mode for HTML output (default: mathml)')
    parser.add_argument('--docx-engine', choices=['pandoc', 'calibre'], default='pandoc',
                        help='DOCX generator. pandoc gives native editable equations; '
                             'calibre drops math entirely (default: pandoc)')
    parser.add_argument('--pdf-engine', choices=['chromium', 'calibre'],
                        default='chromium',
                        help='PDF renderer. chromium uses the local Chrome/Edge '
                             'print engine and honours the @page print CSS; '
                             'calibre is the legacy path and ignores it '
                             '(default: chromium)')
    parser.add_argument('--section-break', dest='section_break',
                        action='store_true', default=None,
                        help='Start every top-level heading on a new page. '
                             'Off by default: in this pipeline h1 is every '
                             'section, so it costs ~20%% more pages')
    parser.add_argument('--no-section-break', dest='section_break',
                        action='store_false',
                        help='Explicitly keep sections running on (the default)')
    parser.add_argument('--print-profile', choices=sorted(layout.PRINT_PROFILES),
                        default=layout.DEFAULT_PRINT_PROFILE,
                        help='Page geometry and base type size for the PDF. '
                             'Changing it needs --force-html, because HTML '
                             'regeneration is keyed on output.md mtime '
                             f'(default: {layout.DEFAULT_PRINT_PROFILE})')

    args = parser.parse_args()
    temp_dir = args.temp_dir

    if not os.path.isdir(temp_dir):
        print(f"Error: Temp directory not found: {temp_dir}")
        sys.exit(1)

    cover = args.cover
    if cover:
        if not os.path.isfile(cover):
            print(f"Error: Cover image not found: {cover}")
            sys.exit(1)
        cover = os.path.abspath(cover)

    export_name = None
    if args.export_name:
        try:
            export_name = _validate_export_name(args.export_name)
        except ValueError as e:
            print(f"Error: {e}")
            sys.exit(1)

    # Load config as base, CLI args override
    config = load_config(temp_dir)

    lang_code = args.lang or config.get('output_lang', 'zh')
    lang_cfg = get_lang_config(lang_code)
    print_cfg = layout.get_print_profile(
        args.print_profile,
        None if args.section_break is None else {'section_break': args.section_break})

    title = args.title or config.get('original_title', 'Translated Book')
    author = args.author or config.get('creator', 'Unknown Author')

    print(f"=== Merge and Build ===")
    print(f"Temp directory: {temp_dir}")
    print(f"Title: {title}")
    print(f"Author: {author}")
    print(f"Language: {lang_code} (attr: {lang_cfg['lang_attr']})")
    if args.pdf_engine == 'chromium':
        print(f"PDF: {args.print_profile} "
              f"({print_cfg['page_size']}, {layout.page_margin_css(print_cfg)}, "
              f"{print_cfg['base_font_size_pt']:g}pt)")

    # Step 4: Merge
    if args.build_only:
        if not os.path.exists(os.path.join(temp_dir, 'output.md')):
            print("Error: --build-only requires an existing output.md")
            sys.exit(1)
        print("=== Skipping merge (--build-only) ===")
    elif not merge_markdown_files(temp_dir):
        sys.exit(1)

    if not check_image_refs_resolve(temp_dir):
        sys.exit(1)

    # Step 5: Convert to HTML
    if not convert_md_to_html(temp_dir, title, lang_cfg, author,
                              allow_degraded=args.allow_degraded_html,
                              math_mode=args.math,
                              force=args.force_html or args.build_only,
                              print_cfg=print_cfg):
        sys.exit(1)

    # Step 6: Add TOC
    add_toc(temp_dir, lang_cfg)

    # Step 7: Generate formats
    all_formats_ok = generate_formats(temp_dir, lang_cfg['lang_attr'], cover=cover,
                                      title=title, author=author,
                                      docx_engine=args.docx_engine,
                                      pdf_engine=args.pdf_engine,
                                      print_cfg=print_cfg,
                                      lang_cfg=lang_cfg)

    if export_name:
        if all_formats_ok:
            aliases = export_named_aliases(temp_dir, export_name)
            if aliases:
                print("\nExport aliases:")
                for name in aliases:
                    print(f"  {name}")
        else:
            print("\nSkipping export aliases — some formats failed.")

    print("\n=== Build Complete ===")
    print(f"All outputs saved to: {temp_dir}")

    # The corpus census grows here, and only here. Made a step someone has to
    # remember, it would be skipped — and an advisor whose evidence stops
    # growing is back to guessing about the paper in front of it.
    try:
        import corpus_census
        corpus_census.record(temp_dir)
    except Exception as exc:                      # never fail a good build
        print(f"Corpus census: not recorded ({exc})")

    # And the referee, for the same reason and with better evidence against it:
    # its tally sat at ten runs while nine books were rebuilt and one was
    # re-translated, because calling it was left to whoever ran the build and
    # nobody did. An advisor you have to remember to consult is a document, not
    # an advisor. It speaks here whether or not anyone asked.
    try:
        import referee
        run = referee.collect(temp_dir, lang_code)
        if run['chunks']:
            data = referee.load()
            # Keyed on the edition, not the paper. Keying on the paper alone
            # let one book's second language replace its first, and the row it
            # erased was the one carrying a brief fault. `referee.py` was
            # fixed for that and this inline copy was not, which is K114's
            # shape a second time in one session.
            history = [r for r in data['runs']
                       if referee.edition_of(r) != referee.edition_of(run)]
            lines, flags = referee.judge(run, history)
            print()
            for line in lines:
                print(line)
            for kind, key, n, total in flags:
                if kind == 'brief':
                    print("REFEREE/BRIEF: `%s` fired on %d of %d chunks. Every "
                          "instance of a role reads the same prompt; fix the "
                          "prompt before you fault the agents." % (key, n, total))
                else:
                    print("REFEREE/CHRONIC: `%s` has now fired in %d runs. "
                          "Nobody has fixed it — it belongs in KNOWLEDGE, not "
                          "in another re-translation." % (key, n))
            # One row per paper: a re-run replaces its own row rather than
            # voting twice.
            data['runs'] = [r for r in data['runs']
                            if referee.edition_of(r) != referee.edition_of(run)]
            data['runs'].append(run)
            referee.save(data)
    except Exception as exc:                      # never fail a good build
        print(f"Referee: not recorded ({exc})")

    # An advisor nobody consults leaves no trace, so nobody — including the
    # agent meant to be calling it — can tell it is being skipped. Say it here.
    try:
        import advisors
        note = advisors.build_note()
        if note:
            print(note)
    except Exception as exc:                      # never fail a good build
        print(f"Advisors: status unavailable ({exc})")

    # List generated files
    for ext in ['book.html', 'book_doc.html', 'book.docx', 'book.epub', 'book.pdf']:
        filepath = os.path.join(temp_dir, ext)
        if os.path.exists(filepath):
            size = os.path.getsize(filepath)
            print(f"  {ext}: {size:,} bytes")

    # Cleanup intermediate artifacts if requested (skip if any format failed)
    if args.cleanup:
        if all_formats_ok:
            cleanup_intermediate_files(temp_dir)
        else:
            print("\nSkipping cleanup — some formats failed. Intermediate files kept for diagnosis/retry.")


def cleanup_intermediate_files(temp_dir):
    """Remove intermediate artifacts, keeping only final outputs."""
    print("\n=== Cleaning up intermediate files ===")

    removed = []

    # Remove chunk*.md and output_chunk*.md.
    # NOTE: chunk*.math.json sidecars are deliberately KEPT — output.md is
    # already restored, but a later --build-only re-merge would need them.
    for pattern in ['chunk*.md', 'output_chunk*.md']:
        for filepath in glob.glob(os.path.join(temp_dir, pattern)):
            os.remove(filepath)
            removed.append(os.path.basename(filepath))

    # Remove specific intermediate files
    for name in ['input.html', 'input.md', 'output.html']:
        filepath = os.path.join(temp_dir, name)
        if os.path.exists(filepath):
            os.remove(filepath)
            removed.append(name)

    if removed:
        print(f"Removed {len(removed)} intermediate file(s):")
        # Summarize chunk files instead of listing each one
        chunk_files = [f for f in removed if 'chunk' in f]
        other_files = [f for f in removed if 'chunk' not in f]
        if chunk_files:
            print(f"  {len(chunk_files)} chunk files (chunk*.md, output_chunk*.md)")
        for f in other_files:
            print(f"  {f}")
    else:
        print("No intermediate files to remove.")


if __name__ == "__main__":
    main()
